"""Structural route must reach publication without a second Markdown extraction."""
from openpyxl import Workbook

from app.db.models import DocumentGenerationState, DocumentTable, DocumentTableRow, OkfConcept, DocumentChunk
from app.db.session import session_scope
from tests.test_generation_pipeline import pipeline_env as _pipeline_env

pipeline_env = _pipeline_env


def test_excel_pipeline_publishes_full_cells_and_exact_refs(pipeline_env, tmp_path):
    pipeline, _, _ = pipeline_env
    pipeline.settings.okf_table_llm_classify = False
    def forbidden(*args, **kwargs):
        raise AssertionError('Structural Excel must not call the prose/Markdown generator')
    pipeline.okf_generator.generate_chunk = forbidden
    doc_id = 'structure-pipeline'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Wide'
    ws.append([f'Column {i}' for i in range(1, 317)])
    ws.append([f'value-{i}' for i in range(1, 317)])
    ws.cell(2, 2, 0)
    ws.cell(2, 3, False)
    source = tmp_path / 'wide.xlsx'
    wb.save(source)
    pipeline.registry.create(doc_id, 'wide.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', source.stat().st_size)
    pipeline._process(doc_id, source, 'wide.xlsx', [], resume=False)
    assert pipeline.registry.get(doc_id)['status'] == 'done'
    with session_scope() as s:
        active = s.get(DocumentGenerationState, doc_id).active_generation_id
        tables = s.query(DocumentTable).filter_by(doc_id=doc_id).all()
        assert len(tables) == 1 and tables[0].generation_id == active
        assert s.query(DocumentTableRow).filter_by(doc_id=doc_id).count() == 2
        concepts = s.query(OkfConcept).filter_by(doc_id=doc_id).all()
        chunks = s.query(DocumentChunk).filter_by(doc_id=doc_id).all()
        assert concepts and all(c.table_refs for c in concepts)
        assert chunks and all(c.table_refs for c in chunks)
        assert any('value-316' in c.content for c in concepts)
        assert {(r['row_number'], col) for c in concepts for r in c.table_refs for col in r['columns']} == {(row, col) for row in (1, 2) for col in range(1, 317)}
        assert all(r['generation_id'] == active for c in concepts for r in c.table_refs)
    bundle = pipeline.settings.okf_dir / doc_id / 'generations' / active
    assert (bundle / 'tables' / 'manifest.json').is_file()


def test_table_chunk_index_failure_preserves_previous_publication(pipeline_env,tmp_path,monkeypatch):
    from app.services.errors import VectorStoreError
    pipeline,_,_=pipeline_env
    pipeline.settings.okf_table_llm_classify=False
    pipeline.settings.llm_chunk_retry_attempts=1
    source=tmp_path/'swap.xlsx';wb=Workbook();ws=wb.active;ws.append(['Code','Value']);ws.append(['K',7]);wb.save(source)
    doc_id='d123456789abcdef'
    pipeline.registry.create(doc_id,'swap.xlsx','application/octet-stream',source.stat().st_size)
    pipeline._process(doc_id,source,'swap.xlsx',[],resume=False)
    with session_scope() as s: before=s.get(DocumentGenerationState,doc_id).active_generation_id
    ws.cell(2,2,9);wb.save(source)
    def fail(*args,**kw):raise VectorStoreError('injected chunk upsert failure')
    original_index=pipeline.vector_store.index_chunks
    monkeypatch.setattr(pipeline.vector_store,'index_chunks',fail)
    pipeline._process(doc_id,source,'swap.xlsx',[],resume=False)
    with session_scope() as s: assert s.get(DocumentGenerationState,doc_id).active_generation_id==before
    assert pipeline.registry.get(doc_id)['status']=='paused'
    monkeypatch.setattr(pipeline.vector_store,'index_chunks',original_index)
    pipeline._process(doc_id,source,'swap.xlsx',[],resume=True)
    assert pipeline.registry.get(doc_id)['status']=='done'
    with session_scope() as s:
        active=s.get(DocumentGenerationState,doc_id).active_generation_id
        assert active != before
        assert s.query(DocumentTableRow).filter_by(doc_id=doc_id,generation_id=active).count()==2
        assert all(r['generation_id']==active for chunk in s.query(DocumentChunk).filter_by(doc_id=doc_id) for r in chunk.table_refs)


def test_supervised_excel_spool_roundtrips_into_candidate(pipeline_env,tmp_path):
    pipeline,_,_=pipeline_env
    pipeline.settings.parser_supervisor_enabled=True
    pipeline.settings.okf_table_llm_classify=False
    source=tmp_path/'worker.xlsx';wb=Workbook();ws=wb.active;ws.append(['Code','Value']);ws.append(['K',0]);wb.save(source)
    doc_id='e123456789abcdef';pipeline.registry.create(doc_id,'worker.xlsx','application/octet-stream',source.stat().st_size)
    pipeline._process(doc_id,source,'worker.xlsx',[],resume=False)
    assert pipeline.registry.get(doc_id)['status']=='done'
    with session_scope() as s:
        state=s.get(DocumentGenerationState,doc_id)
        assert s.query(DocumentTable).filter_by(doc_id=doc_id,generation_id=state.active_generation_id).one().quality['cell_count']==4
        assert s.query(DocumentTableRow).filter_by(doc_id=doc_id).count()==2


def test_structural_repair_carries_tables_and_rebinds_refs(pipeline_env,tmp_path):
    from app.services.canonical_repair import repair_published_document
    pipeline,_,_=pipeline_env;pipeline.settings.okf_table_llm_classify=False
    source=tmp_path/'repair.xlsx';wb=Workbook();ws=wb.active;ws.append(['Code','Value']);ws.append(['K',7]);wb.save(source)
    doc_id='f123456789abcdef';pipeline.registry.create(doc_id,'repair.xlsx','application/octet-stream',source.stat().st_size)
    pipeline._process(doc_id,source,'repair.xlsx',[],resume=False)
    def edit(snapshot):
        snapshot['concepts'][0]['title']+=' repaired'
    result=repair_published_document(doc_id,pipeline.settings,pipeline.okf_generator,pipeline.embedder,pipeline.vector_store,edit)
    with session_scope() as s:
        generation=result['generation_id'];table=s.query(DocumentTable).filter_by(doc_id=doc_id,generation_id=generation).one()
        assert table.quality['cell_count']==4
        assert s.query(DocumentTableRow).filter_by(doc_id=doc_id,generation_id=generation).count()==2
        for concept in s.query(OkfConcept).filter_by(doc_id=doc_id):
            assert concept.table_refs and all(r['generation_id']==generation for r in concept.table_refs)
        for chunk in s.query(DocumentChunk).filter_by(doc_id=doc_id):
            assert chunk.table_refs and all(r['generation_id']==generation for r in chunk.table_refs)

def test_structural_embedding_budget_includes_the_title(pipeline_env,tmp_path,monkeypatch):
    pipeline,_,_=pipeline_env
    pipeline.settings.okf_table_llm_classify=False
    pipeline.settings.table_embedding_tokens=512
    pipeline.settings.table_context_tokens=512
    original=pipeline.embedder.embed_texts
    inputs=[]
    def bounded(texts):
        inputs.extend(texts)
        assert all(len(text.encode('utf8')) <= 512 for text in texts)
        return original(texts)
    monkeypatch.setattr(pipeline.embedder,'embed_texts',bounded)
    source=tmp_path/'budget.xlsx';wb=Workbook();ws=wb.active;ws.title='Budgets';ws.append(['Code','Text']);ws.append(['K','😀'*400]);wb.save(source)
    doc='c123456789abcdef';pipeline.registry.create(doc,'budget.xlsx','application/octet-stream',source.stat().st_size)
    pipeline._process(doc,source,'budget.xlsx',[],resume=False)
    assert pipeline.registry.get(doc)['status']=='done'
    assert inputs and max(len(text.encode('utf8')) for text in inputs)<=512


def test_excel_pipeline_uses_one_based_structural_classifier_prompt(pipeline_env, tmp_path):
    pipeline, _, _ = pipeline_env
    pipeline.settings.okf_table_llm_classify = True
    seen = []
    def classify(system, user, **kwargs):
        seen.append(system)
        assert '0-based' not in system
        assert 'title_col=0' not in system
        assert '"title_col": 0' not in system
        assert 'one-based' in system
        return {'concept_per_row': True, 'title_col': 1, 'concept_type': 'reference', 'extraction_mode': 'whole'}
    pipeline.okf_generator.llm.chat_json = classify
    source = tmp_path / 'classifier.xlsx'
    wb = Workbook(); ws = wb.active
    ws.append(['Code', 'Value']); ws.append(['K', 7]); wb.save(source)
    doc_id = 'a234567890abcdef'
    pipeline.registry.create(doc_id, 'classifier.xlsx', 'application/octet-stream', source.stat().st_size)
    pipeline._process(doc_id, source, 'classifier.xlsx', [], resume=False)
    assert seen
    with session_scope() as session:
        table = session.query(DocumentTable).filter_by(doc_id=doc_id).one()
        assert table.quality['classifier_cause'] is None
        assert table.quality['extraction_mode'] == 'whole'


def test_structural_classifier_uses_fixture_cache_not_shared_data(pipeline_env):
    import app.services.table_classification as classification
    pipeline, _, _ = pipeline_env
    assert classification.get_settings().cache_dir == pipeline.settings.cache_dir

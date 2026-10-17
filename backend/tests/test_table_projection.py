"""Projection completeness is by original cells and exact text ranges."""
import importlib
from types import SimpleNamespace

import pytest

from app.config import Settings
from app.services import field_table
from app.services.okf_generator import OKFGenerator


from docparser.table_model import CellRange, TableCell, TableColumn, TableDescriptor, TableRow


def module():
    assert importlib.util.find_spec('app.services.table_projection') is not None, 'projection module is missing'
    return importlib.import_module('app.services.table_projection')


def table(width=316, long_value='value'):
    from openpyxl.utils.cell import get_column_letter
    columns = [TableColumn(i, get_column_letter(i), ['Group', f'Property {i}']) for i in range(1, width + 1)]
    descriptor = TableDescriptor('root', 'a' * 32, 0, 'Directory', CellRange(0, 1, 1, 4, width),
                                 header_rows=[1], columns=columns, method='structural')
    cells = {i: TableCell(f'{get_column_letter(i)}2', long_value if i == width else f'v{i}',
                         long_value if i == width else f'v{i}', 'string') for i in range(1, width + 1)}
    cells[1] = TableCell('A2', None, '', 'string', state='blank')
    cells[2] = TableCell('B2', 0, '0', 'number')
    cells[3] = TableCell('C2', False, 'False', 'boolean')
    return descriptor, [TableRow(2, cells), TableRow(4, {i: __import__('dataclasses').replace(c, address=f'{get_column_letter(i)}4') for i, c in cells.items()})]


def test_preview_is_bounded_and_mapping_keeps_original_right_edge():
    m = module()
    descriptor, rows = table(long_value='z' * 50000)
    preview = m.build_classification_preview(descriptor, rows, m.ProjectionLimits())
    assert len(preview.column_index_map) <= 12
    assert 316 in preview.column_index_map and 1 in preview.column_index_map
    assert len(preview.text) <= 12000 and 'truncated' in preview.text.lower()
    assert m.BudgetEstimator().estimate(preview.text).tokens <= m.ProjectionLimits().classification_tokens


def test_wide_fragments_cover_every_nonblank_cell_and_repeat_right_key():
    m = module()
    descriptor, rows = table()
    limits = m.ProjectionLimits(embedding_tokens=900, context_tokens=700)
    fragments = list(m.iter_table_fragments(descriptor, rows, limits, generation_id='b' * 32, title_column=316))
    assert len(fragments) > 2 and len({f.fragment_id for f in fragments}) == len(fragments)
    actual = {(r.row_number, column) for f in fragments for r in f.coverage_refs for column in r.columns}
    expected = {(row, column) for row in (2, 4) for column in range(2, 317)}
    assert actual == expected
    assert all('value' in f.content and 'Directory' in f.content for f in fragments)
    assert all(m.BudgetEstimator().estimate(f.content).tokens <= 700 for f in fragments)
    assert any('False' in f.content for f in fragments) and any('B2' in f.content for f in fragments)


def test_long_cell_continuations_cover_exact_ranges_without_truncation():
    m = module()
    value = '<script>|\\nкириллица' * 500
    descriptor, rows = table(width=3, long_value=value)
    # The third cell in the synthetic helper is boolean; install the long cell independently.
    rows = [TableRow(2, {3: TableCell('C2', value, value, 'string')})]
    limits = m.ProjectionLimits(embedding_tokens=900, context_tokens=700)
    fragments = list(m.iter_table_fragments(descriptor, rows, limits, generation_id='b' * 32, title_column=2))
    refs = [r for f in fragments for r in f.coverage_refs if r.columns == [3]]
    assert len(refs) > 2 and refs[0].text_start == 0 and refs[-1].text_end == len(value)
    assert all(a.text_end == b.text_start for a, b in zip(refs, refs[1:]))
    assert all('<script>' not in f.content for f in fragments)
    assert all(m.BudgetEstimator().estimate(f.content).tokens <= 700 for f in fragments)


def test_formula_missing_is_searchable_and_leading_zeros_survive():
    m = module()
    descriptor, _ = table(width=2)
    rows = [TableRow(2, {1: TableCell('A2', 7, '00007', 'number', '00000'),
                         2: TableCell('B2', None, '', 'number', formula='1+1', state='formula_missing')})]
    fragments = list(m.iter_table_fragments(descriptor, rows, m.ProjectionLimits(), generation_id='b' * 32))
    assert any('00007' in f.content for f in fragments)
    assert any('formula_result_missing' in f.content for f in fragments)
    assert {(r.row_number, c) for f in fragments for r in f.coverage_refs for c in r.columns} == {(2, 1), (2, 2)}

def test_small_whole_table_is_one_fragment_but_large_whole_is_bounded():
    from app.services.table_projection import project_table, ProjectionLimits
    from tests.test_structural_classification import fixture_table
    descriptor, rows = fixture_table()
    fragments = list(project_table(descriptor, iter(rows), ProjectionLimits(), generation_id='a'*32, extraction_mode='whole'))
    assert len(fragments) == 1 and fragments[0].kind == 'table_whole'
    assert len(fragments[0].coverage_refs) == 2


def prompt_for(text):
    generator = OKFGenerator(llm=SimpleNamespace())
    return generator._build_prompt(text, "source.xlsx", 1, 1)


def test_generation_prompt_removes_excel_empty_tail_without_changing_checkpoint(tmp_path):
    width = 16_384
    text = "\n".join([
        "| Name | Value " + "| " * (width - 2) + "|",
        "|" + "---|" * width,
        "| Alpha | 0 " + "| " * (width - 2) + "|",
        "| Beta | False " + "| " * (width - 2) + "|",
    ])
    checkpoint = tmp_path / "chunk_118.md"
    checkpoint.write_text(text, encoding="utf-8")
    before = checkpoint.read_bytes()
    prompt = prompt_for(text)
    assert len(prompt) < 10_000
    assert "| Alpha | 0 |" in prompt
    assert "| Beta | False |" in prompt
    assert checkpoint.read_bytes() == before


def test_classifier_compacts_empty_tail_and_validates_transmitted_columns(tmp_path, monkeypatch):
    monkeypatch.setattr(field_table, "get_settings", lambda: Settings(_env_file=None, data_dir=tmp_path))
    header = ["Name", "Value"] + [""] * 200
    rows = ["| Alpha | value " + "| " * 200 + "|"]
    seen = []

    def classify(system, user, **kwargs):
        seen.append(user)
        return {"concept_per_row": True, "title_col": 2, "concept_type": "reference", "extraction_mode": "per_row"}

    with pytest.raises(ValueError, match="outside"):
        field_table._llm_classify_table(header, rows, SimpleNamespace(chat_json=classify), "test", 0)
    assert len(seen[0]) < 1000
    assert header == ["Name", "Value"] + [""] * 200


def test_classifier_uses_last_row_to_determine_width():
    header = ["Name", "", ""]
    rows = ["| Alpha | | |"] * 5 + ["| Beta | | Late value |"]
    user = field_table._classification_user(header, rows)
    assert "['Name', '', '']" in user
    assert "| --- | --- | --- |" in user
    assert "Late value" not in user


@pytest.mark.parametrize("cell", [r"A\|B", "`A|B`", "``A`|B``", "line<br>next", "0", "False"])
def test_projection_preserves_cell_syntax_and_interior_empty_column(cell):
    text = f"| Name | | Value | |\n|---|---|---|---|\n| {cell} | | end | |"
    prompt = prompt_for(text)
    assert f"| {cell} | | end |" in prompt
    assert "|---|---|---|---|" not in prompt


@pytest.mark.parametrize("text", [
    "| Name | Value | |\n|---|---|---|\n| Alpha | too | many | cells |",
    "| Name | Value | |\n|---|---|---|\n| Alpha | unfinished\ncontinuation | |",
    "```markdown\n| Name | Value | |\n|---|---|---|\n| Alpha | 0 | |\n```",
    "~~~\n| Name | Value | |\n|---|---|---|\n| Alpha | 0 | |\n~~~",
])
def test_ambiguous_or_fenced_tables_are_sent_unchanged(text):
    assert text in prompt_for(text)


def test_nonempty_header_preserves_empty_body_column():
    text = "| Name | Meaning | |\n|---|---|---|\n| Alpha | | |"
    prompt = prompt_for(text)
    assert "| Name | Meaning |" in prompt
    assert "| Alpha | |" in prompt


def test_table_projection_preserves_other_text_and_multiple_tables():
    text = "Before\n\n| A | |\n|---|---|\n| One | |\n\nBetween\n\n| B | |\n|---|---|\n| Two | |\n\nAfter"
    prompt = prompt_for(text)
    assert "Before\n\n| A |\n|---|\n| One |\n\nBetween" in prompt
    assert "| B |\n|---|\n| Two |\n\nAfter" in prompt


def test_generation_retains_a_column_populated_only_in_last_row():
    text = "| Name | | | |\n|---|---|---|---|\n| Alpha | | | |\n| Beta | | late | |"
    prompt = prompt_for(text)
    assert "| Name | | |\n|---|---|---|" in prompt
    assert "| Beta | | late |" in prompt


def test_generation_evidence_uses_canonical_chunk_offsets(monkeypatch):
    from app.services import okf_generator
    from app.services.source_evidence import chunk_digest

    source = "| Name | Value | |\n|---|---|---|\n| Alpha | Unique retained source evidence | |"
    seen = []

    def answer(system, user, **kwargs):
        seen.append(user)
        return [{"id": "alpha", "title": "Alpha", "type": "reference", "content": "Summary", "source_quotes": ["Unique retained source evidence"]}]

    monkeypatch.setattr(okf_generator, "extract_table_concepts", lambda text, **kw: ([], text))
    generator = OKFGenerator(llm=SimpleNamespace(chat_json=answer))
    concepts = generator.generate_chunk(source, "source.xlsx", 1, 1)
    assert "|---|---|---|" not in seen[0]
    span = concepts[0].source_spans[0]
    assert source[span.start:span.end] == "Unique retained source evidence"
    assert span.chunk_hash == chunk_digest(source)


def test_classifier_cache_tracks_width_changes_beyond_preview(monkeypatch):
    monkeypatch.setattr(field_table, "_get_prompt_hash", lambda: "fixed")
    rows = ["| Alpha | | |"] * 6
    original = field_table._build_cache_key(["Name", "", ""], rows)
    rows[-1] = "| Alpha | | Late value |"
    assert field_table._build_cache_key(["Name", "", ""], rows) != original


def test_projection_preserves_crlf_and_alignment():
    text = "  | Name | Value | |\r\n  |:---|---:|---|\r\n  | Alpha | False | |\r\n"
    assert "  | Name | Value |\r\n  |:---|---:|\r\n  | Alpha | False |\r\n" in prompt_for(text)


@pytest.mark.parametrize("mode", ["per_row", "whole"])
def test_programmatic_concepts_drop_empty_tail_but_keep_original_source_spans(mode, monkeypatch):
    from app.services.source_evidence import chunk_digest

    width = 300
    text = "\n".join([
        "| Name | Value " + "| " * (width - 2) + "|",
        "|" + "---|" * width,
        "| Alpha | kept " + "| " * (width - 2) + "|",
        "| Beta | retained " + "| " * (width - 2) + "|",
        "| Gamma | third " + "| " * (width - 2) + "|",
    ])
    monkeypatch.setattr(field_table, "_min_rows", lambda: 1)
    block = field_table.detect_tables(text)[0]
    cls = field_table.TableClassification(concept_per_row=True, title_col=0, extraction_mode=mode)
    concepts = field_table.build_row_concepts(block, cls, None, lines=text.split("\n"))
    assert all(len(c.content) < 1000 for c in concepts)
    assert all("Колонка 299" not in c.content for c in concepts)
    assert all(span.chunk_hash == chunk_digest(text) for c in concepts for span in c.source_spans)
    assert "retained" in "\n".join(c.content for c in concepts)
    assert len(block.header) == width




@pytest.mark.parametrize('unit', ['\ufdfa', '\u337f', 'ﬄ'])
def test_estimated_budget_covers_compatibility_normalization(unit):
    import unicodedata
    text = unit * 80
    estimate = module().BudgetEstimator().estimate(text)
    assert estimate.mode == 'estimated'
    assert estimate.tokens >= len(unicodedata.normalize('NFKC', text).encode('utf8'))


def test_normalized_long_cell_keeps_exact_original_ranges_and_bounded_identity():
    from dataclasses import replace
    import unicodedata
    m = module(); value = '\ufdfa' * 3000
    descriptor, _ = table(width=1)
    descriptor = replace(descriptor, sheet_name='\ufdfa' * 31,
                         columns=[TableColumn(1, 'A', ['\ufdfa' * 100])])
    rows = [TableRow(2, {1: TableCell('A2', value, value, 'string')})]
    title = descriptor.sheet_name[:64] + ' — 2 [ffffffff]\n'
    reserve = len(unicodedata.normalize('NFKC', title).encode('utf8'))
    limits = m.ProjectionLimits(embedding_tokens=6000-reserve, context_tokens=6000-reserve,
                                max_fragment_chars=1744)
    fragments = list(m.project_table(descriptor, rows, limits, generation_id='b' * 32))
    assert len(fragments) > 3
    ranges = sorted((ref.text_start, ref.text_end) for fragment in fragments for ref in fragment.coverage_refs)
    assert ranges[0][0] == 0 and ranges[-1][1] == len(value)
    assert all(left[1] == right[0] for left, right in zip(ranges, ranges[1:]))
    assert all(len(unicodedata.normalize('NFKC', title + fragment.content).encode('utf8')) <= 6000
               for fragment in fragments)
    assert ''.join(value[start:end] for start, end in ranges) == value

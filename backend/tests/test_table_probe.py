"""Acceptance checks use coordinates and typed values, never concept counts."""
from copy import deepcopy
import importlib.util
from pathlib import Path

import pytest


def _probe():
    path = Path(__file__).parents[1] / 'test_scripts' / 'probe_tables.py'
    assert path.is_file(), 'table acceptance probe is missing'
    spec = importlib.util.spec_from_file_location('table_probe', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_probe_fails_on_one_missing_cell():
    baseline = {'schema_version': 1, 'cells': [
        {'source_id': 'root', 'sheet': 0, 'address': 'A2', 'value': '001'},
        {'source_id': 'root', 'sheet': 0, 'address': 'LD2', 'value': 'right edge'},
    ], 'concept_count': 2}
    after = deepcopy(baseline)
    after['cells'].pop()
    after['concept_count'] = 100
    report = _probe().compare_reports(baseline, after)
    assert report['passed'] is False
    assert report['gates']['cell_coverage']['missing'] == 1


def test_probe_preserves_every_failed_gate():
    before = {'schema_version': 1, 'cells': [
        {'source_id': 'root', 'sheet': 0, 'address': 'A2', 'value': 0},
    ], 'headers': [{'source_id': 'root', 'sheet': 0, 'value': 'Code'}]}
    after = {'schema_version': 1, 'cells': [
        {'source_id': 'root', 'sheet': 0, 'address': 'A2', 'value': False},
    ], 'headers': []}
    report = _probe().compare_reports(before, after)
    assert report['passed'] is False
    assert report['gates']['cell_coverage']['changed'] == 1
    assert report['gates']['header_coverage']['passed'] is False


def test_probe_unknown_evidence_is_not_success():
    assert not _probe().compare_reports({'schema_version': 1}, {'schema_version': 1})['passed']


def test_probe_parse_cli_uses_original_coordinates_and_unknown_headers(tmp_path):
    import hashlib
    import json
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['Code', 'Value'])
    sheet.append(['001', False])
    sheet.append(['002', 0])
    source = tmp_path / 'input.xlsx'
    workbook.save(source)
    workbook.close()
    manifest = {'documents': [{'path': 'input.xlsx', 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
                              'source_id': 'root/0', 'expected_cells': [
                                  {'source_id': 'root/0', 'sheet': 0, 'address': 'A2', 'value': '001'},
                                  {'source_id': 'root/0', 'sheet': 0, 'address': 'B2', 'value': 'False'},
                                  {'source_id': 'root/0', 'sheet': 0, 'address': 'B3', 'value': '0'},
                              ], 'expected_headers': [{'source_id': 'root/0', 'sheet': 0,
                                                       'value': ['Code', 'Value'], 'status': 'legacy_unverified'}]}]}
    manifest_path = tmp_path / 'manifest.json'
    manifest_path.write_text(json.dumps(manifest), encoding='utf-8')
    output = tmp_path / 'report.json'
    probe = _probe()
    assert probe.main(['--manifest', str(manifest_path), '--output', str(output), '--mode', 'parse']) == 1
    report = json.loads(output.read_text(encoding='utf-8'))
    assert report['gates']['0']['gates']['cell_coverage']['passed']
    assert report['gates']['0']['gates']['header_correctness']['status'] == 'unknown'


def test_probe_cell_identity_includes_document():
    before = {'schema_version': 1, 'cells': [
        {'document_id': 'a', 'source_id': 'root/0', 'sheet': 0, 'address': 'A2', 'value': 'same'},
        {'document_id': 'b', 'source_id': 'root/0', 'sheet': 0, 'address': 'A2', 'value': 'same'},
    ], 'headers': [{'value': 'Code'}]}
    after = deepcopy(before)
    after['cells'].pop()
    report = _probe().compare_reports(before, after)
    assert not report['passed']
    assert report['gates']['cell_coverage']['missing'] == 1


def test_compare_preserves_failed_original_gates():
    before = {'schema_version': 1, 'cells': [
        {'source_id': 'root', 'sheet': 0, 'address': 'A2', 'value': 'same'},
    ], 'headers': [{'value': 'Code'}], 'passed': False,
        'gates': {'original_oracle': {'passed': False, 'missing': 1}}}
    report = _probe().compare_reports(before, deepcopy(before))
    assert not report['passed']
    assert report['no_regression'] is True
    assert report['gates']['source_acceptance']['baseline']['original_oracle']['missing'] == 1


def test_probe_preserves_sheet_identity_with_empty_and_reordered_sheets(tmp_path):
    import hashlib
    import zipfile
    from lxml import etree
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.title = 'Empty'
    for name in ('Second', 'Third'):
        sheet = workbook.create_sheet(name)
        sheet.append(['Code', 'Value'])
        sheet.append([name, name + '-value'])
    source = tmp_path / 'input.xlsx'
    workbook.save(source)
    workbook.close()
    with zipfile.ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(entries['xl/workbook.xml'])
    sheets = root.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheets')
    third = sheets[-1]
    sheets.remove(third)
    sheets.insert(1, third)
    entries['xl/workbook.xml'] = etree.tostring(root)
    with zipfile.ZipFile(source, 'w') as archive:
        for name, payload in entries.items():
            archive.writestr(name, payload)
    manifest = {'documents': [{'document_id': 'book', 'path': 'input.xlsx',
                              'sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'source_id': 'root',
                              'expected_cells': [
                                  {'source_id': 'root', 'sheet': 1, 'address': 'A2', 'value': 'Third'},
                                  {'source_id': 'root', 'sheet': 2, 'address': 'A2', 'value': 'Second'},
                              ], 'expected_headers': []}]}
    report = _probe().parse_manifest(manifest, tmp_path)
    assert report['gates']['0']['gates']['cell_coverage']['passed']
    assert {c['sheet']: c['value'] for c in report['cells'] if c['address'] == 'A2'} == {1: 'Third', 2: 'Second'}


def test_retrieval_probe_uses_the_same_query_aware_context_path_as_api(monkeypatch):
    from types import SimpleNamespace
    from app.config import Settings
    from app.services import context_builder, embedder, retrieval_hydration, vector_store

    query = 'Для кода 7601 укажи значение поля «Описание»'
    settings = Settings(_env_file=None, glossary_query_expansion_enabled=False)
    monkeypatch.setattr('app.config.get_settings', lambda: settings)
    hit = SimpleNamespace(payload={'doc_id': 'book', 'content': 'expected-value'})

    class Store:
        def __init__(self):
            self.client = SimpleNamespace(close=lambda: None)

        def search_composite(self, **kwargs):
            return [hit]

    monkeypatch.setattr(vector_store, 'VectorStore', Store)
    monkeypatch.setattr(embedder.Embedder, 'embed', lambda self, text: [])
    received = {}

    def hydrate(hits, **kwargs):
        received['hydrate_query'] = kwargs.get('query')
        return hits, {'book': {'filename': 'input.xlsx'}}

    def merge(hits, config, **kwargs):
        received['merge_query'] = kwargs.get('query')
        content = 'expected-value' if received['hydrate_query'] == query and kwargs.get('query') == query else 'other-value'
        return [{'doc_id': 'book', 'content': content}]

    monkeypatch.setattr(retrieval_hydration, 'load_visible_retrieval_hits', hydrate)
    monkeypatch.setattr(context_builder, 'merge_and_format', merge)
    monkeypatch.setattr(context_builder, 'drop_unmatched_blocks', lambda blocks, *args, **kwargs: blocks)
    monkeypatch.setattr(context_builder, 'drop_partial_title_matches', lambda blocks, *args, **kwargs: blocks)
    monkeypatch.setattr(_probe_module := _probe(), 'metadata', lambda: {})
    report = _probe_module.run_retrieval({'cases': [{'id': 'exact-field', 'category': 'table', 'query': query,
                                                  'evidence': [{'doc_id': 'book', 'value': 'expected-value'}]}]})
    assert report['passed'], report['cases'][0]['final_blocks']
    assert received == {'hydrate_query': query, 'merge_query': query}


def _retrieval_fixture(monkeypatch, payloads, hydrate):
    """Replace external retrieval I/O; probe snapshots/matching stay real."""
    from types import SimpleNamespace
    from app.config import Settings
    from app.services import context_builder, embedder, retrieval_hydration, vector_store
    settings = Settings(_env_file=None, glossary_query_expansion_enabled=False)
    monkeypatch.setattr('app.config.get_settings', lambda: settings)
    hits = [SimpleNamespace(payload=deepcopy(p)) for p in payloads]
    class Store:
        def __init__(self):
            self.client = SimpleNamespace(close=lambda: None)
        def search_composite(self, **kwargs):
            return hits
    monkeypatch.setattr(vector_store, 'VectorStore', Store)
    monkeypatch.setattr(embedder.Embedder, 'embed', lambda self, text: [])
    monkeypatch.setattr(retrieval_hydration, 'load_visible_retrieval_hits', hydrate)
    monkeypatch.setattr(context_builder, 'merge_and_format',
                        lambda current, *args, **kwargs: [deepcopy(h.payload) for h in current])
    monkeypatch.setattr(context_builder, 'drop_unmatched_blocks', lambda blocks, *args, **kwargs: blocks)
    monkeypatch.setattr(context_builder, 'drop_partial_title_matches', lambda blocks, *args, **kwargs: blocks)
    probe = _probe()
    monkeypatch.setattr(probe, 'metadata', lambda: {})
    return probe


def test_probe_raw_snapshot_precedes_nested_hydration_mutation(monkeypatch):
    payload = {'doc_id': 'book', 'content': 'index-value', 'table_refs': [{'columns': [1]}]}
    def hydrate(hits, **kwargs):
        hits[0].payload['content'] = 'canonical-value'
        hits[0].payload['table_refs'][0]['columns'].append(2)
        return hits, {}
    probe = _retrieval_fixture(monkeypatch, [payload], hydrate)
    result = probe.run_retrieval({'cases': [{'id': 'one', 'category': 'table', 'query': 'value',
                                            'evidence': [{'doc_id': 'book', 'value': 'canonical-value'}]}]})
    case = result['cases'][0]
    assert case['raw_candidates'] == [payload]
    assert case['visible_candidates'][0]['content'] == 'canonical-value'


def test_probe_removed_raw_hit_does_not_count_as_visible_recall(monkeypatch):
    probe = _retrieval_fixture(monkeypatch, [{'doc_id': 'book', 'content': '312'}],
                               lambda hits, **kwargs: ([], {}))
    result = probe.run_retrieval({'cases': [{'id': 'stale', 'category': 'table', 'query': '312',
                                            'evidence': [{'doc_id': 'book', 'value': '312'}]}]})
    case = result['cases'][0]
    assert case['raw_candidate_recall'] == [True]
    assert case['candidate_recall'] == [False]
    assert case['context_value_and_source'] == [False]


def test_probe_empty_value_and_numeric_substring_cannot_pass(monkeypatch):
    probe = _retrieval_fixture(monkeypatch, [{'doc_id': 'book', 'content': 'Code: 3120; Header'}],
                               lambda hits, **kwargs: (hits, {}))
    report = probe.run_retrieval({'cases': [
        {'id': 'empty', 'category': 'table', 'query': 'Code', 'evidence': [{'doc_id': 'book', 'value': ''}]},
        {'id': 'substring', 'category': 'table', 'query': 'Code',
         'evidence': [{'doc_id': 'book', 'value': '312', 'header': 'Header'}]},
    ]})
    assert [case['passed'] for case in report['cases']] == [False, False]


def test_physical_evidence_survives_rechunking_but_rejects_other_row_and_field():
    probe = _probe()
    evidence = {'source_sha256': 'a' * 64, 'source_id': 'root', 'sheet_index': 0,
                'address': 'B7', 'state': 'value', 'value': '312', 'header': 'Amount'}
    binding = {'doc_id': 'book', 'source_id': 'root', 'source_sha256': 'a' * 64,
               'sheet_index': 0, 'table_id': 'table', 'generation_id': 'new'}
    block = {'doc_id': 'book', 'source_id': 'root', 'generation_id': 'new',
             '_canonical_verified': True, 'chunk_index': 999,
             'table_refs': [{'source_id': 'root', 'generation_id': 'new', 'table_id': 'table',
                             'row_number': 7, 'columns': [2]}],
             'content': 'Sheet: Example; row: 7; key: K\nB7 | Amount: 312'}
    assert probe.matches_evidence(block, evidence, bindings=[binding], require_value=True)
    wrong = deepcopy(block)
    wrong['table_refs'][0]['row_number'] = 8
    assert not probe.matches_evidence(wrong, evidence, bindings=[binding], require_value=True)
    wrong = deepcopy(block)
    wrong['content'] = 'B7 | Other: 312\nC7 | Amount: 999'
    assert not probe.matches_evidence(wrong, evidence, bindings=[binding], require_value=True)


def test_oracle_audit_rejects_wrong_cell_and_unqualified_empty(tmp_path):
    import hashlib
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(['Code', 'Amount'])
    workbook.active.append(['K', 312])
    source = tmp_path / 'original.xlsx'
    workbook.save(source)
    workbook.close()
    sha = hashlib.sha256(source.read_bytes()).hexdigest()
    evidence = {'source_sha256': sha, 'source_id': 'root', 'sheet_index': 0,
                'address': 'B2', 'state': 'value', 'value': '312',
                'row_key': [{'address': 'A2', 'value': 'K'}],
                'header_anchors': [{'address': 'B1', 'value': 'Amount'}]}
    cases = []
    for name, changes in [('correct', {}), ('wrong', {'address': 'A2'}), ('empty', {'value': ''})]:
        cases.append({'id': name, 'category': 'table', 'intent': 'lookup', 'query': 'For K, Amount?',
                      'evidence': [{**evidence, **changes}]})
    manifest = {'schema_version': 2, 'documents': [{'path': 'original.xlsx', 'sha256': sha, 'source_id': 'root'}],
                'cases': cases}
    report = _probe().audit_manifest(manifest, tmp_path)
    assert [case['status'] for case in report['cases']] == ['source-valid', 'invalid-oracle', 'invalid-oracle']
    assert report['passed'] is False


def test_probe_numeric_literal_does_not_match_signed_or_decimal_value(monkeypatch):
    probe = _retrieval_fixture(monkeypatch, [{'doc_id': 'book', 'content': '-312 312.5 312,5'}],
                               lambda hits, **kwargs: (hits, {}))
    result = probe.run_retrieval({'cases': [{'id': 'numeric', 'category': 'table', 'query': 'Amount',
                                           'evidence': [{'doc_id': 'book', 'value': '312'}]}]})
    assert result['cases'][0]['passed'] is False


def test_physical_evidence_rejects_unverified_source_and_mixed_generation():
    probe = _probe()
    evidence = {'source_sha256': 'a' * 64, 'source_id': 'root', 'sheet_index': 0,
                'address': 'B7', 'state': 'value', 'value': '312', 'header': 'Amount'}
    binding = {'doc_id': 'book', 'source_id': 'root', 'source_sha256': 'a' * 64,
               'sheet_index': 0, 'table_id': 'table', 'generation_id': 'new'}
    block = {'doc_id': 'book', 'source_id': 'root', 'generation_id': 'new', '_canonical_verified': True,
             'table_refs': [{'source_id': 'root', 'generation_id': 'old', 'table_id': 'table',
                             'row_number': 7, 'columns': [2]}], 'content': 'B7 | Amount: 312'}
    assert not probe.matches_evidence(block, evidence, bindings=[binding], require_value=True)
    block['table_refs'][0]['generation_id'] = 'new'
    block['_canonical_verified'] = False
    assert not probe.matches_evidence(block, evidence, bindings=[binding], require_value=True)
    block['_canonical_verified'] = True
    binding['source_sha256'] = 'b' * 64
    assert not probe.matches_evidence(block, evidence, bindings=[binding], require_value=True)


def test_physical_evidence_matches_multiline_cell_without_neighbor_value():
    probe = _probe()
    evidence = {'source_sha256': 'a' * 64, 'source_id': 'root', 'sheet_index': 0,
                'address': 'B7', 'state': 'value', 'value': 'first\nsecond', 'header': 'Header\ncontinued'}
    binding = {'doc_id': 'book', 'source_id': 'root', 'source_sha256': 'a' * 64,
               'sheet_index': 0, 'table_id': 'table', 'generation_id': 'new'}
    block = {'doc_id': 'book', 'source_id': 'root', 'generation_id': 'new', '_canonical_verified': True,
             'table_refs': [{'source_id': 'root', 'generation_id': 'new', 'table_id': 'table',
                             'row_number': 7, 'columns': [2, 3]}],
             'content': 'Sheet: Example; row: 7; key: K\nB7 | Header\ncontinued: first\nsecond\nC7 | Other: 312'}
    assert probe.matches_evidence(block, evidence, bindings=[binding], require_value=True)
    evidence['value'] = '312'
    assert not probe.matches_evidence(block, evidence, bindings=[binding], require_value=True)


def test_physical_evidence_rejects_block_from_another_source():
    evidence = {'source_sha256': 'a' * 64, 'source_id': 'root', 'sheet_index': 0,
                'address': 'B7', 'state': 'value', 'value': '312'}
    binding = {'doc_id': 'book', 'source_id': 'root', 'source_sha256': 'a' * 64,
               'sheet_index': 0, 'table_id': 'table', 'generation_id': 'new'}
    block = {'doc_id': 'book', 'source_id': 'root/other', 'generation_id': 'new', '_canonical_verified': True,
             'table_refs': [{'source_id': 'root', 'generation_id': 'new', 'table_id': 'table',
                             'row_number': 7, 'columns': [2]}], 'content': 'B7 | Amount: 312'}
    assert not _probe().matches_evidence(block, evidence, bindings=[binding], require_value=True)


def test_v2_retrieval_cli_rejects_invalid_oracle_before_search(tmp_path, monkeypatch):
    import hashlib
    import json
    from openpyxl import Workbook
    source = tmp_path / 'input.xlsx'
    workbook = Workbook()
    workbook.active.append(['Code', 'Amount'])
    workbook.active.append(['K', 312])
    workbook.save(source)
    workbook.close()
    sha = hashlib.sha256(source.read_bytes()).hexdigest()
    manifest = {'schema_version': 2, 'documents': [{'path': 'input.xlsx', 'sha256': sha, 'source_id': 'root'}],
                'cases': [{'id': 'wrong-oracle', 'category': 'table', 'intent': 'lookup', 'query': 'For K, Amount?',
                           'evidence': [{'source_sha256': sha, 'source_id': 'root', 'sheet_index': 0,
                                         'address': 'B2', 'value': '999', 'state': 'value',
                                         'row_key': [{'address': 'A2', 'value': 'K'}],
                                         'header_anchors': [{'address': 'B1', 'value': 'Amount'}]}]}]}
    file = tmp_path / 'manifest.json'
    file.write_text(json.dumps(manifest), encoding='utf-8')
    output = tmp_path / 'output.json'
    probe = _probe()
    monkeypatch.setattr(probe, 'run_retrieval', lambda manifest: {'passed': True})
    assert probe.main(['--manifest', str(file), '--mode', 'retrieval', '--output', str(output)]) == 1
    result = json.loads(output.read_text())
    assert result['status'] == 'invalid'
    assert result['oracle_audit']['cases'][0]['status'] == 'invalid-oracle'


@pytest.mark.parametrize('state,marker', [('blank', '[blank]'), ('empty_string', '[empty_string]')])
def test_blank_evidence_requires_explicit_same_cell_canonical_state(state, marker):
    probe = _probe()
    evidence = {'source_sha256': 'a' * 64, 'source_id': 'root', 'sheet_index': 0,
                'address': 'B7', 'state': 'blank', 'value': '', 'header': 'Amount'}
    binding = {'doc_id': 'book', 'source_id': 'root', 'source_sha256': 'a' * 64,
               'sheet_index': 0, 'table_id': 'table', 'generation_id': 'new'}
    ref = {'source_id': 'root', 'generation_id': 'new', 'table_id': 'table', 'row_number': 7, 'columns': [2]}
    block = {'doc_id': 'book', 'source_id': 'root', 'generation_id': 'new', '_canonical_verified': True,
             'table_refs': [ref], 'content': f'B7 | Amount: {marker}',
             '_table_cell_states': [{**ref, 'address': 'B7', 'column': 2, 'state': state}]}
    assert probe.matches_evidence(block, evidence, bindings=[binding], require_value=True)
    del block['_table_cell_states']
    assert not probe.matches_evidence(block, evidence, bindings=[binding], require_value=True)


@pytest.mark.parametrize('changes', [
    {'state': 'empty_formula_result'}, {'state': 'formula_missing'}, {'state': 'value'},
    {'source_id': 'root/1'}, {'generation_id': 'old'}, {'table_id': 'other'},
    {'row_number': 8}, {'column': 3}, {'address': 'B8'}, {'columns': [3]},
])
def test_blank_marker_with_wrong_or_formula_state_is_not_empty_cell_evidence(changes):
    evidence = {'source_sha256': 'a' * 64, 'source_id': 'root', 'sheet_index': 0,
                'address': 'B7', 'state': 'blank', 'value': '', 'header': 'Amount'}
    binding = {'doc_id': 'book', 'source_id': 'root', 'source_sha256': 'a' * 64,
               'sheet_index': 0, 'table_id': 'table', 'generation_id': 'new'}
    ref = {'source_id': 'root', 'generation_id': 'new', 'table_id': 'table', 'row_number': 7, 'columns': [2]}
    proof = {**ref, 'address': 'B7', 'column': 2, 'state': 'blank', **changes}
    block = {'doc_id': 'book', 'source_id': 'root', 'generation_id': 'new', '_canonical_verified': True,
             'table_refs': [ref], 'content': 'B7 | Amount: [blank]', '_table_cell_states': [proof]}
    assert not _probe().matches_evidence(block, evidence, bindings=[binding], require_value=True)


@pytest.mark.parametrize('state,value,marker', [('value', '312', '312'), ('blank', None, '[blank]'), ('empty_string', None, '[empty_string]')])
def test_physical_evidence_accepts_only_same_cell_state_annotation(state, value, marker):
    evidence = {'source_sha256': 'a' * 64, 'source_id': 'root', 'sheet_index': 0,
                'address': 'B7', 'state': 'blank' if value is None else 'value', 'value': value, 'header': 'Amount'}
    binding = {'doc_id': 'book', 'source_id': 'root', 'source_sha256': 'a' * 64,
               'sheet_index': 0, 'table_id': 'table', 'generation_id': 'new'}
    ref = {'source_id': 'root', 'generation_id': 'new', 'table_id': 'table', 'row_number': 7, 'columns': [2]}
    proof = {**ref, 'address': 'B7', 'column': 2, 'state': state}
    block = {'doc_id': 'book', 'source_id': 'root', 'generation_id': 'new', '_canonical_verified': True,
             'table_refs': [ref], '_table_cell_states': [proof],
             'content': f'B7 | Amount: {marker} (cell state: {state})\nA7 | Code: 0001 (cell state: value)'}
    assert _probe().matches_evidence(block, evidence, bindings=[binding], require_value=True)
    for fault in ('address', 'source_id', 'generation_id', 'table_id', 'state'):
        wrong = deepcopy(block)
        wrong['_table_cell_states'][0][fault] = 'different'
        assert not _probe().matches_evidence(wrong, evidence, bindings=[binding], require_value=True)
    wrong = deepcopy(block); wrong['_table_cell_states'] = []
    assert not _probe().matches_evidence(wrong, evidence, bindings=[binding], require_value=True)
    wrong = deepcopy(block); wrong['content'] = f'B7 | Amount: wrong (cell state: {state})\nA7 | Code: {marker}'
    assert not _probe().matches_evidence(wrong, evidence, bindings=[binding], require_value=True)

"""A trusted server flag selects new representations without reinterpreting resumes."""
import pytest
from app.config import Settings
from app.db.models import Document, DocumentGeneration, DocumentGenerationState
from app.db.session import session_scope
from app.services.generation_store import GenerationConflict, begin_generation, prepare_generation_attempt
from app.services.errors import VectorStoreError
from app.services.staging import StagingStore
from tests.test_generation_pipeline import DOC_ID, pipeline_env as pipeline_env


def test_default_writer_stays_disabled_and_incompatible_index_is_rejected():
    assert Settings(_env_file=None).chunk_index_projection_enabled is False
    with pytest.raises(ValueError, match="chunk"):
        Settings(_env_file=None, chunk_index_projection_enabled=True, search_index_chunks_enabled=False)


@pytest.mark.parametrize("invalid", [True, False, 0, 999, "1"])
def test_unsupported_requested_version_does_not_replace_candidate(invalid):
    with session_scope() as s:
        s.add(Document(id="policy", filename="source.docx")); s.flush()
        old = begin_generation(s, "policy")
        with pytest.raises(GenerationConflict):
            prepare_generation_attempt(s, "policy", resume=False, chunk_projection_version=invalid)
        assert s.get(DocumentGenerationState, "policy").candidate_generation_id == old.id
        assert old.phase == "preparing" and old.chunk_projection_version is None


@pytest.mark.parametrize("existing, requested", [(None, 1), (1, None)])
def test_resume_keeps_existing_marker_across_server_flag_change(existing, requested):
    with session_scope() as s:
        s.add(Document(id="resume-policy", filename="source.docx")); s.flush()
        old = begin_generation(s, "resume-policy", chunk_projection_version=existing)
        selected = prepare_generation_attempt(s, "resume-policy", resume=True, chunk_projection_version=requested)
        assert selected.id == old.id and selected.chunk_projection_version == existing


def test_new_candidate_can_select_format_without_private_hook():
    with session_scope() as s:
        s.add(Document(id="new-policy", filename="source.docx")); s.flush()
        generation = prepare_generation_attempt(s, "new-policy", resume=False, chunk_projection_version=1)
        assert generation.chunk_projection_version == 1


@pytest.mark.parametrize("structural", [False, True])
def test_public_pipeline_flag_selects_format_without_injecting_marker(pipeline_env, structural):
    pipeline, source, write = pipeline_env
    pipeline.settings.chunk_index_projection_enabled = True
    pipeline.settings.table_structure_enabled = structural
    write("Public server selected projection with literal 0 and False.")
    pipeline._process(DOC_ID, source, "decision.eml", [], resume=False)
    assert pipeline.registry.get(DOC_ID)["status"] == "done"
    with session_scope() as s:
        state = s.get(DocumentGenerationState, DOC_ID)
        assert state.candidate_generation_id is None
        assert s.get(DocumentGeneration, state.active_generation_id).chunk_projection_version == 1


def test_public_format_resume_survives_disabling_new_candidate_flag(pipeline_env, monkeypatch):
    pipeline, source, write = pipeline_env
    pipeline.settings.chunk_index_projection_enabled = True
    write("Resume retains the originally bound source representation.")
    original = pipeline.vector_store.index_chunk_projection_fragments
    def fail(*a, **kw):
        raise VectorStoreError("synthetic public format index interruption")
    monkeypatch.setattr(pipeline.vector_store, "index_chunk_projection_fragments", fail)
    pipeline._process(DOC_ID, source, "decision.eml", [], resume=False)
    assert pipeline.registry.get(DOC_ID)["status"] == "paused"
    candidate = StagingStore(DOC_ID).load()["generation_id"]
    with session_scope() as s:
        assert s.get(DocumentGeneration, candidate).chunk_projection_version == 1
    pipeline.settings.chunk_index_projection_enabled = False
    monkeypatch.setattr(pipeline.vector_store, "index_chunk_projection_fragments", original)
    pipeline._process(DOC_ID, source, "decision.eml", [], resume=True)
    assert pipeline.registry.get(DOC_ID)["status"] == "done"
    with session_scope() as s:
        state = s.get(DocumentGenerationState, DOC_ID)
        assert state.active_generation_id == candidate and state.candidate_generation_id is None
        assert s.get(DocumentGeneration, candidate).chunk_projection_version == 1



def test_public_table_ref_only_route_has_no_generic_parents(pipeline_env, tmp_path):
    from openpyxl import Workbook
    from app.db.models import DocumentChunk, DocumentChunkProjection
    pipeline, _, _ = pipeline_env
    pipeline.settings.chunk_index_projection_enabled = True
    pipeline.settings.table_structure_enabled = True
    pipeline.settings.okf_table_llm_classify = False
    source = tmp_path / "public-only-table.xlsx"
    wb = Workbook(); sheet = wb.active
    sheet.append(["Code", "Value"]); sheet.append(["A", 0]); sheet.append(["B", False])
    wb.save(source)
    doc = "public-marker-table"
    pipeline.registry.create(doc, source.name, "application/octet-stream", source.stat().st_size)
    pipeline._process(doc, source, source.name, [], resume=False)
    assert pipeline.registry.get(doc)["status"] == "done"
    with session_scope() as s:
        state = s.get(DocumentGenerationState, doc)
        assert s.get(DocumentGeneration, state.active_generation_id).chunk_projection_version == 1
        rows = s.query(DocumentChunk).filter_by(doc_id=doc).all()
        assert rows and all(row.table_refs for row in rows)
        assert not s.query(DocumentChunkProjection).filter_by(doc_id=doc).all()


def test_public_enable_does_not_reinterpret_legacy_pipeline_resume(pipeline_env, monkeypatch):
    pipeline, source, write = pipeline_env
    write("A legacy paused candidate retains its original representation.")
    original = pipeline.vector_store.index_chunks
    def fail(*a, **kw):
        raise VectorStoreError("synthetic legacy index interruption")
    monkeypatch.setattr(pipeline.vector_store, "index_chunks", fail)
    pipeline._process(DOC_ID, source, "decision.eml", [], resume=False)
    assert pipeline.registry.get(DOC_ID)["status"] == "paused"
    candidate = StagingStore(DOC_ID).load()["generation_id"]
    with session_scope() as s:
        assert s.get(DocumentGeneration, candidate).chunk_projection_version is None
    pipeline.settings.chunk_index_projection_enabled = True
    monkeypatch.setattr(pipeline.vector_store, "index_chunks", original)
    pipeline._process(DOC_ID, source, "decision.eml", [], resume=True)
    assert pipeline.registry.get(DOC_ID)["status"] == "done"
    with session_scope() as s:
        assert s.get(DocumentGenerationState, DOC_ID).active_generation_id == candidate
        assert s.get(DocumentGeneration, candidate).chunk_projection_version is None

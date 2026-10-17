"""One synthetic mail combines generic parts, short chunks and structural tables."""
from copy import deepcopy
from email.message import EmailMessage
from io import BytesIO
import json
import pytest
from sqlalchemy import select
from app.db.models import DocumentChunk,DocumentGeneration,DocumentGenerationState,DocumentChunkFragment
from app.db.session import session_scope
from app.services.vector_store import VectorStore,chunk_point_id
from app.services.fusion import Hit
from app.services.retrieval_hydration import load_visible_retrieval_hits
from app.services.context_builder import merge_and_format
from app.services.chat_source_selection import snapshot_blocks,restore_blocks
from app.services.canonical_repair import repair_published_document
from app.services.export_okf import export_okf_bundle
from app.api.errors import ApiError
from tests.test_generation_pipeline import pipeline_env as pipeline_env
from tests.test_chunk_fragment_pipeline import explicit_format


DOC='g6-mixed-format1'


def points(store,gen):
    from qdrant_client import models as qm
    records,offset=store.client.scroll(store.collection,limit=1000,with_vectors=False,
        scroll_filter=qm.Filter(must=[qm.FieldCondition(key='generation_id',match=qm.MatchValue(value=gen))]))
    assert offset is None
    return records


def test_mixed_fresh_search_snapshot_restart_repair_export(pipeline_env,tmp_path,monkeypatch):
    from openpyxl import Workbook
    from docx import Document as WordDocument
    pipeline,_,_=pipeline_env
    explicit_format(monkeypatch)
    pipeline.settings.table_structure_enabled=True
    pipeline.settings.okf_table_llm_classify=False
    pipeline.settings.embedding_batch_size=2
    book=Workbook();sheet=book.active
    for row in [['Code','Value'],['A',0],['B',False]]:sheet.append(row)
    xlsx=BytesIO();book.save(xlsx)
    word=WordDocument();word.add_paragraph('Short standalone rule 0 False.')
    docx=BytesIO();word.save(docx)
    message=EmailMessage();message['Subject']='Mixed lifecycle'
    message.set_content('physical rule 0 False.\n'*600+'Final mail value.')
    message.add_attachment(xlsx.getvalue(),maintype='application',subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename='table.xlsx')
    message.add_attachment(docx.getvalue(),maintype='application',subtype='vnd.openxmlformats-officedocument.wordprocessingml.document',filename='short.docx')
    source=pipeline.settings.uploads_dir/(DOC+'.eml');source.write_bytes(message.as_bytes())
    pipeline.registry.create(DOC,source.name,'message/rfc822',source.stat().st_size)
    pipeline._process(DOC,source,source.name,[],resume=False)
    assert pipeline.registry.get(DOC)['status']=='done'
    with session_scope() as s:
        gen=s.get(DocumentGenerationState,DOC).active_generation_id
        assert s.get(DocumentGeneration,gen).chunk_projection_version==1
        rows=list(s.scalars(select(DocumentChunk).where(DocumentChunk.doc_id==DOC)))
        table_indices={row.chunk_index for row in rows if row.table_refs}
        part_ids=set(s.scalars(select(DocumentChunkFragment.point_id).where(DocumentChunkFragment.generation_id==gen)))
        short_indices={row.chunk_index for row in rows if not row.table_refs and len(row.content.encode('utf-8'))<1000}
    assert table_indices and part_ids and short_indices
    from app.services.sparse import to_sparse_vector
    candidates=pipeline.vector_store.search_composite(dense_vec=None,sparse_vec=to_sparse_vector('physical'),
        tags=None,branches={'bm25'},top_k=40,per_branch_top_k=40,doc_ids=[DOC],mail_mode='only')
    found,_=load_visible_retrieval_hits(candidates,mail_mode='only')
    assert found and any(h.payload.get('chunk_fragment') for h in found)
    assert all(h.payload['doc_id']==DOC for h in found)
    original=points(pipeline.vector_store,gen)
    hits=[Hit(str(p.id),1.,deepcopy(p.payload),rank=i) for i,p in enumerate(original)]
    visible,_=load_visible_retrieval_hits(hits,mail_mode='only')
    blocks=merge_and_format(visible,mail_mode='only',limit_total_chars=False)
    assert any(b.get('table_refs') for b in blocks)
    assert any((b.get('_evidence') or {}).get('chunk_fragment') for b in blocks)
    snapshots=snapshot_blocks(blocks)
    restored=restore_blocks(snapshots,mail_mode='only')
    assert [b['content'] for b in restored]==[b['content'] for b in blocks]
    excluded,_=load_visible_retrieval_hits([Hit(str(p.id),1.,deepcopy(p.payload)) for p in original],mail_mode='exclude')
    assert not excluded
    restarted=VectorStore.__new__(VectorStore)
    restarted.settings=pipeline.settings;restarted.client=pipeline.vector_store.client
    assert restarted.backfill_chunks(pipeline.embedder,batch_size=2)==0
    missing=next(iter(part_ids));restarted.client.delete(restarted.collection,[missing])
    assert restarted.backfill_chunks(pipeline.embedder,batch_size=2)==1
    assert {str(p.id) for p in points(restarted,gen)}=={str(p.id) for p in original}
    table_id=chunk_point_id(DOC,next(iter(table_indices)),generation_id=gen)
    restarted.client.delete(restarted.collection,[table_id])
    assert restarted.backfill_chunks(pipeline.embedder,batch_size=2)==1
    def edit(snapshot):
        row=next(row for row in snapshot['chunks'] if not row.get('table_refs') and len(row['content'])>3000)
        row['content']+='\nRepaired final value.'
    result=repair_published_document(DOC,pipeline.settings,pipeline.okf_generator,pipeline.embedder,restarted,edit)
    assert result['generation_id']!=gen
    with pytest.raises(ApiError) as caught:restore_blocks(snapshots,mail_mode='only')
    assert caught.value.status_code==409
    dest=tmp_path/'mixed-export'
    files=export_okf_bundle(DOC,dest)
    assert 'tables/manifest.json' in files and 'chunks/projection.json' in files
    metadata=json.loads((dest/'chunks/projection.json').read_text(encoding='utf-8'))
    assert metadata['chunk_projection_version']==1 and {r['route'] for r in metadata['chunks']}=={'generic','table'}

"""Preview excluded originals without admitting generation or changing selection."""
import io
import importlib
import zipfile

import pytest
from openpyxl import Workbook

from tests import test_processing_scope_api as scope_fixtures

upload_client = scope_fixtures.upload_client
base_upload_client = scope_fixtures.base_upload_client


def test_labels_are_read_only_and_available_before_preview(upload_client, monkeypatch):
    from email.message import EmailMessage
    import struct
    import app.api.document_processing_scope as api
    client, settings, started = upload_client
    settings.attachment_review_enabled = True
    settings.dedup_enabled = False
    monkeypatch.setattr(api, 'get_settings', lambda: settings)
    import app.services.processing_scope_labels as labels_service
    calls = []
    real_worker = labels_service.parse_document_supervised
    def count_worker(*args, **kwargs):
        calls.append(kwargs.get('_inspection_options'))
        return real_worker(*args, **kwargs)
    monkeypatch.setattr(labels_service, 'parse_document_supervised', count_worker)
    text = 'Real source name.pdf'.encode('utf-16-le')
    header = bytearray(88)
    struct.pack_into('<II', header, 0, 1, 88)
    header[40:44] = b' EMF'
    record = bytearray(76 + len(text) + (-len(text) % 4))
    struct.pack_into('<II', record, 0, 84, len(record))
    struct.pack_into('<II', record, 44, len(text) // 2, 76)
    record[76:76 + len(text)] = text
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w') as archive:
        archive.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"><w:object><v:shape><v:imagedata r:id="icon"/></v:shape><o:OLEObject r:id="file"/></w:object></w:document>')
        archive.writestr('word/_rels/document.xml.rels', '<Relationships><Relationship Id="icon" Target="media/icon.emf"/><Relationship Id="file" Target="embeddings/oleObject1.bin"/></Relationships>')
        archive.writestr('word/media/icon.emf', header + record)
        from pypdf import PdfWriter
        pdf = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(pdf)
        archive.writestr('word/embeddings/oleObject1.bin', pdf.getvalue())
    mail = EmailMessage()
    mail.set_content('Root')
    mail.add_attachment(output.getvalue(), maintype='application', subtype='octet-stream', filename='nested.docx')
    book = Workbook()
    for _ in range(600):
        book.active.append(['HEAVY'])
    heavy = io.BytesIO()
    book.save(heavy)
    mail.add_attachment(heavy.getvalue(), maintype='application', subtype='octet-stream', filename='large.xlsx')
    response = client.post('/api/documents', files={'file': ('root.eml', mail.as_bytes(), 'message/rfc822')})
    assert response.status_code == 202, response.text
    route = f"/api/documents/{response.json()['id']}/processing-scope"
    before = client.get(route).json()
    review = before['pending']
    assert any(n['display_name'] == 'oleObject1.bin' for n in review['inventory']['nodes']), review['inventory']
    node = next(n for n in review['inventory']['nodes'] if n['display_name'] == 'oleObject1.bin')
    # Simulate an already saved inventory, created before icon-caption support.
    from app.db.session import session_scope
    from app.db.models import DocumentProcessingReview
    import copy
    with session_scope() as session:
        row = session.get(DocumentProcessingReview, response.json()['id'])
        inventory = copy.deepcopy(row.inventory_json)
        for item in inventory['nodes']:
            item.pop('display_label', None)
        row.inventory_json = inventory
    before = client.get(route).json()
    result = client.get(route + '/labels', params={'scope_type': 'pending', 'scope_id': review['id']})
    assert result.status_code == 200, result.text
    assert result.headers['cache-control'] == 'no-store'
    assert result.json()['labels'][node['source_key']]['display_label'] == 'Real source name.pdf'
    worker_calls = len(calls)
    assert worker_calls == 1
    assert client.get(route + '/labels', params={'scope_type': 'pending', 'scope_id': review['id']}).json() == result.json()
    assert len(calls) == worker_calls, 'Repeated label reads must reuse metadata after verifying the original hash'
    assert client.get(route).json() == before
    assert not started
    assert client.get(route + '/labels', params={'scope_type': 'pending', 'scope_id': 'stale'}).status_code == 409
    from app.db.models import DocumentGeneration, DocumentGenerationState
    generation_id = 'a' * 32
    with session_scope() as session:
        session.add(DocumentGeneration(id=generation_id, doc_id=response.json()['id'], phase='active',
                                      processing_scope_json={'schema_version': 1, 'review_id': review['id'], 'inventory': inventory}))
        state = session.get(DocumentGenerationState, response.json()['id'])
        if state is None:
            state = DocumentGenerationState(doc_id=response.json()['id'])
            session.add(state)
        state.active_generation_id = generation_id
    active = client.get(route + '/labels', params={'scope_type': 'active', 'scope_id': generation_id})
    assert active.status_code == 200, active.text
    assert active.json()['labels'] == result.json()['labels']
    assert len(calls) == worker_calls, 'The same verified bytes can supply active and pending display labels'
    assert client.get(route + '/labels', params={'scope_type': 'active', 'scope_id': 'stale'}).status_code == 409
    with session_scope() as session:
        row = session.get(DocumentProcessingReview, response.json()['id'])
        changed = copy.deepcopy(row.inventory_json)
        next(n for n in changed['nodes'] if n['kind'] == 'docx')['content_fingerprint'] = '0' * 64
        row.inventory_json = changed
    assert client.get(route + '/labels', params={'scope_type': 'pending', 'scope_id': review['id']}).status_code == 409


@pytest.mark.parametrize('kind',['docx','pdf','eml','msg','zip','7z','rar'])
def test_supported_formats_preview_real_worker_without_admission(upload_client,monkeypatch,kind):
    from email.message import EmailMessage
    from pathlib import Path
    import app.api.document_processing_scope as api
    client,settings,started=upload_client
    settings.attachment_review_enabled=True;settings.dedup_enabled=False
    monkeypatch.setattr(api,'get_settings',lambda:settings)
    output=io.BytesIO()
    if kind=='docx':
        with zipfile.ZipFile(output,'w') as z:
            z.writestr('word/document.xml','<document><p><t>DOCUMENT_CANARY</t></p></document>')
            z.writestr('docProps/core.xml','<core xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Real document title</dc:title></core>')
    elif kind=='pdf':
        from pypdf import PdfWriter
        writer=PdfWriter();writer.add_blank_page(width=100,height=100);writer.add_metadata({'/Title':'Real PDF title'});writer.write(output)
    elif kind=='eml':
        inner=EmailMessage();inner['Subject']='Real mail title';inner.set_content('MAIL_CANARY');output.write(inner.as_bytes())
    elif kind=='msg':
        output.write((Path(__file__).parents[2]/'doc-parser/tests/fixtures/mail/synthetic-unicode-attachment.msg').read_bytes())
    elif kind=='zip':
        with zipfile.ZipFile(output,'w') as z:z.writestr('file.txt','ARCHIVE_PAYLOAD')
    elif kind=='7z':
        import py7zr
        with py7zr.SevenZipFile(output,'w') as z:z.writestr(b'ARCHIVE_PAYLOAD','file.txt')
    else:
        output.write((Path(__file__).parents[2]/'doc-parser/tests/fixtures/archives/rar5-solid.rar').read_bytes())
    root=EmailMessage();root.set_content('Root body')
    book=Workbook()
    for _ in range(600):book.active.append(['HEAVY'])
    heavy=io.BytesIO();book.save(heavy)
    root.add_attachment(heavy.getvalue(),maintype='application',subtype='octet-stream',filename='heavy.xlsx')
    root.add_attachment(output.getvalue(),maintype='application',subtype='octet-stream',filename='sample.'+kind)
    uploaded=client.post('/api/documents',files={'file':('root.eml',root.as_bytes())})
    assert uploaded.status_code==202,uploaded.text
    route=f"/api/documents/{uploaded.json()['id']}/processing-scope"
    before=client.get(route).json();review=before['pending']
    node=next(n for n in review['inventory']['nodes'] if n['display_name']=='sample.'+kind)
    response=client.get(route+'/attachments/'+node['source_key']+'/preview',params={'review_id':review['id']})
    assert response.status_code==200,response.text
    assert response.headers['cache-control']=='no-store'
    assert response.json()['kind']==('archive' if kind in {'zip','7z','rar'} else 'text')
    if kind in {'docx','pdf','eml'}:assert response.json()['title']==node['title'] and node['title']
    assert client.get(route).json()==before
    assert not started


def preview_function():
    return importlib.import_module("docparser.xlsx_preview").preview_xlsx


def test_preview_bounds_rows_columns_and_preserves_sheet_choice(tmp_path):
    book = Workbook()
    book.active.title = "Summary"
    book.active.append([])
    for i in range(35):
        book.active.append([f"row {i}", "<script>alert(1)</script>", *range(25)])
    second = book.create_sheet("Details")
    second.append(["Second sheet", "=1+2", "x" * 1000])
    path = tmp_path / "book.xlsx"
    book.save(path)
    preview = preview_function()(path)
    assert [s["name"] for s in preview["sheets"]] == ["Summary", "Details"]
    assert len(preview["rows"]) == 30
    assert len(preview["rows"][0]["cells"]) == 20
    assert preview["rows"][0]["number"] == 2
    assert preview["rows"][0]["cells"][1] == "<script>alert(1)</script>"
    assert preview["limited"] is True
    detail = preview_function()(path, sheet_index=1)
    assert detail["rows"][0]["cells"][:2] == ["Second sheet", "=1+2"]
    assert len(detail["rows"][0]["cells"][2]) <= 500
    with pytest.raises(ValueError):
        preview_function()(path, sheet_index=2)


def test_preview_resolves_shared_strings_and_ignores_stale_dimensions(tmp_path):
    path = tmp_path / "shared.xlsx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("xl/workbook.xml", '<workbook xmlns:r="urn:r"><sheets><sheet name="Shared" r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", '<Relationships><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
        z.writestr("xl/sharedStrings.xml", '<sst><si><t>Header</t></si><si><r><t>Rich </t></r><r><t>text</t></r></si></sst>')
        z.writestr("xl/worksheets/sheet1.xml", '<worksheet><dimension ref="A1"/><sheetData><row r="9"><c r="B9" t="s"><v>1</v></c><c r="A9" t="s"><v>0</v></c></row></sheetData></worksheet>')
    preview = preview_function()(path)
    assert preview["rows"] == [{"number": 9, "cells": ["Header", "Rich text"]}]
    assert preview["limited"] is False


def test_preview_byte_budget_is_reported(tmp_path):
    book = Workbook()
    book.active.append(["x" * 2000])
    path = tmp_path / "book.xlsx"
    book.save(path)
    assert preview_function()(path, scan_bytes=256)["limited"] is True


def test_shared_text_survives_exhausted_worksheet_budget(tmp_path):
    path = tmp_path / "tail.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("xl/workbook.xml", '<workbook xmlns:r="urn:r"><sheets><sheet name="Main" r:id="rId1"/></sheets></workbook>')
        archive.writestr("xl/_rels/workbook.xml.rels", '<Relationships><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
        archive.writestr("xl/sharedStrings.xml", '<sst><si><t>Visible header</t></si></sst>')
        archive.writestr("xl/worksheets/sheet1.xml", '<worksheet><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c></row>' + '<row r="2"/>' * 500 + '</sheetData></worksheet>')
    preview = preview_function()(path, scan_bytes=500)
    assert preview["limited"] is True
    assert preview["rows"][0]["cells"] == ["Visible header"]


@pytest.mark.parametrize("epoch_1904", [False, True])
def test_preview_raw_value_contract_dates_percent_codes_booleans(tmp_path, epoch_1904):
    from datetime import datetime
    from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel

    book = Workbook()
    if epoch_1904:
        book.epoch = CALENDAR_MAC_1904
    sheet = book.active
    sheet.append([datetime(2026, 10, 17), .15, 42, True, "=1+2", "<b>literal</b>"])
    sheet["B1"].number_format = "0%"
    sheet["C1"].number_format = "000000"
    path = tmp_path / "raw-values.xlsx"
    book.save(path)
    data = preview_function()(path)
    assert data["limited"] is False
    assert data["rows"][0]["cells"] == [str(int(to_excel(datetime(2026, 10, 17), book.epoch))),
                                          "0.15", "42", "1", "=1+2", "<b>literal</b>"]


def test_preview_distinguishes_empty_sheet_from_data_beyond_visible_columns(tmp_path):
    book = Workbook()
    path = tmp_path / "empty.xlsx"
    book.save(path)
    assert preview_function()(path)["rows"] == []
    book.active["U9"] = "outside sample"
    book.save(path)
    data = preview_function()(path)
    assert data["rows"] == [{"number": 9, "cells": []}]
    assert data["limited"] is True


def test_excluded_attachment_preview_is_read_only(upload_client, monkeypatch):
    import app.api.document_processing_scope as api

    client, settings, started = upload_client
    settings.attachment_review_enabled = True
    settings.dedup_enabled = False
    monkeypatch.setattr(api, "get_settings", lambda: settings)
    doc = scope_fixtures.upload_mail(client, 600).json()
    route = f"/api/documents/{doc['id']}/processing-scope"
    before = client.get(route).json()
    review = before["pending"]
    node = review["inventory"]["nodes"][0]
    url = route + "/attachments/" + node["source_key"] + "/preview"
    response = client.get(url, params={"review_id": review["id"]})
    assert response.status_code == 200, response.text
    assert response.headers["cache-control"] == "no-store"
    assert "content-disposition" not in response.headers
    assert response.json()["rows"][0]["cells"] == ["TABLE_CANARY"]
    assert response.json()["limited"] is True
    assert client.get(route).json() == before
    assert started == []
    assert client.get(url, params={"review_id": "stale"}).status_code == 409
    assert client.get(url, params={"review_id": review["id"], "sheet": 999}).status_code == 422
    assert client.get(route + "/attachments/" + "f" * 64 + "/preview", params={"review_id": review["id"]}).status_code == 404
    client.cookies.clear()
    assert client.get(url, params={"review_id": review["id"]}).status_code == 401


@pytest.mark.parametrize("error_name,status", [
    ("ParserBusyError", 429), ("ParserTimeoutError", 422),
    ("ParserMemoryLimitError", 422), ("ParserIsolationError", 503), ("ParserWorkerError", 422),
])
def test_preview_worker_errors_leave_review_unchanged(upload_client, monkeypatch, error_name, status):
    import app.api.document_processing_scope as api
    import app.services.parser_supervisor as supervisor

    client, settings, started = upload_client
    settings.attachment_review_enabled = True
    settings.dedup_enabled = False
    monkeypatch.setattr(api, "get_settings", lambda: settings)
    doc = scope_fixtures.upload_mail(client, 600).json()
    route = f"/api/documents/{doc['id']}/processing-scope"
    before = client.get(route).json()
    review = before["pending"]
    key = review["inventory"]["nodes"][0]["source_key"]

    def fail(*args, **kwargs):
        raise getattr(supervisor, error_name)("PRIVATE_PROVIDER_DETAILS")

    monkeypatch.setattr(supervisor, "parse_document_supervised", fail)
    response = client.get(f"{route}/attachments/{key}/preview", params={"review_id": review["id"]})
    assert response.status_code == status
    assert "PRIVATE_PROVIDER_DETAILS" not in response.text
    assert client.get(route).json() == before
    assert started == []


@pytest.mark.parametrize("kind", ["zip", "7z", "rar"])
def test_preview_excel_inside_nested_mail_and_archive(upload_client, monkeypatch, kind):
    from email.message import EmailMessage
    import py7zr
    import app.api.document_processing_scope as api
    from tests.test_processing_scope_end_to_end import stored_rar

    client, settings, started = upload_client
    settings.attachment_review_enabled = True
    settings.dedup_enabled = False
    monkeypatch.setattr(api, "get_settings", lambda: settings)
    book = Workbook()
    for _ in range(600):
        book.active.append(["NESTED_PREVIEW_CANARY"])
    data = io.BytesIO()
    book.save(data)
    packed = io.BytesIO()
    if kind == "zip":
        with zipfile.ZipFile(packed, "w") as archive:
            archive.writestr("nested.xlsx", data.getvalue())
    elif kind == "7z":
        with py7zr.SevenZipFile(packed, "w") as archive:
            archive.writestr(data.getvalue(), "nested.xlsx")
    else:
        packed.write(stored_rar([("nested.xlsx", data.getvalue())]))
    inner = EmailMessage()
    inner.set_content("Inner letter")
    inner.add_attachment(packed.getvalue(), maintype="application", subtype="octet-stream", filename="data." + kind)
    outer = EmailMessage()
    outer.set_content("Outer letter")
    outer.add_attachment(inner.as_bytes(), maintype="application", subtype="octet-stream", filename="inner.eml")
    response = client.post("/api/documents", files={"file": ("outer.eml", outer.as_bytes())})
    assert response.status_code == 202, response.text
    route = f"/api/documents/{response.json()['id']}/processing-scope"
    before = client.get(route).json()
    review = before["pending"]
    node = next(n for n in review["inventory"]["nodes"] if n["kind"] == "xlsx")
    assert not next(d for d in review["decisions"] if d["source_key"] == node["source_key"])["effective_included"]
    preview = client.get(route + "/attachments/" + node["source_key"] + "/preview", params={"review_id": review["id"]})
    assert preview.status_code == 200, preview.text
    assert preview.json()["rows"][0]["cells"] == ["NESTED_PREVIEW_CANARY"]
    assert client.get(route).json() == before
    assert not started

"""Search/merge must retain exact fragments and immutable table coordinates."""

import pytest
from openpyxl import Workbook

from app.db.models import DocumentGenerationState, OkfConcept
from app.db.session import session_scope
from app.services.fusion import Hit
from app.services.retrieval_hydration import load_visible_retrieval_hits
from app.services.context_builder import merge_and_format
from app.services.chat_source_selection import snapshot_blocks, restore_blocks
from tests.test_generation_pipeline import pipeline_env as _pipeline_env

pipeline_env = _pipeline_env

DOC = 'b123456789abcdef'


def publish(pipeline, tmp_path):
    pipeline.settings.okf_table_llm_classify = False
    wb = Workbook()
    ws = wb.active
    ws.title = 'Directory'
    ws.append(['Code', 'Quantity'])
    ws.append(['same-key', 7])
    source = tmp_path / 'tables.xlsx'
    wb.save(source)
    pipeline.registry.create(DOC, 'tables.xlsx', 'application/octet-stream', source.stat().st_size)
    pipeline._process(DOC, source, 'tables.xlsx', [], resume=False)
    assert pipeline.registry.get(DOC)['status'] == 'done'
    with session_scope() as s:
        active = s.get(DocumentGenerationState, DOC).active_generation_id
        concept = s.query(OkfConcept).filter_by(doc_id=DOC, chunk_index=1).one()
        return active, concept.slug, concept.chunk_index


def test_hydration_does_not_clip_table_fragment_or_trust_payload_refs(pipeline_env, tmp_path):
    pipeline, _, _ = pipeline_env
    generation, slug, index = publish(pipeline, tmp_path)
    hits = [Hit('c', 1, {'doc_id': DOC, 'point_type': 'concept', 'slug': slug,
                         'generation_id': generation, 'table_refs': [{'table_id': 'forged'}]})]
    kept, _ = load_visible_retrieval_hits(hits, max_concept_chars=10)
    assert kept and 'B2' in kept[0].payload['content'] and '7' in kept[0].payload['content']
    assert kept[0].payload['table_refs'][0]['generation_id'] == generation
    assert kept[0].payload['table_refs'][0]['table_id'] != 'forged'
    blocks = merge_and_format(kept, settings=pipeline.settings)
    assert blocks[0]['table_refs'] == kept[0].payload['table_refs']


def test_distinct_row_parts_are_not_merged_with_giant_chunk():
    refs = [{'kind': 'table', 'generation_id': 'a' * 32, 'source_id': 'root', 'table_id': 'b' * 32,
             'row_number': 2, 'columns': [1]}]
    hits = [Hit('first', 2, {'doc_id': DOC, 'point_type': 'concept', 'slug': 'first', 'title': 'First',
            'chunk_index': 1, 'content': 'A2 exact value', 'table_refs': refs, '_canonical_verified': True}),
            Hit('second', 1, {'doc_id': DOC, 'point_type': 'concept', 'slug': 'second', 'title': 'Second',
            'chunk_index': 1, 'content': 'B2 other value', 'table_refs': [{**refs[0], 'columns': [2]}], '_canonical_verified': True}),
            Hit('chunk', 3, {'doc_id': DOC, 'point_type': 'chunk', 'chunk_index': 1, 'content': 'giant unrelated chunk'})]
    blocks = merge_and_format(hits)
    assert any(b['content'] == 'A2 exact value' and b.get('table_refs') == refs for b in blocks)
    assert any(b['content'] == 'B2 other value' for b in blocks)


def test_table_snapshot_uses_saved_fragment_and_detects_missing_generation(pipeline_env, tmp_path):
    pipeline, _, _ = pipeline_env
    generation, slug, index = publish(pipeline, tmp_path)
    kept, _ = load_visible_retrieval_hits([Hit('c', 1, {'doc_id': DOC, 'point_type': 'concept', 'slug': slug,
                                                      'generation_id': generation})])
    blocks = merge_and_format(kept, settings=pipeline.settings)
    assert blocks[0].get('table_refs'), 'table refs must reach snapshots'
    saved = snapshot_blocks(blocks)
    with session_scope() as s:
        s.query(OkfConcept).filter_by(doc_id=DOC, slug=slug).one().content = 'new projection'
    restored = restore_blocks(saved, mail_mode='all')
    assert restored[0]['content'] == blocks[0]['content']
    from app.db.models import DocumentTableRow, DocumentTable
    with session_scope() as s:
        s.query(DocumentTableRow).filter_by(doc_id=DOC, generation_id=generation).delete()
        s.query(DocumentTable).filter_by(doc_id=DOC, generation_id=generation).delete()
    with pytest.raises(Exception):
        restore_blocks(saved, mail_mode='all')



def _rank_table_blocks(blocks, query):
    from app.services.table_retrieval import prioritize_table_identifier_matches
    return prioritize_table_identifier_matches(blocks, query)


def _table_block(row, code, business_unit, *, prefix='', kind='row'):
    return {'title': 'Directory', 'doc_id': DOC, 'content': prefix + f'A{row} | Код: {code}\nB{row} | Код БЕ: {business_unit}',
            '_table_identifier_cells': [{'value': code, 'labels': ['Код'], 'row_number': row, 'table_id': 'b' * 32},
                {'value': business_unit, 'labels': ['Код БЕ'], 'row_number': row, 'table_id': 'b' * 32}],
            'table_refs': [{'generation_id': 'a' * 32, 'source_id': 'root', 'table_id': 'b' * 32,
                            'row_number': row, 'columns': [1, 2], 'kind': kind}]}


def test_table_code_and_company_reach_context_before_cutoff():
    blocks = [_table_block(42, '7601', '2400'), _table_block(54, '7601', '2500'),
              _table_block(32, '7601', '3100'), _table_block(92, '7601', '3100')]
    ranked = _rank_table_blocks(blocks, 'Для кода 7601 и БЕ 3100 какое значение?')
    assert [b['table_refs'][0]['row_number'] for b in ranked[:2]] == [32, 92]
    assert len(ranked) == len(blocks) and {id(b) for b in ranked} == {id(b) for b in blocks}


def test_table_identifier_priority_preserves_prose_and_equal_order():
    prose1 = {'content': '7601 3100', 'kind': 'concept'}
    prose2 = {'content': 'Ordinary mail and review evidence', 'kind': 'review'}
    wrong = _table_block(42, '7601', '2400')
    right1, right2 = _table_block(32, '7601', '3100'), _table_block(92, '7601', '3100')
    blocks = [prose1, wrong, prose2, right1, right2]
    ranked = _rank_table_blocks(blocks, 'Код 7601 для БЕ 3100')
    assert ranked[0] is prose1 and ranked[2] is prose2
    assert [b for b in ranked if b.get('table_refs')] == [right1, right2, wrong]


def test_table_identifier_priority_ignores_row_metadata_and_substrings():
    wrong_row = _table_block(3100, '7601', '2400', prefix='Sheet: 3100; row: 3100; key: 3100\n')
    wrong_substring = _table_block(2, '7601', '3100005000')
    right = _table_block(3, '7601', '3100')
    ranked = _rank_table_blocks([wrong_row, wrong_substring, right], 'Код 7601 и БЕ 3100')
    assert ranked[0] is right


def test_single_identifier_and_non_table_queries_keep_original_order():
    blocks = [_table_block(42, '7601', '2400'), _table_block(32, '7601', '3100')]
    assert _rank_table_blocks(blocks, 'Код 3100') == blocks
    assert _rank_table_blocks(blocks, 'Сравни компании') == blocks


def test_table_identifier_priority_applies_before_character_budget():
    from app.config import Settings
    wrong = _table_block(42, '7601', '2400')
    right = _table_block(32, '7601', '3100')
    def hit(name, score, block):
        return Hit(name, score, {**block, 'point_type': 'chunk', 'chunk_index': int(name),
                                'section_title': 'Directory', '_canonical_verified': True})
    settings = Settings(_env_file=None, chat_max_context_chars=len(wrong['content']))
    merged = merge_and_format([hit('1', 2, wrong), hit('2', 1, right)], settings=settings, query='Код 7601 для БЕ 3100')
    assert len(merged) == 1 and merged[0]['table_refs'][0]['row_number'] == 32



def test_decimal_signed_and_header_numbers_are_not_exact_identifier_values():
    decimal = _table_block(2, '7601.50', '3100')
    signed = _table_block(3, '-7601', '3100')
    header_only = _table_block(4, '7601', '2400')
    header_only['content'] += '\nC4 | Year 3100: unrelated'
    right = _table_block(5, '7601', '3100')
    assert _rank_table_blocks([decimal, signed, header_only, right], 'Код 7601 для БЕ 3100')[0] is right


def test_amount_value_does_not_confirm_company_qualifier():
    wrong = _table_block(2, '7601', '2400')
    wrong['content'] += '\nC2 | Amount: 3100'
    wrong['_table_identifier_cells'] = [{'value': '7601', 'labels': ['Код'], 'row_number': 2, 'table_id': 'b' * 32},
        {'value': '2400', 'labels': ['Код БЕ'], 'row_number': 2, 'table_id': 'b' * 32}, {'value': '3100', 'labels': ['Amount'], 'row_number': 2, 'table_id': 'b' * 32}]
    right = _table_block(3, '7601', '3100')
    assert _rank_table_blocks([wrong, right], 'Код 7601 для БЕ 3100')[0] is right


@pytest.mark.parametrize('literal', ['7601.50', '-7601', '+7601'])
def test_non_integer_query_does_not_promote_positive_identifier(literal):
    first = _table_block(2, '2410', '2400')
    second = _table_block(3, '7601', '3100')
    blocks = [first, second]
    assert _rank_table_blocks(blocks, f'Код {literal} БЕ 3100') == blocks


def test_identifier_qualification_comes_from_canonical_cells_not_payload(pipeline_env, tmp_path):
    pipeline, _, _ = pipeline_env
    pipeline.settings.okf_table_llm_classify = False
    wb = Workbook();ws = wb.active;ws.title = 'Directory'
    ws.append(['Code', 'Company', 'Year 3100'])
    for row in ([7601.50, 3100, 1], [-7601, 3100, 1], [7601, 2400, 3100], [7601, 3100, 1]):
        ws.append(row)
    source = tmp_path / 'numeric-tables.xlsx';wb.save(source)
    pipeline.registry.create(DOC, 'numeric-tables.xlsx', 'application/octet-stream', source.stat().st_size)
    pipeline._process(DOC, source, 'numeric-tables.xlsx', [], resume=False)
    assert pipeline.registry.get(DOC)['status'] == 'done'
    with session_scope() as session:
        generation = session.get(DocumentGenerationState, DOC).active_generation_id
        slugs = {c.chunk_index: c.slug for c in session.query(OkfConcept).filter_by(doc_id=DOC)}
    hits = [Hit(str(index), float(10-index), {'doc_id': DOC, 'point_type': 'concept', 'slug': slugs[index],
            'generation_id': generation, '_table_identifier_cells': [{'value': '7601', 'labels': ['Code']},
                {'value': '3100', 'labels': ['Company']}]}) for index in sorted(slugs)]
    kept, _ = load_visible_retrieval_hits(hits, query="Code 7601 company 3100")
    blocks = merge_and_format(kept, settings=pipeline.settings, query='Code 7601 company 3100')
    assert blocks[0]['table_refs'][0]['row_number'] == 5
    decimal = next(h for h in kept if h.payload['table_refs'][0]['row_number'] == 2)
    assert not any(c['value'] == '7601' for c in decimal.payload['_table_identifier_cells'])



def test_identifiers_from_different_rows_cannot_confirm_one_qualified_row():
    wrong = _table_block(2, '7601', '2400')
    other = _table_block(3, '7701', '3100')
    wrong['_table_identifier_cells'] += other['_table_identifier_cells']
    wrong['content'] += '\n' + other['content']
    right = _table_block(4, '7601', '3100')
    assert _rank_table_blocks([wrong, right], 'Код 7601 для БЕ 3100')[0] is right


def test_shared_header_caption_does_not_confirm_specific_column():
    wrong = _table_block(2, '7601', '2400')
    right = _table_block(3, '7601', '3100')
    for block, amount in [(wrong, '3100'), (right, '1000')]:
        row = block['table_refs'][0]['row_number']
        block['_table_identifier_cells'] = [
            {'value': block['_table_identifier_cells'][0]['value'], 'labels': ['Company directory', 'Code'], 'row_number': row, 'table_id': 'b' * 32},
            {'value': block['_table_identifier_cells'][1]['value'], 'labels': ['Company directory', 'Company'], 'row_number': row, 'table_id': 'b' * 32},
            {'value': amount, 'labels': ['Company directory', 'Amount'], 'row_number': row, 'table_id': 'b' * 32}]
    assert _rank_table_blocks([wrong, right], 'Code 7601 company 3100')[0] is right


@pytest.mark.parametrize('query', ['', 'Unqualified 7601', 'Unknown 7601 company 3100'])
def test_unqualified_query_avoids_extra_canonical_row_reads(pipeline_env, tmp_path, query):
    from sqlalchemy import event
    from app.db.session import get_engine
    pipeline, _, _ = pipeline_env
    generation, slug, _ = publish(pipeline, tmp_path)
    statements = []
    def observe(_connection, _cursor, statement, *_args):
        if statement.lstrip().upper().startswith('SELECT') and 'FROM document_table_rows' in statement:
            statements.append(statement)
    event.listen(get_engine(), 'before_cursor_execute', observe)
    try:
        kept, _ = load_visible_retrieval_hits([Hit('c', 1., {'doc_id': DOC, 'point_type': 'concept',
            'slug': slug, 'generation_id': generation})], query=query)
    finally:
        event.remove(get_engine(), 'before_cursor_execute', observe)
    assert kept and kept[0].payload['table_refs']
    assert statements == [], 'Unqualified search must not load full canonical table rows'


def test_single_field_identifier_precedes_fast_context_cutoff_with_quoted_year():
    wrong = _table_block(2, '7602', '3100')
    right = _table_block(3, '7601', '2400')
    ranked = _rank_table_blocks([wrong, right], 'Для кода 7601 значение поля «Rules for 2026»?')
    assert ranked[0] is right


def test_single_field_identifier_uses_specific_leaf_over_generic_shared_word():
    wrong = _table_block(2, '7602', '7601')
    right = _table_block(3, '7601', '2400')
    assert _rank_table_blocks([wrong, right], 'Код 7601')[0] is right


def test_invalid_signed_query_cannot_promote_remaining_company_alone():
    first = _table_block(2, '7602', '2400')
    second = _table_block(3, '7601', '3100')
    blocks = [first, second]
    assert _rank_table_blocks(blocks, 'Код -7601 для БЕ 3100') == blocks


def test_single_code_qualification_uses_descriptor_even_when_cell_missing(pipeline_env, tmp_path):
    pipeline, _, _ = pipeline_env
    pipeline.settings.okf_table_llm_classify = False
    wb = Workbook();ws = wb.active
    ws.append(['Code', 'Code BE'])
    ws.append([None, 7601])
    ws.append([7601, 2400])
    source = tmp_path / 'missing-code.xlsx';wb.save(source)
    pipeline.registry.create(DOC, 'missing-code.xlsx', 'application/octet-stream', source.stat().st_size)
    pipeline._process(DOC, source, 'missing-code.xlsx', [], resume=False)
    with session_scope() as session:
        generation = session.get(DocumentGenerationState, DOC).active_generation_id
        entries = [(c.chunk_index,c.slug) for c in session.query(OkfConcept).filter_by(doc_id=DOC)]
    hits = [Hit(str(i),float(10-i),{'doc_id':DOC,'point_type':'concept','slug':slug,
        'generation_id':generation}) for i,slug in entries]
    kept,_ = load_visible_retrieval_hits(hits,query='Code 7601')
    blocks = merge_and_format(kept,settings=pipeline.settings,query='Code 7601')
    assert blocks[0]['table_refs'][0]['row_number']==3


def test_specific_code_leaf_in_other_header_group_does_not_hide_primary_code():
    wrong = _table_block(2, '7602', '3100')
    right = _table_block(3, '7601', '2400')
    for block in [wrong,right]:
        block['_table_identifier_cells'][0]['labels']=['Code VO']
        block['_table_identifier_cells'][1]['labels']=['Accounting','Code']
    assert _rank_table_blocks([wrong,right],'Code 7601')[0] is right


@pytest.mark.parametrize('literal',['-12','+12','12.50'])
def test_short_invalid_literal_cannot_promote_partial_company_match(literal):
    first = _table_block(2,'7602','2400')
    second = _table_block(3,'7601','3100')
    blocks = [first,second]
    assert _rank_table_blocks(blocks,f'Код {literal} БЕ 3100')==blocks


def test_company_name_leaf_does_not_hide_company_code_suffix(pipeline_env, tmp_path):
    pipeline, _, _ = pipeline_env
    pipeline.settings.okf_table_llm_classify=False
    wb=Workbook();ws=wb.active
    ws.append(['Code','BE','Code BE'])
    ws.append([7601,'Company Alpha',2400])
    ws.append([7601,'Company Beta',3100])
    source=tmp_path/'company-code.xlsx';wb.save(source)
    pipeline.registry.create(DOC,'company-code.xlsx','application/octet-stream',source.stat().st_size)
    pipeline._process(DOC,source,'company-code.xlsx',[],resume=False)
    with session_scope() as session:
        generation=session.get(DocumentGenerationState,DOC).active_generation_id
        entries=[(c.chunk_index,c.slug) for c in session.query(OkfConcept).filter_by(doc_id=DOC)]
    hits=[Hit(str(i),float(10-i),{'doc_id':DOC,'point_type':'concept','slug':slug,'generation_id':generation}) for i,slug in entries]
    kept,_=load_visible_retrieval_hits(hits,query='Code 7601 BE 3100')
    blocks=merge_and_format(kept,settings=pipeline.settings,query='Code 7601 BE 3100')
    assert blocks[0]['table_refs'][0]['row_number']==3


@pytest.mark.parametrize("header_count", [0, 1])
def test_parsed_numeric_code_and_headerless_first_record_reach_query_context(pipeline_env, tmp_path, header_count):
    import zipfile
    from openpyxl.worksheet.table import Table, TableColumn
    pipeline, _, _ = pipeline_env
    pipeline.settings.okf_table_llm_classify = False
    wb = Workbook()
    ws = wb.active
    ws.title = "Контроль"
    if header_count:
        ws.append(["Код", "Значение"])
    row = 1 + header_count
    ws.append([7, 42])
    ws.append([8, 43])
    ws.cell(row, 1).number_format = "00000"
    ws.add_table(Table(displayName="Directory", ref=f"A1:B{row + 1}", headerRowCount=header_count,
                       tableColumns=[TableColumn(id=1, name="Код"), TableColumn(id=2, name="Значение")]))
    source = tmp_path / "codes.xlsx"
    wb.save(source)
    with zipfile.ZipFile(source) as archive:
        entries = {n: archive.read(n) for n in archive.namelist()}
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(b"<v>7</v>", b"<v>7.0</v>")
    with zipfile.ZipFile(source, "w") as archive:
        for name, payload in entries.items():
            archive.writestr(name, payload)
    pipeline.registry.create(DOC, "codes.xlsx", "application/octet-stream", source.stat().st_size)
    pipeline._process(DOC, source, "codes.xlsx", [], resume=False)
    assert pipeline.registry.get(DOC)["status"] == "done"
    with session_scope() as session:
        generation = session.get(DocumentGenerationState, DOC).active_generation_id
        concepts = session.query(OkfConcept).filter_by(doc_id=DOC).all()
        assert any("00007" in c.content for c in concepts)
        hits = [Hit(c.slug, 1, {"doc_id":DOC, "point_type":"concept", "slug":c.slug, "generation_id":generation}) for c in concepts]
    query = f'На листе «Контроль» для кода 00007 (строка {row}) укажи значение поля «Значение» (колонка B).'
    kept, _ = load_visible_retrieval_hits(hits, query=query)
    assert any(cell["value"] == "00007" and cell["row_number"] == row for hit in kept for cell in hit.payload.get("_table_identifier_cells", []))
    exact = [hit for hit in kept if hit.payload.get("_table_cell_states")]
    assert len(exact) == 1
    assert f"B{row} | Значение: 42" in exact[0].payload["content"]
    assert all(ref["row_number"] == row and ref["generation_id"] == generation for ref in exact[0].payload["table_refs"])
    blocks = merge_and_format(exact, settings=pipeline.settings, query=query)
    restored = restore_blocks(snapshot_blocks(blocks), mail_mode="all")
    assert restored[0]["content"] == blocks[0]["content"]

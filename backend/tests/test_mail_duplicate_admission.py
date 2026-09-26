"""D05: mail twins are admitted atomically and trash does not block re-upload."""
from concurrent.futures import ThreadPoolExecutor, wait
from email.message import EmailMessage
from threading import Barrier, Event

import pytest
from sqlalchemy import select

from app.db.models import Document
from app.db.session import session_scope
from app.services.registry import get_registry
from tests.test_mail_upload_admission import FIXTURES, _client


def _payload(extension, revision=False):
    if extension == "msg":
        # An unused final CFB sector changes the file SHA, not the MAPI streams.
        original = (FIXTURES / "synthetic-unicode-attachment.msg").read_bytes()
        return original + (b"\0" * 512 if revision else b"")
    message = EmailMessage()
    message["From"] = "sender@example.test"
    message["Subject"] = "Согласование"
    message["Message-ID"] = "<decision@example.test>"
    message["X-Transport-Trace"] = "second" if revision else "first"
    message.set_content("Лимит проверки составляет двенадцать календарных дней после получения заявки.")
    return message.as_bytes()


@pytest.mark.parametrize("extension", ["eml", "msg"])
@pytest.mark.parametrize("identical", [True, False], ids=["exact", "semantic"])
def test_concurrent_mail_twins_have_one_admission(tmp_path, monkeypatch, extension, identical):
    from app.api import documents

    client, settings, started = _client(tmp_path, monkeypatch, dedup_enabled=True)
    both_parsed = Barrier(2)
    first_indexing, release_indexing = Event(), Event()
    real_parse, real_index = documents.parse_document, documents.index_document

    def parse(*args, **kwargs):
        result = real_parse(*args, **kwargs)
        both_parsed.wait(timeout=15)
        return result

    def index(*args):
        if not first_indexing.is_set():
            first_indexing.set()
            assert release_indexing.wait(15)
        return real_index(*args)

    monkeypatch.setattr(documents, "parse_document", parse)
    monkeypatch.setattr(documents, "index_document", index)
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(client.post, "/api/documents", files={"file": (
            f"copy-{i}.{extension}", _payload(extension, revision=bool(i) and not identical),
        )}) for i in range(2)]
        try:
            assert first_indexing.wait(15)
            wait(futures, timeout=0.3)
        finally:
            release_indexing.set()
        responses = [future.result(timeout=15) for future in futures]
    assert sorted(response.status_code for response in responses) == [200, 409]
    rejection = next(response.json() for response in responses if response.status_code == 409)
    assert rejection["code"] == ("duplicate" if identical else "similar_document")
    with session_scope() as session:
        rows = session.scalars(select(Document)).all()
        assert len(rows) == 1
        assert started == [rows[0].id]
    assert len(list(settings.uploads_dir.iterdir())) == 1


@pytest.mark.parametrize("extension", ["eml", "msg"])
@pytest.mark.parametrize("identical", [True, False], ids=["exact", "semantic"])
def test_trashed_mail_twin_does_not_reject_reupload(tmp_path, monkeypatch, extension, identical):
    client, _settings, started = _client(tmp_path, monkeypatch, dedup_enabled=True)
    first = client.post("/api/documents", files={"file": (f"first.{extension}", _payload(extension))})
    assert first.status_code == 200, first.text
    first_id = first.json()["id"]
    get_registry().soft_delete(first_id, "demo.editor")
    second = client.post("/api/documents", files={"file": (
        f"second.{extension}", _payload(extension, revision=not identical),
    )})
    assert second.status_code == 200, second.text
    assert second.json()["id"] != first_id
    assert not second.json()["has_duplicates"]
    assert started == [first_id, second.json()["id"]]


def test_same_msg_standalone_and_inside_docx_keeps_both_document_origins(tmp_path, monkeypatch):
    from io import BytesIO

    from docx import Document as DocxDocument
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from lxml import etree

    from app.db.models import DocumentSource, OkfConcept
    from app.models.schemas import Concept
    from app.services.pipeline import Pipeline
    from app.services.source_store import fetch_source_paths

    client, settings, _started = _client(tmp_path, monkeypatch, dedup_enabled=True)
    settings.embedding_provider = "fake"
    settings.embedding_dimensions = 8
    payload = _payload("msg")
    word = DocxDocument()
    word.add_paragraph("Контейнер исходного письма.")
    part = Part(PackURI("/word/embeddings/decision.msg"), "application/vnd.ms-outlook", payload, word.part.package)
    relation_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rid = word.part.relate_to(part, relation_ns + "/oleObject")
    run = etree.SubElement(word.add_paragraph()._p, "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r")
    obj = etree.SubElement(run, "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}object")
    ole = etree.SubElement(obj, "{urn:schemas-microsoft-com:office:office}OLEObject")
    ole.set("Type", "Embed")
    ole.set("ProgID", "Outlook.File.msg.15")
    ole.set(f"{{{relation_ns}}}id", rid)
    buffer = BytesIO()
    word.save(buffer)

    for module in ("app.services.staging", "app.services.embedder", "app.services.llm_client",
                   "app.services.okf_generator", "app.services.vector_store"):
        monkeypatch.setattr(f"{module}.get_settings", lambda: settings)
    pipeline = Pipeline()
    pipeline.okf_generator.generate_chunk = lambda text, *args, **kwargs: [
        Concept(id="source", title="Исходный фрагмент", content=text),
    ]
    for method in ("ensure_collection", "delete_document", "delete_orphaned_points"):
        monkeypatch.setattr(pipeline.vector_store, method, lambda *args, **kwargs: None)
    for method in ("index_concepts", "index_chunks"):
        monkeypatch.setattr(pipeline.vector_store, method, lambda *args, **kwargs: set())

    ids = []
    for filename, data in (("direct.msg", payload), ("container.docx", buffer.getvalue())):
        response = client.post("/api/documents", files={"file": (filename, data)}, data={"allow_similar": "true"})
        assert response.status_code == 200, response.text
        doc_id = response.json()["id"]
        ids.append(doc_id)
        pipeline._process(doc_id, settings.uploads_dir / f"{doc_id}.{filename.split('.')[-1]}", filename, [], resume=False)
        assert get_registry().get(doc_id)["status"] == "done"
    assert ids[0] != ids[1]
    with session_scope() as session:
        paths = fetch_source_paths(session, {(ids[0], "root"), (ids[1], "root/0")})
        assert [node["display_name"] for node in paths[(ids[0], "root")]] == ["direct.msg"]
        assert paths[(ids[1], "root/0")][0]["display_name"] == "container.docx"
        assert paths[(ids[1], "root/0")][-1]["mail"] is True
        for doc_id, source_id in ((ids[0], "root"), (ids[1], "root/0")):
            assert session.query(OkfConcept).filter_by(doc_id=doc_id, source_id=source_id).count() > 0
        nested = session.query(DocumentSource).filter_by(doc_id=ids[1], source_id="root/0").one()
        assert (settings.uploads_dir / ids[1] / nested.saved_path).read_bytes() == payload
        assert {(row.id, row.filename) for row in session.scalars(select(Document))} == {
            (ids[0], "direct.msg"), (ids[1], "container.docx"),
        }

"""Upload admission must ask before processing similar content."""
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument
from sqlalchemy import select

from app.db.models import Document
from app.db.session import session_scope
from app.services.registry import get_registry
from tests.test_upload_development import make_client, login


TEXT = "Порядок согласования заявок и обработки документов в информационной системе предприятия."


def docx_bytes(revision, text=TEXT):
    doc = DocxDocument()
    doc.core_properties.title = revision
    doc.add_paragraph(text)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


@pytest.fixture
def upload_client(tmp_path, monkeypatch):
    from app.api import documents
    client = make_client(tmp_path, monkeypatch, dedup_enabled=True)
    settings = documents.get_settings()
    for module in (documents, __import__("app.services.pipeline", fromlist=[""]),
                   __import__("app.services.deduplication", fromlist=[""])):
        monkeypatch.setattr(module, "get_settings", lambda: settings)
    started = []
    monkeypatch.setattr(documents, "get_pipeline", lambda: SimpleNamespace(
        ingest=lambda doc_id, *a, **kw: started.append(doc_id)))
    login(client)
    return client, settings, started


def upload(client, revision, **data):
    return client.post("/api/documents", data=data,
                       files={"file": (f"{revision}.docx", docx_bytes(revision),
                                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})


def test_similar_upload_waits_without_document_file_tags_or_pipeline(upload_client):
    client, settings, started = upload_client
    first = upload(client, "first")
    assert first.status_code == 200, first.text
    before = set(settings.uploads_dir.rglob("*"))

    second = upload(client, "revision", tags="must-not-be-created")
    assert second.status_code == 409, second.text
    assert second.json()["code"] == "similar_document"
    candidate = second.json()["duplicates"]["level2"][0]
    assert candidate["doc"]["id"] == first.json()["id"]
    assert candidate["doc"]["created_at"]
    assert candidate["jaccard"] == 1.0
    with session_scope() as s:
        assert len(s.scalars(select(Document)).all()) == 1
    assert set(settings.uploads_dir.rglob("*")) == before
    assert started == [first.json()["id"]]
    assert not get_registry().get(first.json()["id"])["has_duplicates"]
    from app.services.tag_registry import TagRegistry
    assert "must-not-be-created" not in {t["name"] for t in TagRegistry().all()}


def test_confirmation_admits_and_marks_both_even_before_pipeline_starts(upload_client):
    client, settings, started = upload_client
    first = upload(client, "first").json()
    assert upload(client, "revision").status_code == 409
    accepted = upload(client, "revision", allow_similar="true")
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["has_duplicates"]
    assert get_registry().get(first["id"])["has_duplicates"]
    assert started == [first["id"], accepted.json()["id"]]


def test_confirmation_never_bypasses_identical_file(upload_client):
    client, settings, started = upload_client
    data = docx_bytes("same")
    kwargs = {"files": {"file": ("same.docx", data)}}
    assert client.post("/api/documents", **kwargs).status_code == 200
    rejected = client.post("/api/documents", data={"allow_similar": "true"}, **kwargs)
    assert rejected.status_code == 409
    assert rejected.json()["code"] == "duplicate"
    assert len(started) == 1


def test_trash_similarity_does_not_block(upload_client):
    client, settings, started = upload_client
    first = upload(client, "first").json()
    get_registry().soft_delete(first["id"], "demo.editor")
    second = upload(client, "revision")
    assert second.status_code == 200, second.text
    assert not second.json()["has_duplicates"]


def test_bad_file_does_not_silently_bypass_check(upload_client):
    client, settings, started = upload_client
    response = client.post("/api/documents", files={"file": ("broken.docx", b"not a zip")})
    assert response.status_code == 422, response.text
    assert not list(settings.uploads_dir.iterdir())
    assert not started


def test_malformed_root_eml_is_rejected_before_admission(upload_client, monkeypatch):
    client, settings, started = upload_client
    monkeypatch.setattr(settings, "mail_import_enabled", True)

    response = client.post(
        "/api/documents",
        files={"file": ("broken.eml", b"Subject: alone is not a message\n\nopaque bytes")},
    )

    assert response.status_code == 422, response.text
    with session_scope() as session:
        assert not session.scalars(select(Document)).all()
    assert not list(settings.uploads_dir.iterdir())
    assert not started


def test_parser_slot_exhaustion_returns_retryable_429_without_publishing_upload(upload_client, monkeypatch):
    from app.api import documents
    from app.services.parser_supervisor import ParserBusyError

    client, settings, started = upload_client

    def busy(*args, **kwargs):
        raise ParserBusyError("all slots occupied")

    monkeypatch.setattr(documents, "parse_document_supervised", busy)
    response = upload(client, "busy")

    assert response.status_code == 429
    assert response.json()["code"] == "rate_limited"
    assert response.headers["retry-after"] == "1"
    assert not list(settings.uploads_dir.iterdir())
    assert not started


@pytest.mark.parametrize("failure", ["signature", "queue"])
def test_failed_admission_removes_file_row_and_existing_badge(upload_client, monkeypatch, failure):
    from app.api import documents
    from app.services.errors import DomainError

    client, settings, started = upload_client
    first = upload(client, "first").json()
    before = set(settings.uploads_dir.rglob("*"))
    if failure == "signature":
        real_index = documents.index_document

        def fail_index(*args):
            real_index(*args)
            raise RuntimeError("signature persistence failed")

        monkeypatch.setattr(documents, "index_document", fail_index)
    else:
        def reject(*args, **kwargs):
            raise DomainError("queue full", code="queue_overloaded")

        monkeypatch.setattr(documents, "get_pipeline", lambda: SimpleNamespace(ingest=reject))
    client.raise_server_exceptions = False
    response = upload(client, "revision", allow_similar="true")
    assert response.status_code == (500 if failure == "signature" else 503)
    with session_scope() as s:
        assert [d.id for d in s.scalars(select(Document))] == [first["id"]]
    assert set(settings.uploads_dir.rglob("*")) == before
    assert not get_registry().get(first["id"])["has_duplicates"]


@pytest.mark.parametrize("identical", [False, True])
def test_concurrent_uploads_recheck_before_admission(upload_client, monkeypatch, identical):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier, Event
    from app.api import documents

    client, settings, started = upload_client
    first_indexing, release_first = Event(), Event()
    both_parsed = Barrier(2)
    real_index, real_parse = documents.index_document, documents.parse_document
    data = docx_bytes("first")

    def delayed_index(*args):
        if not first_indexing.is_set():
            first_indexing.set()
            assert release_first.wait(10)
        return real_index(*args)

    def observed_parse(*args, **kwargs):
        result = real_parse(*args, **kwargs)
        # Both initial SHA checks see no row; the final check must run again
        # under the admission lock after these parallel parses finish.
        both_parsed.wait(timeout=10)
        return result

    monkeypatch.setattr(documents, "index_document", delayed_index)
    monkeypatch.setattr(documents, "parse_document", observed_parse)
    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(client.post, "/api/documents", files={"file": ("first.docx", data)})
        second = pool.submit(client.post, "/api/documents", files={"file": (
            "second.docx", data if identical else docx_bytes("second"))})
        try:
            assert first_indexing.wait(10)
            # Without the lock the competing request can finish while the
            # first signature is still unpublished.
            from concurrent.futures import wait
            wait([first, second], timeout=0.3)
        finally:
            release_first.set()
        responses = [first.result(), second.result()]
    assert sorted(r.status_code for r in responses) == [200, 409]
    rejected = next(r for r in responses if r.status_code == 409)
    assert rejected.json()["code"] == ("duplicate" if identical else "similar_document")
    assert len(started) == 1


def test_changed_text_still_requires_confirmation(upload_client):
    client, settings, started = upload_client
    original = " ".join(f"term{i}" for i in range(120))
    revised = original + " updated approval procedure"
    for revision, text, status in [("first", original, 200), ("revised", revised, 409)]:
        response = client.post("/api/documents", files={"file": (
            f"{revision}.docx", docx_bytes(revision, text))})
        assert response.status_code == status, response.text
    assert response.json()["code"] == "similar_document"
    assert len(started) == 1


@pytest.mark.parametrize("extension", ["pdf", "xlsx"])
def test_other_supported_formats_check_text_before_admission(upload_client, extension):
    client, settings, started = upload_client

    def content(revision):
        out = BytesIO()
        if extension == "xlsx":
            from openpyxl import Workbook
            book = Workbook()
            book.properties.title = revision
            book.active["A1"] = TEXT
            book.save(out)
        else:
            from pypdf import PdfWriter
            from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
            writer = PdfWriter()
            page = writer.add_blank_page(width=600, height=800)
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                     NameObject("/Subtype"): NameObject("/Type1"),
                                     NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(b"BT /F1 12 Tf 50 750 Td (Procedure for approval and processing of documents in the company system.) Tj ET")
            page[NameObject("/Contents")] = writer._add_object(stream)
            writer.add_metadata({"/Title": revision})
            writer.write(out)
        return out.getvalue()

    for revision, status in [("first", 200), ("revised", 409)]:
        response = client.post("/api/documents", files={"file": (
            f"{revision}.{extension}", content(revision))})
        assert response.status_code == status, response.text
    assert response.json()["code"] == "similar_document"
    assert len(started) == 1

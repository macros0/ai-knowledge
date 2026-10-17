"""Use real canonical stores and Qdrant in memory; forbid provider calls."""
import hashlib
import pytest
from threading import BoundedSemaphore
from types import SimpleNamespace
from unittest.mock import Mock

from qdrant_client import QdrantClient, models as qm
from app.config import Settings
from app.db.models import Document, DocumentGeneration, DocumentGenerationState, DocumentProcessingReview, DocumentSource, OkfConcept
from app.db.session import session_scope
from app.services.vector_store import VectorStore
from app.services.chunk_store import replace_chunks
from app.services.concept_store import replace_concepts
from app.models.schemas import OkfDocument
from app.services.okf_generator import OKFGenerator
from app.services.processing_scope_worker import execute_scope_change
import app.services.processing_scope_jobs as jobs


@pytest.fixture
def source_environment(tmp_path, monkeypatch):
    settings = Settings(_env_file=None, data_dir=tmp_path, qdrant_collection="scope-only-test", embedding_dimensions=2)
    settings.okf_table_llm_classify = False
    settings.uploads_dir.mkdir(parents=True, exist_ok=True)
    doc_id, base_id = "1234567890abcdef", "c" * 32
    root_file = settings.uploads_dir / f"{doc_id}.docx"
    root_file.write_bytes(b"immutable root; deliberately not parsed")
    root_hash = hashlib.sha256(root_file.read_bytes()).hexdigest()
    nodes = [{"source_key": str(i + 1) * 64, "parent_key": "root", "source_id": f"root/{i}",
              "locator": f"part:{i}", "display_name": f"{i}.xlsx", "kind": "xlsx", "supported": True,
              "children_complete": True, "content_fingerprint": str(i + 4) * 64} for i in range(2)]
    decisions = [{"source_key": n["source_key"], "effective_included": True, "reason": "small"} for n in nodes]
    snapshot = {"schema_version": 1, "source_file_hash": root_hash, "inventory": {"nodes": nodes, "complete": True},
                "decisions": decisions, "selection_hash": "b" * 64}
    from app.services.processing_scope_cache import embedding_fingerprint
    snapshot["vector_space"] = {"kind":"proven","fingerprint":embedding_fingerprint(settings)}
    with session_scope() as s:
        s.add(Document(id=doc_id, filename="root.docx", file_hash=root_hash, status="done", total_chunks=2, processed_chunks=2))
        s.flush()
        s.add(DocumentSource(doc_id=doc_id, source_id="root"))
        for n in nodes:
            s.add(DocumentSource(doc_id=doc_id, source_id=n["source_id"], parent_source_id="root",
                ordinal=int(n["source_id"][-1]), metadata_json={"source_key": n["source_key"]},
                content_fingerprint=n["content_fingerprint"], container_locator=n["locator"]))
        s.flush()
        chunks = [{"chunk_index": i * 10, "source_id": n["source_id"], "content": f"unique body {i}"} for i, n in enumerate(nodes)]
        replace_chunks(s, doc_id, chunks)
        docs = [OkfDocument(filepath=f"concept-{i}.md", content=f"concept body {i}", markdown=f"concept body {i}",
                            metadata={"source_id": n["source_id"], "chunk_index": i * 10,
                                      "title": f"Concept {i}", "generated_at": "2026-10-01T00:00:00+00:00"}) for i, n in enumerate(nodes)]
        replace_concepts(s, doc_id, docs)
        s.add(DocumentGeneration(id=base_id, doc_id=doc_id, phase="active", processing_scope_json=snapshot))
        s.add(DocumentGenerationState(doc_id=doc_id, active_generation_id=base_id))
    store = VectorStore()
    store.settings = settings
    store.client = QdrantClient(":memory:")
    store.client.create_collection(collection_name=store.collection, vectors_config=qm.VectorParams(size=2, distance=qm.Distance.COSINE),
                                   sparse_vectors_config={"sparse": qm.SparseVectorParams()})
    store.index_concepts(doc_id, docs, [[1., 0.], [0., 1.]], generation_id=base_id)
    store.index_chunks(doc_id, "root.docx", [c["content"] for c in chunks], [], [[1., 0.], [0., 1.]],
                       generation_id=base_id, source_ids=[c["source_id"] for c in chunks], chunk_indices=[c["chunk_index"] for c in chunks])
    pipeline = SimpleNamespace(settings=settings, vector_store=store, okf_generator=OKFGenerator(),
                               embedder=Mock(side_effect=AssertionError("No embeddings expected")))
    monkeypatch.setattr("app.config.get_settings", lambda: settings)
    monkeypatch.setattr("app.services.pipeline.get_pipeline", lambda: pipeline)
    monkeypatch.setattr(jobs, "_dispatch", lambda *args: None)
    monkeypatch.setattr(jobs, "_fast_slots", BoundedSemaphore(1))
    return locals()


@pytest.mark.parametrize("legacy", [False, True])
@pytest.mark.parametrize("write_bundles",[False,True])
def test_exclude_restore_multiple_times_without_provider_calls(source_environment, monkeypatch, legacy,write_bundles):
    env=source_environment
    doc_id,nodes,snapshot,root_hash,decisions,pipeline=[env[k] for k in ("doc_id","nodes","snapshot","root_hash","decisions","pipeline")]
    env["settings"].okf_write_bundles=write_bundles
    with session_scope() as s:
        s.query(OkfConcept).filter_by(doc_id=doc_id,slug="concept-1").one().relations=["concept-0","external"]
    from app.services.vector_store import concept_point_id
    env["store"].client.set_payload(collection_name=env["store"].collection,
        points=[concept_point_id(doc_id,"concept-1",generation_id=env["base_id"])],
        payload={"relations":["concept-0","external"]})
    if legacy:
        with session_scope() as s:
            generation=s.get(DocumentGeneration,env["base_id"])
            generation.processing_scope_json={k:v for k,v in generation.processing_scope_json.items() if k!="vector_space"}
    for turn, included in enumerate((False, True, False, True)):
        jobs._fast_slots = BoundedSemaphore(1)
        with session_scope() as s:
            state = s.get(DocumentGenerationState, doc_id)
            active = state.active_generation_id
            old = s.get(DocumentProcessingReview, doc_id)
            if old:
                s.delete(old)
                s.flush()
            chosen = [{**d, "effective_included": included if i == 0 else True} for i, d in enumerate(decisions)]
            s.add(DocumentProcessingReview(doc_id=doc_id, id=f"{turn + 1:032x}", revision=1, state="draft",
                source_file_hash=root_hash, parser_version="test", assessor_version="test", policy_json={},
                inventory_json=snapshot["inventory"], selection_json=chosen, selection_hash=f"{turn + 5:064x}", base_generation_id=active))
        preview = jobs.create_change_preview(doc_id, f"{turn + 1:032x}", 1, active, "editor")
        assert preview["mode"] == ("reuse_only" if included else "fast_exclude")
        applied = jobs.apply_scope_change(doc_id, preview["plan_id"], f"{turn + 10:032x}", active, "editor", {preview["mode"]})
        from app.db.models import DocumentScopeChange
        with session_scope() as s:
            owner = s.get(DocumentScopeChange, preview["plan_id"]).owner_token
        execute_scope_change(doc_id, preview["plan_id"], owner)
        if not included:
            # An obsolete staging chunk may survive an interrupted old attempt.
            # A published exclusion must not fall back to that private text.
            from app.api import documents
            from app.api.errors import ApiError
            stale = env["tmp_path"] / "old-staging"
            stale.mkdir(exist_ok=True)
            (stale / "chunk_00.md").write_text("EXCLUDED_STAGING_CANARY", encoding="utf-8")
            with monkeypatch.context() as patch:
                patch.setattr(documents, "StagingStore", lambda *args:SimpleNamespace(exists=lambda:True, dir=stale))
                with pytest.raises(ApiError) as error:
                    documents.get_chunk(doc_id, 0)
                assert error.value.code == "chunk_source_excluded"
        with session_scope() as s:
            assert s.get(DocumentScopeChange, preview["plan_id"]).state == "published"
            assert s.query(OkfConcept).filter_by(doc_id=doc_id).count() == (2 if included else 1)
            assert s.query(OkfConcept).filter_by(doc_id=doc_id, slug="concept-1").one().chunk_index == 10
            relations=["concept-0","external"] if included else ["external"]
            assert s.query(OkfConcept).filter_by(doc_id=doc_id,slug="concept-1").one().relations==relations
            pid=concept_point_id(doc_id,"concept-1",generation_id=s.get(DocumentGenerationState,doc_id).active_generation_id)
            point=env["store"].client.retrieve(collection_name=env["store"].collection,ids=[pid])[0]
            assert point.payload["relations"]==relations
            if turn:
                restored_point=env["store"].client.retrieve(collection_name=env["store"].collection,
                    ids=[concept_point_id(doc_id,"concept-0",generation_id=s.get(DocumentGenerationState,doc_id).active_generation_id)]) if included else []
                assert not restored_point or restored_point[0].payload["source_locale"]=="en"
            active_generation=s.get(DocumentGeneration,s.get(DocumentGenerationState,doc_id).active_generation_id)
            from app.db.models import DocumentChunk
            from app.services.deduplication import content_hash
            ordered=s.query(DocumentChunk).filter_by(doc_id=doc_id).order_by(DocumentChunk.display_order).all()
            assert s.get(Document,doc_id).content_hash==content_hash("\n\n".join(c.content for c in ordered))
            from app.services.generation_files import generation_paths
            bundle=generation_paths(env["settings"],doc_id,active_generation.id).bundle
            assert bool(list(bundle.rglob("*.md")))==write_bundles
            assert active_generation.processing_scope_summary_json["chunk_counts"] == {
                "active":2 if included else 1,"archived":0 if included else 1,"stored":2}
        if turn==0:
            with session_scope() as s:
                s.get(Document,doc_id).source_locale="en"
                s.get(Document,doc_id).source_locale_source="manual"
            env["store"].set_document_source_locale_payload(doc_id,"en")
    pipeline.embedder.assert_not_called()
    from app.services.export_okf import export_okf_bundle
    import json
    exported=env["tmp_path"]/"exported"
    export_okf_bundle(doc_id,exported)
    assert len(json.loads((exported/"_files.json").read_text()))==2
    assert len(json.loads((exported/"chunks/manifest.json").read_text()))==2


@pytest.mark.parametrize("version", [None, 1])
@pytest.mark.parametrize("salvage", [False, True])
def test_reenable_with_cached_chunks_generates_only_missing_concepts(source_environment, monkeypatch, version, salvage):
    from app.services.processing_scope_cache import SourceResultCache, embedding_fingerprint
    from app.services.generation_files import generation_paths
    from app.services.processing_scope_change import load_scope_base
    from app.models.schemas import Concept
    from threading import BoundedSemaphore, RLock
    from concurrent.futures import ThreadPoolExecutor
    env = source_environment
    doc_id, pipeline, nodes, snapshot, root_hash = [env[k] for k in ("doc_id", "pipeline", "nodes", "snapshot", "root_hash")]
    with session_scope() as s:
        gen = s.get(DocumentGeneration, env["base_id"])
        scope = dict(gen.processing_scope_json)
        scope["decisions"] = [{**d, "effective_included": i != 0} for i, d in enumerate(scope["decisions"])]
        gen.processing_scope_json = scope
        s.query(OkfConcept).filter_by(doc_id=doc_id, source_id="root/0").delete()
        from app.db.models import DocumentChunk
        s.query(DocumentChunk).filter_by(doc_id=doc_id, source_id="root/0").delete()
        if version:
            from app.db.models import DocumentChunkProjection
            from app.services.chunk_index_projection import build_chunk_index_projection
            gen.chunk_projection_version = version
            for c in s.query(DocumentChunk).filter_by(doc_id=doc_id):
                c.content_hash = hashlib.sha256(c.content.encode()).hexdigest()
                p = build_chunk_index_projection(c.section_title or "", c.content, version=version)
                s.add(DocumentChunkProjection(doc_id=doc_id, generation_id=gen.id, chunk_index=c.chunk_index,
                    source_id=c.source_id, version=version, source_hash=p.source_hash, title_hash=p.title_hash,
                    projection_hash=p.projection_hash, fragment_count=0, complete=True))
        s.add(DocumentProcessingReview(doc_id=doc_id, id="e" * 32, revision=1, state="draft",
            source_file_hash=root_hash, parser_version="test", assessor_version="test", policy_json={},
            inventory_json=snapshot["inventory"], selection_json=snapshot["decisions"], selection_hash="f" * 64, base_generation_id=env["base_id"]))
    cache = SourceResultCache(generation_paths(env["settings"], doc_id, env["base_id"]).bundle / "source-cache", doc_id)
    chunks = [{"chunk_index": 0, "source_id": "root/0", "content": "cached original", "local_order": 0}]
    if salvage:
        chunks.append({"chunk_index": 1, "source_id": "root/0", "content": "completed sibling", "local_order": 1})
    cache.save_stage(nodes[0], "parse", {"chunks": chunks,
                                       "source": {"source_id": "root/0", "parent_source_id": "root", "ordinal": 0}})
    from app.services import gen_quality
    attempts = []
    def generate_concepts(content, *args, **kwargs):
        attempts.append(content)
        if salvage and attempts == ["cached original"]:
            gen_quality.record(gen_quality.LLM_SALVAGE)
        return [Concept(id="new", title="new concept", content="generated only once")]
    generate = Mock(side_effect=generate_concepts)
    monkeypatch.setattr(pipeline.okf_generator, "generate_chunk", generate)
    pipeline.embedder = SimpleNamespace(embed_texts=Mock(side_effect=lambda values: [[1., 0.] for _ in values]))
    pipeline._threads, pipeline._start_lock = {}, RLock()
    pipeline._pipeline_slots, pipeline._queue_limit = BoundedSemaphore(1), lambda admin: 0
    pipeline._executor = ThreadPoolExecutor(max_workers=1)
    monkeypatch.setattr("app.services.parser_supervisor.parse_document_supervised", Mock(side_effect=AssertionError("Already parsed")))
    preview = jobs.create_change_preview(doc_id, "e" * 32, 1, env["base_id"], "editor")
    assert preview["mode"] == "incremental_generate"
    applied = jobs.apply_scope_change(doc_id, preview["plan_id"], "1" * 32, env["base_id"], "editor", {"incremental_generate"})
    from app.db.models import DocumentScopeChange
    with session_scope() as s:
        operation = s.get(DocumentScopeChange, preview["plan_id"])
        owner = operation.owner_token
        operation.state = "processing"
    if salvage:
        with pytest.raises(ValueError, match="[Ii]ncomplete"):
            execute_scope_change(doc_id, preview["plan_id"], owner)
        with session_scope() as s:
            assert s.get(DocumentGenerationState, doc_id).active_generation_id == env["base_id"]
        stage_cache = SourceResultCache(env["settings"].uploads_dir / doc_id / "scope-stage-cache", doc_id)
        assert not stage_cache.inspect(nodes[0])["stages"]["concepts"]
        assert stage_cache.read_stage(nodes[0], "concepts", allow_partial=True)["completed_chunks"] == [1]
        pipeline.embedder.embed_texts.assert_not_called()
    execute_scope_change(doc_id, preview["plan_id"], owner)
    assert generate.call_count == (3 if salvage else 1)
    assert attempts == (["cached original", "completed sibling", "cached original"] if salvage else ["cached original"])
    assert generate.call_args.args[0] == "cached original"
    with session_scope() as s:
        assert s.query(OkfConcept).filter_by(doc_id=doc_id).count() == (3 if salvage else 2)
        old = s.query(OkfConcept).filter_by(doc_id=doc_id, slug="concept-1").one()
        assert old.chunk_index == 10 and old.content == "concept body 1"
    if version:
        from app.services.processing_scope_cache import inspect_source_cache
        descriptor = inspect_source_cache(doc_id, nodes[0])
        stage = SourceResultCache(descriptor["location"], doc_id).read_stage(nodes[0], "embeddings")
        assert stage.get("projections"), "A completed embedding stage must carry its SQL projection for recovery"
    pipeline._executor.shutdown()


def test_added_excel_uses_supervised_branch_and_keeps_existing_results(source_environment, monkeypatch):
    from copy import deepcopy
    from concurrent.futures import ThreadPoolExecutor
    from threading import RLock
    from openpyxl import Workbook
    from docparser import PARSER_VERSION
    from app.services.generation_files import processing_review_paths
    from app.db.models import DocumentScopeChange, DocumentChunk
    env = source_environment
    doc_id, pipeline = env["doc_id"], env["pipeline"]
    review_id = "e" * 32
    folder = processing_review_paths(env["settings"], doc_id, review_id) / "attachments"
    folder.mkdir(parents=True)
    book = Workbook()
    book.active.append(["Code", "Meaning"])
    book.active.append(["NEW_ONLY", "new row"])
    original = folder / "new.xlsx"
    book.save(original)
    node = {"source_key": "3" * 64, "parent_key": "root", "locator": "part:new", "display_name": "new.xlsx", "kind": "xlsx",
            "saved_path": "new.xlsx", "content_fingerprint": hashlib.sha256(original.read_bytes()).hexdigest(),
            "supported": True, "children_complete": True}
    inventory = deepcopy(env["snapshot"]["inventory"])
    inventory["nodes"].append(node)
    decisions = [*env["decisions"], {"source_key": node["source_key"], "effective_included": True, "reason": "small"}]
    with session_scope() as s:
        s.add(DocumentProcessingReview(doc_id=doc_id, id=review_id, revision=1, state="draft", source_file_hash=env["root_hash"],
            parser_version=PARSER_VERSION + "-scope-v1", assessor_version="test", policy_json={}, inventory_json=inventory,
            selection_json=decisions, selection_hash="f" * 64, base_generation_id=env["base_id"]))
    pipeline.embedder = SimpleNamespace(embed_texts=Mock(side_effect=lambda values: [[1., 0.] for _ in values]))
    pipeline._threads, pipeline._start_lock = {}, RLock()
    pipeline._pipeline_slots, pipeline._queue_limit = BoundedSemaphore(1), lambda admin: 0
    pipeline._executor = ThreadPoolExecutor(max_workers=1)
    monkeypatch.setattr(pipeline.okf_generator, "generate_chunk", Mock(side_effect=AssertionError("Table-only source is deterministic")))
    preview = jobs.create_change_preview(doc_id, review_id, 1, env["base_id"], "editor")
    assert preview["mode"] == "incremental_generate"
    result = jobs.apply_scope_change(doc_id, preview["plan_id"], "1" * 32, env["base_id"], "editor", {"incremental_generate"})
    with session_scope() as s:
        op = s.get(DocumentScopeChange, preview["plan_id"])
        op.state = "processing"
        owner = op.owner_token
    execute_scope_change(doc_id, preview["plan_id"], owner)
    with session_scope() as s:
        assert s.query(OkfConcept).filter_by(doc_id=doc_id, slug="concept-0").one().content == "concept body 0"
        assert s.query(OkfConcept).filter_by(doc_id=doc_id, slug="concept-1").one().chunk_index == 10
        new_chunks = s.query(DocumentChunk).filter(DocumentChunk.doc_id == doc_id, DocumentChunk.source_id.notin_(["root/0", "root/1"])).all()
        assert new_chunks and all(c.chunk_index > 10 for c in new_chunks)
        assert s.get(DocumentScopeChange, preview["plan_id"]).state == "published"
        from app.db.models import DocumentSource
        new_source = s.get(DocumentSource, (doc_id, "root/s333333333333"))
        assert new_source.saved_path, "Newly parsed originals must outlive the review and staging cache"
        registered = env["settings"].uploads_dir / doc_id / new_source.saved_path
        assert hashlib.sha256(registered.read_bytes()).hexdigest() == node["content_fingerprint"]
        assert s.get(DocumentGeneration, result["candidate_generation_id"]).table_schema_version == 1
    import json
    from app.services.generation_files import generation_paths
    quality = generation_paths(env["settings"], doc_id, result["candidate_generation_id"]).bundle / "table-quality.json"
    assert quality.is_file(), "New table extraction must retain its coverage diagnostics"
    assert any(r["table_ref"]["source_id"] == new_source.source_id for r in json.loads(quality.read_text())["reports"])
    pipeline._executor.shutdown()


def test_excluding_highest_chunk_keeps_reserved_index_after_publication(source_environment):
    env = source_environment
    doc_id = env["doc_id"]
    with session_scope() as s:
        s.add(DocumentProcessingReview(doc_id=doc_id, id="e" * 32, revision=1, state="draft",
            source_file_hash=env["root_hash"], parser_version="test", assessor_version="test", policy_json={},
            inventory_json=env["snapshot"]["inventory"],
            selection_json=[{**d, "effective_included": i != 1} for i, d in enumerate(env["decisions"])],
            selection_hash="f" * 64, base_generation_id=env["base_id"]))
    preview = jobs.create_change_preview(doc_id, "e" * 32, 1, env["base_id"], "editor")
    jobs.apply_scope_change(doc_id, preview["plan_id"], "1" * 32, env["base_id"], "editor", {"fast_exclude"})
    from app.db.models import DocumentScopeChange
    with session_scope() as s:
        owner = s.get(DocumentScopeChange, preview["plan_id"]).owner_token
    execute_scope_change(doc_id, preview["plan_id"], owner)
    with session_scope() as s:
        assert s.get(DocumentGenerationState, doc_id).next_chunk_index >= 11


def test_original_changed_during_copy_cannot_publish(source_environment, monkeypatch):
    env = source_environment
    with session_scope() as s:
        s.add(DocumentProcessingReview(doc_id=env["doc_id"], id="e" * 32, revision=1, state="draft",
            source_file_hash=env["root_hash"], parser_version="test", assessor_version="test", policy_json={},
            inventory_json=env["snapshot"]["inventory"],
            selection_json=[{**d,"effective_included":i != 0} for i,d in enumerate(env["decisions"])],
            selection_hash="f" * 64, base_generation_id=env["base_id"]))
    p = jobs.create_change_preview(env["doc_id"], "e" * 32, 1, env["base_id"], "editor")
    jobs.apply_scope_change(env["doc_id"], p["plan_id"], "1" * 32, env["base_id"], "editor", {"fast_exclude"})
    with session_scope() as s:
        from app.db.models import DocumentScopeChange
        owner = s.get(DocumentScopeChange, p["plan_id"]).owner_token
    copy = env["store"].copy_generation_points
    def mutate_original(*args, **kwargs):
        result = copy(*args, **kwargs)
        env["root_file"].write_bytes(b"changed outside the application")
        return result
    monkeypatch.setattr(env["store"], "copy_generation_points", mutate_original)
    from app.services.generation_store import GenerationConflict
    with pytest.raises(GenerationConflict, match="[Oo]riginal"):
        execute_scope_change(env["doc_id"], p["plan_id"], owner)
    with session_scope() as s:
        assert s.get(DocumentGenerationState, env["doc_id"]).active_generation_id == env["base_id"]


def test_registered_source_original_cannot_change_before_copy(source_environment):
    from app.services.canonical_repair import _copy_artifacts
    from app.services.generation_files import generation_paths, GenerationStorageError
    env = source_environment
    paths = generation_paths(env["settings"], env["doc_id"], "f" * 32)
    original = env["settings"].uploads_dir / env["doc_id"] / f'generations/{env["base_id"]}/attachments/source.xlsx'
    original.parent.mkdir(parents=True)
    original.write_bytes(b"changed original")
    proposal = {"sources": [{"saved_path":f'generations/{env["base_id"]}/attachments/source.xlsx',
                             "content_fingerprint":hashlib.sha256(b"expected bytes").hexdigest()}], "attachments":[]}
    with pytest.raises(GenerationStorageError, match="hash changed"):
        _copy_artifacts(env["settings"], env["doc_id"], env["base_id"], "f" * 32, paths, proposal)


def test_excluding_previously_salvaged_source_keeps_cache_incomplete(source_environment):
    from app.services.processing_scope_cache import inspect_source_cache
    env=source_environment
    with session_scope() as s:
        s.get(Document,env["doc_id"]).problem="llm_partial_result"
        s.add(DocumentProcessingReview(doc_id=env["doc_id"],id="e"*32,revision=1,state="draft",
            source_file_hash=env["root_hash"],parser_version="test",assessor_version="test",policy_json={},
            inventory_json=env["snapshot"]["inventory"],
            selection_json=[{**d,"effective_included":i!=0} for i,d in enumerate(env["decisions"])],
            selection_hash="f"*64,base_generation_id=env["base_id"]))
    p=jobs.create_change_preview(env["doc_id"],"e"*32,1,env["base_id"],"editor")
    jobs.apply_scope_change(env["doc_id"],p["plan_id"],"1"*32,env["base_id"],"editor",{"fast_exclude"})
    with session_scope() as s:
        from app.db.models import DocumentScopeChange
        owner=s.get(DocumentScopeChange,p["plan_id"]).owner_token
    execute_scope_change(env["doc_id"],p["plan_id"],owner)
    hit=inspect_source_cache(env["doc_id"],env["nodes"][0])
    assert hit["verified"] and not hit["stages"]["concepts"]
    assert not hit["stages"]["embeddings"]
    with session_scope() as s:
        assert s.get(Document,env["doc_id"]).problem=="llm_partial_result"


@pytest.mark.parametrize("legacy",[False,True])
def test_preview_finds_canceled_new_source_parse_in_staging(source_environment,legacy):
    from copy import deepcopy
    from app.services.processing_scope_cache import SourceResultCache
    env = source_environment
    node = {"source_key": "3" * 64, "parent_key": "root", "locator": "part:new", "display_name": "new.docx",
            "kind": "docx", "content_fingerprint": "6" * 64, "supported": True, "children_complete": True}
    inventory = deepcopy(env["snapshot"]["inventory"])
    inventory["nodes"].append(node)
    with session_scope() as s:
        if legacy:
            s.get(DocumentGeneration,env["base_id"]).processing_scope_json=None
        s.add(DocumentProcessingReview(doc_id=env["doc_id"], id="e" * 32, revision=1, state="draft",
            source_file_hash=env["root_hash"], parser_version="test", assessor_version="test", policy_json={},
            inventory_json=inventory, selection_json=[*env["decisions"], {"source_key":node["source_key"],"effective_included":True}],
            selection_hash="f" * 64, base_generation_id=env["base_id"]))
    cache = SourceResultCache(env["settings"].uploads_dir / env["doc_id"] / "scope-stage-cache", env["doc_id"])
    cache.save_stage(node, "parse", {"chunks": [{"chunk_index": 20, "source_id":"root/s333333333333", "content":"completed parse"}]})
    p = jobs.create_change_preview(env["doc_id"], "e" * 32, 1, env["base_id"], "editor")
    assert p["workset"]["sources"][0]["parse"] == "reuse"

"""Synthetic acceptance inside a home/CI backend container; never writes PROD."""
from io import BytesIO
import json
from pathlib import Path
import sys
import time

import httpx

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.config import get_settings


def main():
    settings = get_settings()
    if settings.deployment_tier not in {"home", "ci"} or settings.model_mode != "stub" or settings.auth_provider != "disabled":
        raise RuntimeError("Synthetic smoke requires an isolated home/CI stub with disabled auth")
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from backup_totals import collect_totals
    from seed_backup_fixture import DOC_ID, ORIGINAL, ATTACHMENT

    with httpx.Client(base_url="http://127.0.0.1:8000", timeout=90) as client:
        for path in ("/health/live", "/health/ready"):
            response = client.get(path)
            response.raise_for_status()
        totals = collect_totals()
        assert totals["documents"] >= 1 and totals["qdrant_points"] >= 2
        assert (settings.uploads_dir / f"{DOC_ID}.txt").read_bytes() == ORIGINAL
        assert (settings.uploads_dir / DOC_ID / "attachments/fixture.txt").read_bytes() == ATTACHMENT
        response = client.post("/api/search", json={"query": "Backup fixture", "mode": "bm25", "use_glossary": False})
        response.raise_for_status()
        assert any(hit["source_filename"] == "fixture.txt" for hit in response.json()["hits"])
        print("PASS readiness, persistent originals/attachments, BM25 and nonempty corpus")

        from docx import Document
        document = Document()
        document.add_heading("Kubernetes synthetic fixture", level=1)
        document.add_paragraph("TEST Kubernetes knowledge fixture. " * 30)
        output = BytesIO()
        document.save(output)
        uploaded = client.post("/api/documents", files={"file": ("kubernetes-smoke.docx", output.getvalue(),
                                   "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, data={"allow_similar": "true"})
        uploaded.raise_for_status()
        identifier = uploaded.json()["id"]
        deadline = time.monotonic() + 180
        while True:
            state = client.get(f"/api/documents/{identifier}")
            state.raise_for_status()
            result = state.json()
            if result["status"] == "done":
                assert not result["problem"]
                assert result["okf_concept_count"] >= 1
                break
            if result["status"] in {"failed", "paused", "error"} or time.monotonic() > deadline:
                raise RuntimeError("Synthetic generation did not complete")
            time.sleep(2)
        sources = client.get(f"/api/documents/{identifier}/sources")
        sources.raise_for_status()
        assert sources.json()["sources"]
        print("PASS DOCX upload, stub generation, indexing and source inventory")

    # Check the real standalone frontend build's frozen backend endpoint and streaming proxy.
    with httpx.Client(base_url="http://frontend:3000", timeout=90) as frontend:
        frontend.get("/health").raise_for_status()
        query = {"query": "Backup fixture", "mode": "bm25", "use_glossary": False}
        with frontend.stream("POST", "/api/chat/stream", json=query) as response:
            response.raise_for_status()
            events = [json.loads(line) for line in response.iter_lines() if line.strip()]
        assert any(event["type"] == "result" for event in events), "Chat stream did not finish"
        assert not any(event["type"] == "error" for event in events)
        print("PASS frontend health rewrite and NDJSON chat proxy with upstream stub SSE")


if __name__ == "__main__":
    main()

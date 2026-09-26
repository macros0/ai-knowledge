"""Create a fixed, synthetic 100-input acceptance corpus (no user documents)."""
import argparse
from datetime import datetime, timezone
from email.message import EmailMessage
from io import BytesIO
import hashlib
import importlib.util
import json
from pathlib import Path
import random
import zipfile

ROOT = Path(__file__).resolve().parents[2]


def canonical_zip(data):
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            entry = zipfile.ZipInfo(name, (2026, 9, 25, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(entry, source.read(name))
    return output.getvalue()


def mail(index, body):
    msg = EmailMessage()
    msg["From"] = f"sender-{index:03d}@example.test"
    msg["To"] = "recipient@example.test"
    msg["Subject"] = f"Проверка LOAD-{index:03d}"
    msg["Date"] = "Fri, 25 Sep 2026 12:00:00 +0300"
    msg["Message-ID"] = f"<load-{index:03d}@example.test>"
    msg.set_content(body)
    return msg


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    task = parser.parse_args().output.resolve()
    from docx import Document
    from docx.shared import Inches
    from PIL import Image
    from reportlab.pdfgen.canvas import Canvas
    spec = importlib.util.spec_from_file_location("load_fixture_helpers", ROOT / "doc-parser/tests/fixtures.py")
    helpers = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(helpers)
    corpus = task / "corpus"
    assert not corpus.exists(), "Corpus is immutable; reuse its manifest, never silently replace"
    corpus.mkdir(parents=True)
    image_data = BytesIO()
    Image.frombytes("RGB", (2048, 1024), random.Random(420026).randbytes(2048 * 1024 * 3)).save(image_data, format="PNG")
    rows = []
    for index in range(100):
        fact = f"Срок обработки заявки LOAD-{index:03d} составляет {12 + index} календарных дней."
        if index < 70:
            category = "ordinary_mail"
            payload = mail(index, fact + "\n\n" + "Для проверки требуется полный комплект документов, согласование заявки и подпись ответственного исполнителя. " * 8).as_bytes()
            path = corpus / f"{index:03d}.eml"
            expected_sources = 1
        elif index < 80:
            category = "attachment_only_mail"
            doc = Document()
            doc.core_properties.created = doc.core_properties.modified = datetime(2026, 9, 25, tzinfo=timezone.utc)
            doc.add_heading(f"Порядок LOAD-{index:03d}", level=1)
            doc.add_paragraph(fact)
            doc.add_picture(BytesIO(image_data.getvalue()), width=Inches(4))
            buffer = BytesIO()
            doc.save(buffer)
            msg = mail(index, "")
            msg.add_attachment(canonical_zip(buffer.getvalue()), maintype="application",
                               subtype="vnd.openxmlformats-officedocument.wordprocessingml.document", filename="instruction.docx")
            msg.set_boundary(f"load-fixed-{index:03d}")
            payload = msg.as_bytes()
            path = corpus / f"{index:03d}.eml"
            expected_sources = 2
        else:
            category = "document_with_embedded_mail"
            path = corpus / f"{index:03d}.docx"
            helpers.make_docx_with_embedded_xlsx(path, mail(index, fact).as_bytes(),
                prog_id="Outlook.FileMsg.15", filename=f"decision-{index:03d}.eml", in_table=bool(index % 2))
            payload = canonical_zip(path.read_bytes())
            expected_sources = 2
        path.write_bytes(payload)
        rows.append({"name": path.name, "category": category, "size": len(payload),
                     "sha256": hashlib.sha256(payload).hexdigest(), "expected_sources": expected_sources,
                     "required_fact": fact})
    controls = []
    for index in range(10):
        doc = Document()
        doc.core_properties.created = doc.core_properties.modified = datetime(2026, 9, 25, tzinfo=timezone.utc)
        for paragraph in range(20):
            doc.add_paragraph(f"Control document {index:03d}, paragraph {paragraph:03d}: certificate verification and processing requirements.")
        buffer = BytesIO()
        doc.save(buffer)
        path = corpus / f"control-{index:03d}.docx"
        path.write_bytes(canonical_zip(buffer.getvalue()))
        controls.append(path)
        pdf = corpus / f"control-{index:03d}.pdf"
        canvas = Canvas(str(pdf), invariant=1)
        for paragraph in range(20):
            canvas.drawString(40, 780 - 25 * paragraph, f"Control PDF {index:03d}, paragraph {paragraph:03d}: processing requirements.")
        canvas.save()
        controls.append(pdf)
    total = sum(row["size"] for row in rows)
    assert 80 * 1024**2 < total < 100 * 1024**2
    manifest = {"kind": "synthetic_mail_load_v1", "count": 100, "input_bytes": total,
                "limits": {"ordinary_mail_parse_p95_seconds": 5, "worker_memory_mib": 1024,
                           "worker_timeout_seconds": 120, "ordinary_docx_pdf_p95_ratio": 1.2,
                           "search_p95_ratio_during_load": 1.2},
                "entries": rows, "controls": [{"name": path.name, "size": path.stat().st_size,
                    "sha256": hashlib.sha256(path.read_bytes()).hexdigest()} for path in controls],
                "repeat_uploads": [rows[index]["name"] for index in (0, 35, 69, 70, 79, 80, 99)]}
    (task / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"count": 100, "bytes": total, "mib": round(total / 1024**2, 3),
                      "ordinary_mail": 70, "attachment_only_mail": 10, "embedded_mail_documents": 20,
                      "control_files": len(controls), "manifest": str(task / "manifest.json")}))


if __name__ == "__main__":
    main()

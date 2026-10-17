"""Header decisions are checked against independent, original coordinates."""

import pytest

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from docparser import parse_document_result
from tests.test_xlsx_structure import RecordingSink


def parse(tmp_path, configure):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    configure(ws)
    path = tmp_path / "headers.xlsx"
    wb.save(path)
    sink = RecordingSink()
    parse_document_result(path, table_sink=sink)
    return list(sink.tables.values())


def test_multilevel_header_preserves_paths_and_addresses(tmp_path):
    def configure(ws):
        ws.append(["Annual report"])
        ws.merge_cells("A1:D1")
        ws.append(["Code", "Personnel", None, "Totals"])
        ws.merge_cells("A2:A3")
        ws.merge_cells("B2:C2")
        ws.append([None, "Name", "Count", "Count"])
        ws.append(["K", "Alice", 0, 1])

    table = parse(tmp_path, configure)[0]
    d = table["descriptor"]
    assert d.header_rows == [2, 3] and d.method == "structural"
    assert [c.header_path for c in d.columns] == [
        ["Code"],
        ["Personnel", "Name"],
        ["Personnel", "Count"],
        ["Totals", "Count"],
    ]
    assert d.columns[2].header_addresses == ["B2", "C3"]
    assert table["rows"][0].cells[1].raw_value == "Annual report"


def test_empty_a1_and_blank_separator_keep_original_indices(tmp_path):
    def configure(ws):
        ws.append([None, "Code", None, "Quantity"])
        ws.append([None, "K", None, 0])
        ws.merge_cells("B1:D1")  # group spans the blank separator

    table = parse(tmp_path, configure)[0]
    assert [c.index for c in table["descriptor"].columns] == [2, 3, 4]
    assert table["rows"][1].cells[4].raw_value == 0


def test_explicit_tables_take_priority_and_keep_undeclared_cells(tmp_path):
    def configure(ws):
        ws.append(["Code", "Value", None, "Other", "Count"])
        ws.append(["K", 7, None, "X", 8])
        ws.append(["K", 7, None, "X", 8])
        ws.add_table(Table(displayName="Left", ref="A1:B3"))
        ws.add_table(Table(displayName="Right", ref="D1:E3"))
        ws["G9"] = "outside"

    tables = parse(tmp_path, configure)
    assert len(tables) == 3
    assert [t["descriptor"].method for t in tables[:2]] == ["explicit", "explicit"]
    assert tables[0]["descriptor"].header_rows == [1]
    assert any(
        c.address == "G9" and c.raw_value == "outside"
        for t in tables
        for r in t["rows"]
        for c in r.cells.values()
    )


def test_ambiguous_header_keeps_all_rows(tmp_path):
    tables = parse(
        tmp_path, lambda ws: (ws.append(["alpha", "beta"]), ws.append(["gamma", "delta"]))
    )
    d = tables[0]["descriptor"]
    assert d.method == "unresolved" and d.header_rows == []
    assert "header_ambiguous" in d.warnings
    assert [c.header_path for c in d.columns] == [["A"], ["B"]]
    assert [r.row_number for r in tables[0]["rows"]] == [1, 2]


def test_duplicates_and_repeated_headers_keep_physical_identity(tmp_path):
    def configure(ws):
        ws.append(["Code", "Value", "Value"])
        ws.append(["K", 0, False])
        ws.append(["Code", "Value", "Value"])
        ws.append(["K", 0, False])

    table = parse(tmp_path, configure)[0]
    assert [c.letter for c in table["descriptor"].columns] == ["A", "B", "C"]
    assert table["descriptor"].header_rows == [1]
    assert len(table["rows"]) == 4
    assert table["rows"][2].cells[2].raw_value == "Value"


def test_vertical_data_merge_keeps_anchor_without_copying_value(tmp_path):
    def configure(ws):
        ws.append(["Code", "Count"])
        ws.append(["K", 1])
        ws.append([None, 2])
        ws.merge_cells("A2:A3")

    table = parse(tmp_path, configure)[0]
    assert table["rows"][1].cells[1].merge_anchor == "A2"
    other = table["rows"][2].cells.get(1)
    assert other is None or other.raw_value is None


def test_blank_row_alone_does_not_split_table(tmp_path):
    def configure(ws):
        ws.append(["Code", "Count"])
        ws.append(["K", 1])
        ws.append([])
        ws.append(["X", 2])

    tables = parse(tmp_path, configure)
    assert len(tables) == 1 and tables[0]["descriptor"].cell_range.max_row == 4


def test_declared_range_never_duplicates_surrounding_cells(tmp_path):
    def configure(ws):
        ws.append(["Report", "Extra"])
        ws.append(["Code", "Count"])
        ws.append(["K", 7])
        ws.append(["Footer", 8])
        ws.add_table(Table(displayName="Middle", ref="A2:B3"))

    tables = parse(tmp_path, configure)
    covered = [
        c.address
        for t in tables
        for row in t["rows"]
        for c in row.cells.values()
        if c.raw_value != ""
    ]
    assert len(covered) == len(set(covered)) == 8


def test_three_header_levels_keep_rightmost_label(tmp_path):
    def configure(ws):
        ws.append(["Code", "Group"])
        ws.merge_cells("A1:A3")
        ws.merge_cells("B1:D1")
        ws.append([None, "Subgroup", None, "Total"])
        ws.merge_cells("B2:C2")
        ws.merge_cells("D2:D3")
        ws.append([None, "Name", "Count"])
        ws.append(["K", "Alice", 0, 1])

    d = parse(tmp_path, configure)[0]["descriptor"]
    assert d.header_rows == [1, 2, 3]
    assert d.columns[2].header_path == ["Group", "Subgroup", "Count"]


def test_horizontal_merges_across_three_levels_keep_leaf_names(tmp_path):
    def configure(ws):
        ws.append(["Group A", None, "Group B"])
        ws.merge_cells("A1:B1")
        ws.merge_cells("C1:D1")
        ws.append(["Section A", None, "Section B"])
        ws.merge_cells("A2:B2")
        ws.merge_cells("C2:D2")
        ws.append(["Code", "Quantity", "Code", "Price"])
        ws.append(["K", 2, "Y", 4])

    d = parse(tmp_path, configure)[0]["descriptor"]
    assert d.header_rows == [1, 2, 3]
    assert d.columns[3].header_path == ["Group B", "Section B", "Price"]


def test_sparse_preamble_is_not_a_level_above_numbered_headers(tmp_path):
    from openpyxl.styles import Font

    def configure(ws):
        ws.append(["Instruction", None, None, "Context"])
        ws.merge_cells("A1:C1")
        ws["A1"].font = Font(bold=True)
        ws.append(["Code", "Name", "Measure", "Group", None, "Other", "Flags", None, "Notes"])
        for address in ["A2:A3", "B2:B3", "C2:C3", "D2:E2", "F2:F3", "G2:H2", "I2:I3"]:
            ws.merge_cells(address)
        ws.append([None, None, None, "Left", "Right", None, "Active", "Count"])
        ws.append(list(range(1, 10)))
        ws.append([1, "Item", 0, 1, 2, 3, False, 4, "Text"])

    descriptor = parse(tmp_path, configure)[0]["descriptor"]
    assert descriptor.header_rows == [2, 3]
    assert descriptor.columns[3].header_path == ["Group", "Left"]
    assert "Instruction" not in descriptor.columns[0].header_path


def test_separate_header_after_gap_starts_a_new_table(tmp_path):
    from openpyxl.styles import Font

    def configure(ws):
        for r, labels in [
            (1, ["Code", "Name", "Quantity"]),
            (6, ["Code", "Name", "Location"]),
            (11, ["Code", "Name", "Location"]),
        ]:
            for c, v in enumerate(labels, 1):
                ws.cell(r, c, v).font = Font(bold=True)
            ws.cell(r + 1, 1, 100 + r)
            ws.cell(r + 1, 2, "Item")
            ws.cell(r + 1, 3, 0)
        ws.cell(10, 1, "Third section")

    tables = parse(tmp_path, configure)
    assert len(tables) == 3
    assert [t["descriptor"].header_rows for t in tables] == [[1], [6], [11]]
    addresses = [
        cell.address
        for t in tables
        for row in t["rows"]
        for cell in row.cells.values()
        if cell.raw_value not in (None, "")
    ]
    assert len(addresses) == len(set(addresses)) == 19


def test_stacked_merged_sections_keep_their_own_labels(tmp_path):
    from openpyxl.styles import Font

    def configure(ws):
        for r, title, labels in [
            (1, "Section A", ["Code", "Name", "Value"]),
            (7, "Section B", ["Code", "Name", "Other"]),
        ]:
            ws.cell(r, 1, title).font = Font(bold=True)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            for c, v in enumerate(labels, 1):
                ws.cell(r + 1, c, v)
            ws.cell(r + 2, 1, 100 + r)
            ws.cell(r + 2, 2, "Item")
            ws.cell(r + 2, 3, 0)

    tables = parse(tmp_path, configure)
    assert len(tables) == 2
    assert [t["descriptor"].header_rows for t in tables] == [[1, 2], [7, 8]]
    assert tables[1]["descriptor"].columns[2].header_path == ["Section B", "Other"]


def test_preamble_and_group_do_not_promote_a_merged_data_row(tmp_path):
    from openpyxl.styles import Font

    def configure(ws):
        ws.append(["Report note"])
        ws.merge_cells("A1:F1")
        ws["A1"].font = Font(bold=True)
        ws.append(["Group"])
        ws.merge_cells("A2:F2")
        ws["A2"].font = Font(bold=True)
        ws.append(["Code", "Previous", None, "New"])
        ws.merge_cells("B3:C3")
        ws.merge_cells("D3:F3")
        ws.append([101, "Previous name", None, "New name"])
        ws.merge_cells("B4:C4")
        ws.merge_cells("D4:F4")

    d = parse(tmp_path, configure)[0]["descriptor"]
    assert d.header_rows == [2, 3]
    assert d.columns[3].header_path == ["Group", "New"]


def test_bold_text_data_after_gap_is_not_a_new_header(tmp_path):
    from openpyxl.styles import Font

    def configure(ws):
        ws.append(["Code", "Description"])
        ws.append([101, "First"])
        ws.cell(5, 1, "K").font = Font(bold=True)
        ws.cell(5, 2, "Second").font = Font(bold=True)
        ws.cell(6, 1, "L")
        ws.cell(6, 2, "Third")

    tables = parse(tmp_path, configure)
    assert len(tables) == 1
    assert tables[0]["descriptor"].header_rows == [1]


def test_full_width_section_after_prose_splits_without_empty_row(tmp_path):
    def configure(ws):
        ws.append(["Initial note"])
        ws.merge_cells("A1:C1")
        ws.append(["Code", "Previous", None])
        ws.merge_cells("B2:C2")
        ws.append([None, "Short", "Full"])
        ws.merge_cells("A2:A3")
        for row in range(4, 10):
            ws.append([row, "Value", "Description"])
        ws.append(["Explanation"])
        ws.append(["Next section"])
        ws.merge_cells("A11:E11")
        ws.append(["Code", "Name", "Count", "Unit", "Other"])
        ws.append([201, "Item", 0, "kg", False])

    tables = parse(tmp_path, configure)
    assert len(tables) == 2
    assert [t["descriptor"].header_rows for t in tables] == [[2, 3], [12]]
    assert tables[0]["descriptor"].cell_range.max_column == 3
    assert tables[1]["descriptor"].cell_range.min_row == 11


def test_column_numbering_bridges_blank_column_under_shared_header(tmp_path):
    def configure(ws):
        ws.append(["Instruction"])
        ws.merge_cells("A1:B1")
        for col in [1, 2, 3, 5, 6, 7]:
            ws.cell(2, col, "Field " + str(col))
            ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
        for number, col in enumerate([1, 2, 3, 5, 6, 7], 1):
            ws.cell(4, col, number)
            ws.cell(5, col, number * 10)

    tables = parse(tmp_path, configure)
    assert len(tables) == 1
    descriptor = tables[0]["descriptor"]
    assert descriptor.header_rows == [2, 3]
    assert [c.index for c in descriptor.columns] == list(range(1, 8))
    assert descriptor.columns[4].header_path == ["Field 5"]
    assert tables[0]["rows"][3].cells[5].raw_value == 4


def test_numbering_does_not_combine_explicit_tables(tmp_path):
    def configure(ws):
        ws.append(["Left", "Value", None, "Right", "Value"])
        ws.append([1, 2, None, 3, 4])
        ws.append([10, 20, None, 30, 40])
        ws.add_table(Table(displayName="Left", ref="A1:B3"))
        ws.add_table(Table(displayName="Right", ref="D1:E3"))

    tables = parse(tmp_path, configure)
    assert len(tables) == 2
    assert all(t["descriptor"].method == "explicit" for t in tables)


def test_section_caption_can_be_narrower_than_next_leaf_level(tmp_path):
    from openpyxl.styles import Font

    def configure(ws):
        for row, title in [(1, "First"), (7, "Second")]:
            ws.cell(row, 1, title).font = Font(bold=True)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            for col, label in enumerate(
                ["Code", "Name", "Value"] + (["Extra"] if row == 7 else []), 1
            ):
                ws.cell(row + 1, col, label)
            for col, value in enumerate(["K", "Item", "Yes"], 1):
                ws.cell(row + 2, col, value)

    tables = parse(tmp_path, configure)
    assert len(tables) == 2
    assert [t["descriptor"].header_rows for t in tables] == [[1, 2], [7, 8]]
    assert tables[1]["descriptor"].columns[3].header_path == ["Extra"]


def test_numeric_year_labels_below_caption_are_headers(tmp_path):
    def configure(ws):
        ws.append(["Code", "Count"])
        ws.append([101, 0])
        for row in range(3, 11):
            ws.cell(row, 1, "Explanation")
        ws.cell(11, 1, "Next section")
        ws.merge_cells("A11:D11")
        ws.append(["Code", "Name", 2025, 2026])
        ws.append([201, "Item", "Yes", "No"])

    tables = parse(tmp_path, configure)
    assert len(tables) == 2
    assert tables[1]["descriptor"].header_rows == [12]
    assert tables[1]["descriptor"].columns[3].header_path == ["2026"]


def test_merged_string_data_does_not_become_a_header_chain(tmp_path):
    def configure(ws):
        for row, values in enumerate(
            [["Code", "Description"], ["K1", "Alice"], ["K2", "Bob"], ["K3", "Carol"]], 1
        ):
            for col, value in enumerate(values, 1):
                ws.cell(row, col, value)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)

    tables = parse(tmp_path, configure)
    assert tables[0]["descriptor"].header_rows == []
    assert tables[0]["descriptor"].method == "unresolved"
    assert len(tables[0]["rows"]) == 4


def test_vertical_span_past_sample_is_ambiguous(tmp_path):
    def configure(ws):
        ws.append(["Category", "Name"])
        ws.append([None, "Alice"])
        ws.merge_cells("A1:A100000")

    descriptor = parse(tmp_path, configure)[0]["descriptor"]
    assert descriptor.header_rows == []
    assert descriptor.method == "unresolved"


def test_bold_style_of_neighbor_table_cannot_split_current_region(tmp_path):
    from openpyxl.styles import Font

    def configure(ws):
        for row, left in [
            (1, ["Code", "Quantity"]),
            (2, ["K1", 2]),
            (5, ["K2", "Other"]),
            (6, ["K3", 3]),
        ]:
            for col, value in enumerate(left, 1):
                ws.cell(row, col, value)
            ws.cell(row, 4, "Neighbor")
            ws.cell(row, 5, row)
        ws["D5"].font = Font(bold=True)

    tables = parse(tmp_path, configure)
    left = [t for t in tables if t["descriptor"].cell_range.min_column == 1]
    assert len(left) == 1
    assert left[0]["descriptor"].header_rows == [1]


def test_late_numeric_data_does_not_validate_repeated_merge_headers(tmp_path):
    def configure(ws):
        for row, values in enumerate(
            [["Code", "Description"], ["K1", "Alice"], ["K2", "Bob"], ["K3", 5]], 1
        ):
            for col, value in enumerate(values, 1):
                ws.cell(row, col, value)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)

    descriptor = parse(tmp_path, configure)[0]["descriptor"]
    assert descriptor.header_rows == []
    assert descriptor.method == "unresolved"


def test_vertical_string_data_before_numeric_row_is_ambiguous(tmp_path):
    def configure(ws):
        ws.append(["Category", "Name"])
        ws.append([None, "Alice"])
        ws.append([None, "Bob"])
        ws.append(["Next", 5])
        ws.merge_cells("A1:A3")

    descriptor = parse(tmp_path, configure)[0]["descriptor"]
    assert descriptor.header_rows == []
    assert descriptor.method == "unresolved"


def test_numeric_first_leaf_above_numbering_keeps_multilevel_header(tmp_path):
    def configure(ws):
        ws['C1'] = 199  # auxiliary note, not a header
        ws.append(['Code', 'Rates', None, 'Accounting', None, None])
        ws.merge_cells('A2:A3')
        ws.merge_cells('B2:C2')
        ws.merge_cells('D2:F2')
        ws.append([None, 0.13, 0.3, 'Cost item', 'Account', 'Report'])
        ws.append([1, 2, 3, 4, 5, 6])
        ws.append(['0001', 'yes', 'yes', 'Payroll', '44', 'included'])
    table = parse(tmp_path, configure)[0]
    descriptor = table['descriptor']
    assert descriptor.header_rows == [2, 3]
    assert descriptor.columns[1].header_path == ['Rates', '0.13']
    assert descriptor.columns[3].header_addresses == ['D2', 'D3']
    assert next(r for r in table['rows'] if r.row_number == 4).cells[6].raw_value == 6


def test_numeric_leaf_without_ordinal_boundary_stays_ambiguous(tmp_path):
    def configure(ws):
        ws.append(['Code', 'Rates', None, 'Accounting', None, None])
        ws.merge_cells('A1:A2')
        ws.merge_cells('B1:C1')
        ws.merge_cells('D1:F1')
        ws.append([None, 0.13, 0.3, 'Cost item', 'Account', 'Report'])
        ws.append(['0001', 'yes', 'yes', 'Payroll', '44', 'included'])
    assert parse(tmp_path, configure)[0]['descriptor'].header_rows == []


def test_corrupt_ordinal_row_does_not_promote_numeric_leaf(tmp_path):
    def configure(ws):
        ws.append(['Code', 'Rates', None, 'Accounting', None, None])
        ws.merge_cells('A1:A2')
        ws.merge_cells('B1:C1')
        ws.merge_cells('D1:F1')
        ws.append([None, 0.13, 0.3, 'Cost item', 'Account', 'Report'])
        ws.append([99, 2, 3, 4, 5, 6])
        ws.append(['0001', 'yes', 'yes', 'Payroll', '44', 'included'])
    assert parse(tmp_path, configure)[0]['descriptor'].header_rows == []


@pytest.mark.parametrize("header_count", [0, 1, None])
@pytest.mark.parametrize("nested", [False, True])
def test_declared_header_count_preserves_first_data_record(tmp_path, header_count, nested):
    from openpyxl.worksheet.table import TableColumn as ExcelColumn
    import tests.fixtures as fixtures
    wb = Workbook()
    ws = wb.active
    ws.append(["001", 7] if header_count == 0 else ["Code", "Value"])
    ws.append(["002", 8])
    ws.add_table(Table(displayName="Declared", ref="A1:B2", headerRowCount=header_count,
                       tableColumns=[ExcelColumn(id=1, name="Code"), ExcelColumn(id=2, name="Value")]))
    path = tmp_path / "declared.xlsx"
    wb.save(path)
    if nested:
        path = fixtures.make_docx_with_embedded_xlsx(tmp_path / "nested.docx", path.read_bytes())
    sink = RecordingSink()
    parse_document_result(path, table_sink=sink)
    table, = sink.tables.values()
    d = table["descriptor"]
    assert d.source_id == ("root/0" if nested else "root")
    assert d.header_rows == ([] if header_count == 0 else [1])
    assert d.method == "explicit" and not d.warnings
    assert [c.header_path for c in d.columns] == [["Code"], ["Value"]]
    assert [c.header_addresses for c in d.columns] == ([[], []] if header_count == 0 else [["A1"], ["B1"]])
    data_rows = [row for row in table["rows"] if row.row_number not in d.header_rows]
    assert [row.row_number for row in data_rows] == ([1, 2] if header_count == 0 else [2])
    if header_count == 0:
        assert data_rows[0].cells[1].raw_value == "001"
        assert data_rows[0].cells[2].raw_value == 7

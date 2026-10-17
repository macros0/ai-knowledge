"""Structural Excel contracts against original OOXML coordinates and values."""
import json
import zipfile

import pytest
from openpyxl import Workbook

from docparser import parse_document_result
from docparser.source_model import ParseContext
import tests.fixtures as fixtures


class RecordingSink:
    def __init__(self):
        self.tables = {}
        self.batches = []

    def begin_table(self, descriptor):
        self.tables[descriptor.table_id] = {'descriptor': descriptor, 'rows': [], 'summary': None}

    def append_rows(self, table_id, rows):
        from dataclasses import asdict
        self.batches.append((len(rows), len(json.dumps([asdict(r) for r in rows]).encode())))
        self.tables[table_id]['rows'].extend(rows)

    def finish_table(self, table_id, summary):
        self.tables[table_id]['summary'] = summary


def workbook(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Codes'
    ws.append(['Code', 'Quantity', 'Enabled', 'Blank', 'Formula'])
    ws.append([7, 0, False, None, '=1+1'])
    ws['A2'].number_format = '00000'
    ws['B3'] = 'repeat'
    ws['B4'] = 'repeat'
    ws['C3'] = '<script>|test\\nnext'
    ws.row_dimensions[3].hidden = True
    ws.column_dimensions['C'].hidden = True
    ws['XFD9'] = 'remote'
    ws['A1000000'].number_format = '0.00'  # formatting alone is not data
    path = tmp_path / 'values.xlsx'
    wb.save(path)
    return path


def structural(tmp_path, path=None):
    sink = RecordingSink()
    result = parse_document_result(path or workbook(tmp_path), table_sink=sink)
    rows = [r for t in sink.tables.values() for r in t['rows']]
    return result, sink, {c.address: c for r in rows for c in r.cells.values()}


def test_excel_cells_keep_original_coordinates(tmp_path):
    result, sink, cells = structural(tmp_path)
    assert cells['XFD9'].raw_value == 'remote'
    assert cells['A2'].raw_value == 7
    assert cells['A2'].display_text == '00007'
    assert cells['A2'].number_format == '00000'
    assert not any(t['descriptor'].cell_range.max_row == 1000000 for t in sink.tables.values())
    assert all(t['summary'].complete for t in sink.tables.values())
    assert all('cells' not in b.meta for b in result.blocks)
    assert max(size for _, size in sink.batches) <= 1024 * 1024


def test_blank_zero_false_are_distinct(tmp_path):
    _, _, cells = structural(tmp_path)
    assert cells['B2'].raw_value == 0 and cells['B2'].state == 'value'
    assert cells['C2'].raw_value is False and cells['C2'].value_type == 'boolean'
    assert cells['E2'].state == 'formula_missing'
    assert cells['E2'].formula == '1+1' and cells['E2'].cached_value is None
    assert cells.get('D2') is None or cells['D2'].state == 'blank'
    assert cells['C3'].hidden


def test_nested_xlsx_keeps_source_id(tmp_path):
    source = workbook(tmp_path)
    docx = fixtures.make_docx_with_embedded_xlsx(tmp_path / 'nested.docx', source.read_bytes())
    result, sink, _ = structural(tmp_path, docx)
    assert sink.tables
    source_ids = {t['descriptor'].source_id for t in sink.tables.values()}
    assert source_ids == {'root/0'}
    assert source_ids <= {s.source_id for s in result.sources}


def test_table_sink_is_optional(tmp_path):
    source = fixtures.make_xlsx(tmp_path / 'plain.xlsx', rows=[['Code', 'Value'], ['K', 0]])
    result = parse_document_result(source)
    assert any(b.type == 'table' and 'K' in b.text for b in result.blocks)
    assert ParseContext('plain.xlsx').table_sink is None


def test_sink_batches_are_bounded_and_repeated_rows_survive(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(['Code', 'Value'])
    for _ in range(600):
        ws.append(['same', 'value' * 500])
    source = tmp_path / 'many.xlsx'
    wb.save(source)
    _, sink, _ = structural(tmp_path, source)
    rows = [r for t in sink.tables.values() for r in t['rows'] if r.row_number > 1]
    assert len(rows) == 600 and len({r.row_number for r in rows}) == 600
    assert all(count <= 256 and size <= 1024 * 1024 for count, size in sink.batches)


def test_failed_sink_is_never_finished(tmp_path):
    class BrokenSink(RecordingSink):
        def append_rows(self, table_id, rows):
            raise RuntimeError('cancelled')
    sink = BrokenSink()
    with pytest.raises(RuntimeError, match='cancelled'):
        parse_document_result(workbook(tmp_path), table_sink=sink)
    assert sink.tables and all(t['summary'] is None for t in sink.tables.values())


def test_formula_cached_empty_and_missing_are_distinct(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(['Empty', 'Missing', 'Error'])
    ws.append(['=IF(1,"",0)', '=1+1', '#DIV/0!'])
    path = tmp_path / 'cache.xlsx'
    wb.save(path)
    with zipfile.ZipFile(path) as z:
        entries = {i.filename: z.read(i) for i in z.infolist()}
    entries['xl/worksheets/sheet1.xml'] = entries['xl/worksheets/sheet1.xml'].replace(b'<c r="A2"><f>', b'<c r="A2" t="str"><f>')
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, payload in entries.items():
            z.writestr(name, payload)
    _, _, cells = structural(tmp_path, path)
    assert cells['A2'].state == 'value' and cells['A2'].cached_value == ''
    assert cells['B2'].state == 'formula_missing'
    assert cells['C2'].state == 'error' and cells['C2'].raw_value == '#DIV/0!'


def test_numeric_xml_precision_is_not_lost_to_binary_float(tmp_path):
    import zipfile
    from decimal import Decimal
    wb=Workbook();ws=wb.active;ws.append(['Code','Value']);ws.append(['K',9.6])
    source=tmp_path/'precise.xlsx';wb.save(source)
    with zipfile.ZipFile(source) as archive:entries={n:archive.read(n) for n in archive.namelist()}
    entries['xl/worksheets/sheet1.xml']=entries['xl/worksheets/sheet1.xml'].replace(b'<v>9.6</v>',b'<v>9.5999999999999996</v>')
    with zipfile.ZipFile(source,'w') as archive:
        for name,value in entries.items():archive.writestr(name,value)
    sink=RecordingSink();parse_document_result(source,table_sink=sink)
    table=next(iter(sink.tables.values()));cell=table['rows'][1].cells[2]
    assert Decimal(str(cell.raw_value))==Decimal('9.5999999999999996')
    assert cell.value_type=='number' and cell.display_text=='9.6'


def test_styled_empty_right_edge_keeps_interior_columns_merges_and_last_row(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.merge_cells('A1:C1')
    ws['A1'] = 'Identity'
    ws['D1'] = 'Amount'
    ws.append(['Code', None, 'Enabled', 'Quantity'])
    ws.append(['K1', None, False, 0])
    ws.append(['K2', None, True, 7])
    ws['D5'] = 'late value'
    ws['XFD5'].number_format = '0.00'
    ws['XFD1000000'].number_format = '0.00'
    source = tmp_path / 'styled-tail.xlsx'
    wb.save(source)

    _, sink, cells = structural(tmp_path, source)
    assert cells['C3'].raw_value is False
    assert cells['D3'].raw_value == 0
    assert cells['D5'].raw_value == 'late value'
    assert 'XFD5' not in cells and 'XFD1000000' not in cells
    assert cells.get('B3') is None or cells['B3'].state == 'blank'
    assert 'C3' in cells and 'D3' in cells  # no shift across the empty interior B
    assert all(t['descriptor'].cell_range.max_column <= 4 for t in sink.tables.values())
    assert all(t['descriptor'].cell_range.max_row <= 5 for t in sink.tables.values())
    assert any(merge.min_column == 1 and merge.max_column == 3 and merge.min_row == 1
               for t in sink.tables.values() for merge in t['descriptor'].merges)


@pytest.mark.parametrize("xml_value,expected", [
    ("7", "00007"), ("7.0", "00007"), ("7E+00", "00007"),
    ("0.0", "00000"), ("7.5", "7.5"),
    ("9007199254740993.0", "9007199254740993"),
])
@pytest.mark.parametrize("formula", [False, True])
def test_zero_padded_integral_xml_preserves_value_and_display(tmp_path, xml_value, expected, formula):
    from decimal import Decimal
    wb = Workbook()
    ws = wb.active
    ws.append(["Code", "Enabled"])
    ws.append([7, False])
    ws["A2"].number_format = "00000"
    ws["B2"].number_format = "00000"
    path = tmp_path / "integral-code.xlsx"
    wb.save(path)
    with zipfile.ZipFile(path) as archive:
        entries = {n: archive.read(n) for n in archive.namelist()}
    payload = ("<f>7</f>" if formula else "") + f"<v>{xml_value}</v>"
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(b"<v>7</v>", payload.encode())
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    _, _, cells = structural(tmp_path, path)
    cell = cells["A2"]
    assert cell.display_text == expected
    assert cell.value_type == "number" and Decimal(str(cell.raw_value)) == Decimal(xml_value)
    assert cell.cached_value == (cell.raw_value if formula else None)
    assert cells["B2"].raw_value is False and cells["B2"].value_type == "boolean"
    assert cells["B2"].display_text == "False"

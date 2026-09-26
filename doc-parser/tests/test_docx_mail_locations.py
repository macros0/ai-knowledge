"""Real OOXML relationships in nested tables, textboxes and header/footer parts."""
from pathlib import Path

import pytest
from docparser import parse_document_result
from docx import Document
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from lxml import etree

from tests.mail_fixtures import unicode_msg

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
O_NS = 'urn:schemas-microsoft-com:office:office'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def _ole(paragraph, part, *, rid=None, external=False):
    if rid is None:
        if external:
            rid = part.relate_to('https://example.invalid/private.msg', R + '/oleObject', is_external=True)
        else:
            binary = Part(PackURI('/word/embeddings/approval.msg'), 'application/vnd.ms-outlook',
                          unicode_msg(body='UNIQUE APPROVAL BODY'), part.package)
            rid = part.relate_to(binary, R + '/oleObject')
    run = etree.SubElement(paragraph, f'{{{W}}}r')
    obj = etree.SubElement(run, f'{{{W}}}object')
    ole = etree.SubElement(obj, f'{{{O_NS}}}OLEObject')
    ole.set('Type', 'Link' if external else 'Embed')
    ole.set('ProgID', 'Outlook.File.msg.15')
    ole.set(f'{{{R}}}id', rid)
    return rid


def _parse(doc, tmp_path):
    path = tmp_path / 'locations.docx'
    doc.save(path)
    return parse_document_result(path, attachments_dir=tmp_path / 'attachments')


def test_nested_table_extracts_each_xml_occurrence_only_once(tmp_path):
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    nested = cell.add_table(rows=1, cols=1)
    _ole(nested.cell(0, 0).paragraphs[0]._p, doc.part)
    result = _parse(doc, tmp_path)
    assert [n.source_id for n in result.sources] == ['root', 'root/0']
    assert sum(b.text == 'UNIQUE APPROVAL BODY' for b in result.blocks) == 1


@pytest.mark.parametrize('story', ['header', 'footer'])
def test_shared_header_footer_uses_own_relationships_once(tmp_path, story):
    doc = Document()
    doc.add_paragraph('BODY REMAINS PRESENT')
    section_part = getattr(doc.sections[0], story)
    _ole(section_part.paragraphs[0]._p, section_part.part)
    doc.add_section()  # The same header/footer is inherited, not a second XML occurrence.
    result = _parse(doc, tmp_path)
    assert [n.source_id for n in result.sources] == ['root', 'root/0']
    assert result.sources[1].metadata['document_location'] == story
    assert sum(b.text == 'UNIQUE APPROVAL BODY' for b in result.blocks) == 1
    marker = next(b for b in result.blocks if b.type == 'attachment')
    assert Path(marker.meta['saved_path']).read_bytes() == unicode_msg(body='UNIQUE APPROVAL BODY')


def test_textbox_inside_table_is_not_visited_twice(tmp_path):
    doc = Document()
    paragraph = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]._p
    run = etree.SubElement(paragraph, f'{{{W}}}r')
    box = etree.SubElement(run, f'{{{W}}}txbxContent')
    inner = etree.SubElement(box, f'{{{W}}}p')
    _ole(inner, doc.part)
    result = _parse(doc, tmp_path)
    assert [n.source_id for n in result.sources] == ['root', 'root/0']
    assert result.sources[1].metadata['document_location'] == 'textbox'


def test_distinct_occurrences_of_same_relationship_are_preserved(tmp_path):
    doc = Document()
    rid = _ole(doc.add_paragraph()._p, doc.part)
    _ole(doc.add_paragraph()._p, doc.part, rid=rid)
    result = _parse(doc, tmp_path)
    assert [n.source_id for n in result.sources] == ['root', 'root/0', 'root/1']
    assert sum(b.text == 'UNIQUE APPROVAL BODY' for b in result.blocks) == 2


def test_external_ole_is_explicit_without_download_or_mail_metadata(tmp_path, monkeypatch):
    import socket
    def deny_network(*args, **kwargs):
        raise AssertionError('Linked OLE must never fetch external bytes')
    monkeypatch.setattr(socket.socket, 'connect', deny_network)
    doc = Document()
    _ole(doc.add_paragraph()._p, doc.part, external=True)
    result = _parse(doc, tmp_path)
    assert len(result.sources) == 2
    assert result.sources[1].kind == 'attachment'
    assert 'sender' not in result.sources[1].metadata
    assert result.warnings == [{'code': 'external_attachment', 'source_id': 'root/0'}]
    marker = next(b for b in result.blocks if b.type == 'attachment')
    assert marker.meta['extraction_status'] == 'unsupported'
    assert not marker.meta.get('saved_path')


@pytest.mark.parametrize('external', [False, True])
def test_docx_node_limit_stops_relationship_resolution_with_one_remainder_marker(tmp_path, monkeypatch, external):
    from docparser import docx_parser, parse_document
    from docparser.embedded import AttachmentBudget
    from docparser.source_model import ParseContext

    doc = Document()
    rid = _ole(doc.add_paragraph()._p, doc.part, external=external)
    for _ in range(7):
        _ole(doc.add_paragraph()._p, doc.part, rid=rid, external=external)
    doc.add_paragraph('BODY AFTER ATTACHMENTS')
    path = tmp_path / 'many-objects.docx'
    doc.save(path)
    resolved = []
    original = docx_parser._resolve_ole_attachment

    def observe(part, ole):
        resolved.append(ole['r_id'])
        return original(part, ole)

    monkeypatch.setattr(docx_parser, '_resolve_ole_attachment', observe)
    context = ParseContext(path.name)
    blocks = parse_document(path, context=context, budget=AttachmentBudget(max_nodes=2))
    assert len(resolved) == 2
    assert len(context.sources) == 4  # root + admitted occurrences + one remainder marker
    assert sum(b.meta.get('extraction_status') == 'skipped_count' for b in blocks) == 1
    assert {'code': 'attachment_count_exceeded', 'source_id': 'root/2'} in context.warnings
    assert any(b.text == 'BODY AFTER ATTACHMENTS' for b in blocks)

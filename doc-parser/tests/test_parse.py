"""Тесты извлечения блоков из DOCX (текст, заголовки, таблицы, комментарии), XLSX и PDF."""
from pathlib import Path

import pytest

from docparser import ParseError, SUPPORTED_EXTENSIONS, parse_document
from tests.fixtures import (
    inject_comment,
    inject_comments,
    make_docx,
    make_docx_with_image,
    make_pdf,
    make_pdf_scanned_like,
    make_pdf_with_image,
    make_xlsx,
)


class TestDispatch:
    def test_supported_extensions(self):
        assert SUPPORTED_EXTENSIONS == {".docx", ".xlsx", ".pdf"}

    def test_unknown_extension_raises(self, tmp_path: Path):
        f = tmp_path / "doc.txt"
        f.write_text("hello", encoding="utf-8")
        with pytest.raises(ParseError):
            parse_document(f, filename="doc.txt")

    def test_extension_from_filename_override(self, tmp_path: Path):
        f = tmp_path / "noext"
        f.write_bytes(b"not a real docx")
        with pytest.raises(Exception):
            parse_document(f, filename="noext.docx")


class TestDocx:
    def test_paragraphs_and_heading(self, tmp_path: Path):
        docx = make_docx(tmp_path / "doc.docx", heading="Заголовок", paragraphs=["Абзац один", "Абзац два"])
        blocks = parse_document(docx)
        types = [b.type for b in blocks]
        assert types == ["heading", "paragraph", "paragraph"]
        assert blocks[0].text == "Заголовок"
        assert blocks[0].level == 1
        assert blocks[1].text == "Абзац один"

    def test_comments(self, tmp_path: Path):
        docx = make_docx(tmp_path / "doc.docx", paragraphs=["Абзац с замечанием"])
        inject_comment(docx, author="Рецензент", text="Поправьте формулировку")
        blocks = parse_document(docx)
        comments = [b for b in blocks if b.type == "comment"]
        assert len(comments) == 1
        assert comments[0].text == "Поправьте формулировку"
        assert comments[0].meta["author"] == "Рецензент"

    def test_comment_single_block_still_has_thread_meta(self, tmp_path: Path):
        """Одиночный комментарий — тред из одной записи (единый формат meta.thread)."""
        docx = make_docx(tmp_path / "doc.docx", paragraphs=["Абзац"])
        inject_comment(docx, author="Рецензент", text="Замечание")
        blocks = parse_document(docx)
        c = [b for b in blocks if b.type == "comment"][0]
        assert len(c.meta["thread"]) == 1
        assert c.meta["thread"][0]["text"] == "Замечание"
        assert "resolved" not in c.meta  # commentsExtended нет — статус неизвестен

    def test_thread_question_and_reply_merged(self, tmp_path: Path):
        """Тред «вопрос → ответ» (commentsExtended paraIdParent) — один блок."""
        docx = make_docx(tmp_path / "doc.docx", paragraphs=["Алгоритм выбора табельного"])
        inject_comments(docx, comments=[
            {"id": "137", "author": "Волкова Анастасия", "text": "Какой ТН считать свежим?",
             "para_id": "552F5CF1", "done": "1"},
            {"id": "138", "author": "Сагитов Алексей", "text": "Наибольший табельный — самый свежий.",
             "para_id": "0E439753", "parent_para_id": "552F5CF1", "done": "1"},
        ])
        blocks = parse_document(docx)
        comments = [b for b in blocks if b.type == "comment"]
        assert len(comments) == 1
        thread = comments[0].meta["thread"]
        assert len(thread) == 2
        assert thread[0]["author"] == "Волкова Анастасия"  # вопрос первым
        assert thread[0]["text"] == "Какой ТН считать свежим?"
        assert thread[1]["text"] == "Наибольший табельный — самый свежий."
        assert comments[0].meta["resolved"] is True
        assert comments[0].meta["authors"] == ["Волкова Анастасия", "Сагитов Алексей"]
        assert comments[0].meta["date"].startswith("2026-08-12")

    def test_thread_unresolved_flag(self, tmp_path: Path):
        docx = make_docx(tmp_path / "doc.docx", paragraphs=["Абзац"])
        inject_comments(docx, comments=[
            {"id": "1", "author": "Рецензент", "text": "Вопрос", "para_id": "AAAA0001", "done": "0"},
            {"id": "2", "author": "Автор", "text": "Ответ", "para_id": "BBBB0002",
             "parent_para_id": "AAAA0001", "done": "0"},
        ])
        blocks = parse_document(docx)
        comments = [b for b in blocks if b.type == "comment"]
        assert len(comments) == 1
        assert comments[0].meta["resolved"] is False

    def test_thread_reply_order_deterministic_no_extended(self, tmp_path: Path):
        """РЕГРЕССИЯ: без commentsExtended порядок комментариев в одном абзаце
        был недетерминирован (итерация set — порядок зависит от hash-seed),
        ответ (id 138) мог печататься раньше вопроса (id 137). Теперь —
        сортировка по числовому id."""
        docx = make_docx(tmp_path / "doc.docx", paragraphs=["Абзац с двумя замечаниями"])
        inject_comments(docx, comments=[
            {"id": "138", "author": "Автор", "text": "Ответ автора"},
            {"id": "137", "author": "Рецензент", "text": "Вопрос рецензента"},
        ])
        blocks = parse_document(docx)
        comments = [b for b in blocks if b.type == "comment"]
        assert [c.text for c in comments] == ["Вопрос рецензента", "Ответ автора"]

    def test_thread_reply_anchored_before_question(self, tmp_path: Path):
        """РЕГРЕССИЯ: ответ физически стоит в document.xml РАНЬШЕ вопроса
        (якорь ответа в первом абзаце, якорь вопроса во втором). Тред всё
        равно начинается с вопроса (структура от paraIdParent, а не от позиции
        в XML) и эмитится на первом физическом якоре."""
        docx = make_docx(tmp_path / "doc.docx", paragraphs=["Первый абзац", "Второй абзац"])
        inject_comments(
            docx,
            comments=[
                {"id": "137", "author": "Рецензент", "text": "Вопрос",
                 "para_id": "AAAA0001", "done": "1"},
                {"id": "138", "author": "Автор", "text": "Ответ",
                 "para_id": "BBBB0002", "parent_para_id": "AAAA0001", "done": "1"},
            ],
            anchors={"137": 1, "138": 0},  # ответ якорится раньше вопроса
        )
        blocks = parse_document(docx)
        comments = [b for b in blocks if b.type == "comment"]
        assert len(comments) == 1
        assert [e["text"] for e in comments[0].meta["thread"]] == ["Вопрос", "Ответ"]
        # тред эмитится на первом физическом якоре: до второго абзаца
        idx_comment = blocks.index(comments[0])
        idx_second = next(
            i for i, b in enumerate(blocks)
            if b.type == "paragraph" and b.text == "Второй абзац"
        )
        assert idx_comment < idx_second

    def test_comment_context_and_section_meta(self, tmp_path: Path):
        docx = make_docx(tmp_path / "doc.docx", heading="Раздел А", paragraphs=["Целевой абзац"])
        inject_comments(docx, comments=[{"id": "1", "author": "Рецензент", "text": "Вопрос"}],
                        anchors={"1": 1})  # 0 — заголовок «Раздел А», 1 — целевой абзац
        blocks = parse_document(docx)
        c = [b for b in blocks if b.type == "comment"][0]
        assert c.meta["context"] == "Целевой абзац"
        assert c.meta["section"] == "Раздел А"

    def test_comment_in_table_cell_emitted_after_table(self, tmp_path: Path):
        """РЕГРЕССИЯ: комментарии в ячейках таблицы сваливались в конец
        документа. Теперь — сразу после таблицы."""
        from docx import Document

        doc = Document()
        doc.add_paragraph("Текст до таблицы")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Ячейка с замечанием"
        table.cell(0, 1).text = "Соседняя"
        doc.add_paragraph("Текст после таблицы")
        path = tmp_path / "t.docx"
        doc.save(str(path))
        inject_comments(path, comments=[{"id": "5", "author": "Рецензент", "text": "Замечание к ячейке"}],
                        anchors={"5": "cell"})
        blocks = parse_document(path)
        comments = [b for b in blocks if b.type == "comment"]
        assert len(comments) == 1
        ti = next(i for i, b in enumerate(blocks) if b.type == "table")
        ci = blocks.index(comments[0])
        pi = next(i for i, b in enumerate(blocks) if b.type == "paragraph" and b.text == "Текст после таблицы")
        assert ti < ci < pi

    def test_unanchored_comment_appended_at_end(self, tmp_path: Path):
        """Комментарий без якоря в теле (маркеры удалены) — в конец документа."""
        docx = make_docx(tmp_path / "doc.docx", paragraphs=["Абзац"])
        inject_comments(docx, comments=[{"id": "9", "author": "Рецензент", "text": "Одинокое замечание"}])
        # удаляем якорь: пересобираем document.xml без commentRangeStart
        import zipfile

        with zipfile.ZipFile(str(docx)) as zf:
            entries = {n: zf.read(n) for n in zf.namelist()}
        doc_xml = entries["word/document.xml"].decode("utf-8")
        import re as _re

        doc_xml = _re.sub(r'<w:commentRangeStart[^/]*/>', "", doc_xml)
        doc_xml = _re.sub(r'<w:commentRangeEnd[^/]*/>', "", doc_xml)
        doc_xml = _re.sub(r'<w:commentReference[^/]*/>', "", doc_xml)
        entries["word/document.xml"] = doc_xml.encode("utf-8")
        with zipfile.ZipFile(str(docx), "w", zipfile.ZIP_DEFLATED) as zf:
            for name, data in entries.items():
                zf.writestr(name, data)
        blocks = parse_document(docx)
        comments = [b for b in blocks if b.type == "comment"]
        assert len(comments) == 1
        assert blocks[-1].type == "comment"

    def test_empty_paragraphs_skipped(self, tmp_path: Path):
        docx = make_docx(tmp_path / "doc.docx", paragraphs=["   ", "", "Текст"])
        blocks = parse_document(docx)
        assert [b.text for b in blocks if b.type == "paragraph"] == ["Текст"]

    def test_code_paragraphs_detected(self, tmp_path: Path):
        from docx import Document

        doc = Document()
        p1 = doc.add_paragraph()
        r1 = p1.add_run("def bridge():")
        r1.font.name = "Consolas"
        p2 = doc.add_paragraph()
        r2 = p2.add_run("    return True")
        r2.font.name = "Consolas"
        doc.add_paragraph("Обычный текст")
        path = tmp_path / "code.docx"
        doc.save(str(path))

        blocks = parse_document(path)
        code = [b for b in blocks if b.type == "code"]
        assert len(code) == 1
        assert "def bridge():" in code[0].text
        assert "    return True" in code[0].text
        assert [b.text for b in blocks if b.type == "paragraph"] == ["Обычный текст"]

    def test_inline_images_extracted(self, tmp_path: Path):
        docx = make_docx_with_image(tmp_path / "img.docx")
        att_dir = tmp_path / "attachments"
        blocks = parse_document(docx, attachments_dir=att_dir)

        paragraphs = [b.text for b in blocks if b.type == "paragraph"]
        assert paragraphs == ["Перед картинкой", "После картинки"]

        images = [b for b in blocks if b.type == "image"]
        assert len(images) == 1
        assert images[0].meta["kind"] == "image"
        saved = Path(images[0].meta["saved_path"])
        assert saved.exists()
        assert saved.suffix in (".png", ".jpg", ".jpeg")

    def test_inline_images_without_attachments_dir(self, tmp_path: Path):
        docx = make_docx_with_image(tmp_path / "img.docx")
        blocks = parse_document(docx)
        images = [b for b in blocks if b.type == "image"]
        assert len(images) == 1
        assert "saved_path" not in images[0].meta


class TestXlsx:
    def test_sheets_to_tables(self, tmp_path: Path):
        xlsx = make_xlsx(tmp_path / "t.xlsx", sheet="Правила", rows=[["Код", "Описание"], [1, "Первое"], [2, "Второе"]])
        blocks = parse_document(xlsx)
        headings = [b for b in blocks if b.type == "heading"]
        tables = [b for b in blocks if b.type == "table"]
        assert headings[0].text == "Таблица: Правила"
        assert len(tables) == 1
        assert "Код" in tables[0].text
        assert "Второе" in tables[0].text


class TestPdf:
    def test_text_by_pages(self, tmp_path: Path):
        pdf = make_pdf(tmp_path / "p.pdf", "Строка из мини-pdf")
        blocks = parse_document(pdf)
        assert blocks
        assert all(b.type == "paragraph" for b in blocks)
        assert any("мини-pdf" in b.text for b in blocks)
        assert blocks[0].meta["page"] == 1

    def test_page_images_extracted(self, tmp_path: Path):
        pdf = make_pdf_with_image(tmp_path / "p.pdf")
        att_dir = tmp_path / "attachments"
        blocks = parse_document(pdf, attachments_dir=att_dir)

        assert any(b.type == "paragraph" for b in blocks)
        images = [b for b in blocks if b.type == "image"]
        assert len(images) == 1
        assert images[0].meta["kind"] == "image"
        assert images[0].meta["page"] == 1
        assert "Страница 1" in images[0].meta["caption"]
        saved = Path(images[0].meta["saved_path"])
        assert saved.exists()
        assert saved.read_bytes() == att_dir.joinpath(saved.name).read_bytes()

    def test_scanned_page_rendered_via_fallback(self, tmp_path: Path):
        pdf = make_pdf_scanned_like(tmp_path / "scan.pdf")
        att_dir = tmp_path / "attachments"
        blocks = parse_document(pdf, attachments_dir=att_dir)

        assert not any(b.type == "paragraph" for b in blocks)
        images = [b for b in blocks if b.type == "image"]
        assert len(images) == 1
        assert images[0].meta["kind"] == "image"
        assert images[0].meta["page"] == 1
        assert "Страница 1" in images[0].meta["caption"]
        saved = Path(images[0].meta["saved_path"])
        assert saved.exists()
        assert saved.suffix == ".jpg"

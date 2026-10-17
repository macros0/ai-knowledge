"""Synthetic inputs for attachment assessment; no customer document content."""

import io
from email.message import EmailMessage

from docx import Document
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from lxml import etree
from openpyxl import Workbook


def spreadsheet(rows=5, columns=2, *, text="included evidence", tail=False):
    book = Workbook()
    sheet = book.active
    for _ in range(rows):
        sheet.append([text] * columns)
    if tail:
        sheet.cell(100000, 300).number_format = "0.00"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def office_container(path, attachments):
    doc = Document()
    doc.add_paragraph("ROOT EVIDENCE REMAINS AVAILABLE")
    namespace = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    for index, (name, data) in enumerate(attachments):
        part = Part(
            PackURI(f"/word/embeddings/{index}/{name}"),
            "application/octet-stream",
            data,
            doc.part.package,
        )
        relation = doc.part.relate_to(part, namespace + "/oleObject")
        paragraph = doc.add_paragraph()._p
        obj = etree.SubElement(
            paragraph, "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}object"
        )
        ole = etree.SubElement(obj, "{urn:schemas-microsoft-com:office:office}OLEObject")
        ole.set("{" + namespace + "}id", relation)
        ole.set("ProgID", "Package")
    doc.save(path)
    return path


def mail_bytes(attachments, *, body="USEFUL MAIL BODY"):
    message = EmailMessage()
    message["Subject"] = "Synthetic attachment review"
    message.set_content(body)
    for name, data in attachments:
        message.add_attachment(data, maintype="application", subtype="octet-stream", filename=name)
    return message.as_bytes()


def msg_bytes(attachments, *, package=None):
    """Synthetic CFB/MAPI with normal FAT streams; no Outlook or customer data."""
    import struct

    free, end, fat = 0xFFFFFFFF, 0xFFFFFFFE, 0xFFFFFFFD
    entries = [{"name": "Root Entry", "type": 5, "parent": None}]

    def storage(name, parent):
        entries.append({"name": name, "type": 1, "parent": parent})
        return len(entries) - 1

    def stream(name, data, parent, header=0):
        length = max(len(data), 4096)
        if header:
            length = header + ((length - header + 15) // 16) * 16
        entries.append(
            {"name": name, "type": 2, "parent": parent, "data": data.ljust(length, b"\0")}
        )

    stream("__properties_version1.0", b"\0" * 32, 0, 32)
    stream("__substg1.0_0037001F", "Synthetic approval\0".encode("utf-16-le"), 0)
    stream("__substg1.0_1000001F", "USEFUL MSG BODY\0".encode("utf-16-le"), 0)
    if package is not None:
        stream('Package',package,0)
    for index, (name, data) in enumerate(attachments):
        node = storage(f"__attach_version1.0_#{index:08X}", 0)
        fixed = b"\0" * 8 + struct.pack("<II8s", 0x37050003, 0, struct.pack("<I", 1) + b"\0" * 4)
        stream("__properties_version1.0", fixed, node, 8)
        stream("__substg1.0_3707001F", (name + "\0").encode("utf-16-le"), node)
        stream("__substg1.0_37010102", data, node)
    sectors = []
    chains = []
    for entry in entries:
        data = entry.get("data", b"")
        entry["start"] = len(sectors) if data else end
        entry["size"] = len(data)
        for offset in range(0, len(data), 512):
            sectors.append(data[offset : offset + 512].ljust(512, b"\0"))
        if data:
            chains.append((entry["start"], len(sectors)))
    directory_start = len(sectors)
    directory = b""
    for index, entry in enumerate(entries):
        siblings = sorted(
            [i for i, e in enumerate(entries) if e["parent"] == entry["parent"] and i],
            key=lambda i: (len(entries[i]["name"]), entries[i]["name"].upper()),
        )
        pos = siblings.index(index) if index in siblings else -1
        right = siblings[pos + 1] if pos >= 0 and pos + 1 < len(siblings) else free
        children = sorted(
            [i for i, e in enumerate(entries) if e["parent"] == index],
            key=lambda i: (len(entries[i]["name"]), entries[i]["name"].upper()),
        )
        name = (entry["name"] + "\0").encode("utf-16-le")
        directory += name.ljust(64, b"\0") + struct.pack(
            "<HBBIII", len(name), entry["type"], 1, free, right, children[0] if children else free
        )
        directory += b"\0" * 36 + struct.pack("<IQ", entry["start"], entry["size"])
    sectors.extend(directory[o : o + 512].ljust(512, b"\0") for o in range(0, len(directory), 512))
    chains.append((directory_start, len(sectors)))
    count = 1
    while count * 128 < len(sectors) + count:
        count += 1
    fat_start = len(sectors)
    links = [free] * (count * 128)
    for start, stop in chains:
        for sector in range(start, stop):
            links[sector] = sector + 1 if sector + 1 < stop else end
    for sector in range(fat_start, fat_start + count):
        links[sector] = fat
    header = bytearray(512)
    header[:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    struct.pack_into("<HHHHH", header, 24, 0x3E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<IIIIIIIII", header, 40, 0, count, directory_start, 0, 4096, end, 0, end, 0)
    struct.pack_into(
        "<109I", header, 76, *(list(range(fat_start, fat_start + count)) + [free] * (109 - count))
    )
    return bytes(header) + b"".join(sectors) + struct.pack("<" + "I" * len(links), *links)

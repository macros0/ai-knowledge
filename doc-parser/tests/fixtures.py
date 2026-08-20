"""Фикстуры: создание реальных файлов на лету для тестов.

- DOCX с комментариями рецензентов (инъекция comments.xml + разметки в document.xml).
- DOCX со встроенным XLSX (OLE-объект, rels + [Content_Types]).
- XLSX через openpyxl.
- PDF через reportlab.
"""
import io
import zipfile
from pathlib import Path

from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
O = "urn:schemas-microsoft-com:office:office"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"


# ---------------------------------------------------------------- helpers
def _tiny_png(width: int = 4, height: int = 4) -> bytes:
    """Генерирует маленький валидный PNG для тестов изображений."""
    import io

    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (width, height), (200, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


def _read_zip(path) -> dict[str, bytes]:
    with zipfile.ZipFile(str(path)) as zf:
        return {n: zf.read(n) for n in zf.namelist()}


def _write_zip(path, entries: dict[str, bytes]) -> None:
    with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)


def _xml(data: bytes) -> etree._Element:
    return etree.fromstring(data)


def _dump(root: etree._Element) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _add_content_type(entries: dict[str, bytes], part_name: str, content_type: str) -> None:
    ct = _xml(entries["[Content_Types].xml"])
    override = etree.SubElement(ct, f"{{{CT}}}Override")
    override.set("PartName", part_name)
    override.set("ContentType", content_type)
    entries["[Content_Types].xml"] = _dump(ct)


def _add_relationship(entries: dict[str, bytes], rels_path: str, rid: str, rel_type: str, target: str) -> None:
    rels = _xml(entries[rels_path])
    rel = etree.SubElement(rels, f"{{{PKG}}}Relationship")
    rel.set("Id", rid)
    rel.set("Type", rel_type)
    rel.set("Target", target)
    entries[rels_path] = _dump(rels)


# ---------------------------------------------------------------- docx
def make_docx(path: Path, paragraphs: list[str] | None = None, heading: str | None = None) -> Path:
    from docx import Document

    doc = Document()
    if heading:
        doc.add_heading(heading, level=1)
    for p in paragraphs or ["Первый абзац", "Настройка VLAN", "Конец документа"]:
        doc.add_paragraph(p)
    doc.save(str(path))
    return path


def make_docx_with_image(path: Path, image_bytes: bytes | None = None) -> Path:
    """DOCX с inline-картинкой между двумя абзацами."""
    import io

    from docx import Document

    doc = Document()
    doc.add_paragraph("Перед картинкой")
    p = doc.add_paragraph()
    p.add_run().add_picture(io.BytesIO(image_bytes or _tiny_png()))
    doc.add_paragraph("После картинки")
    doc.save(str(path))
    return path


def inject_comment(path: Path, author: str = "Иван Иванов", text: str = "Включите Trunk для VLAN 10-20") -> Path:
    entries = _read_zip(path)
    docxml = _xml(entries["word/document.xml"])
    body = docxml.find(f"{{{W}}}body")

    target = None
    for p in body.iter(f"{{{W}}}p"):
        if p.findall(f".//{{{W}}}t"):
            target = p
            break
    assert target is not None, "нет абзаца с текстом для комментария"

    cstart = etree.SubElement(target, f"{{{W}}}commentRangeStart")
    cstart.set(f"{{{W}}}id", "0")
    run = etree.SubElement(target, f"{{{W}}}r")
    cref = etree.SubElement(run, f"{{{W}}}commentReference")
    cref.set(f"{{{W}}}id", "0")
    cend = etree.SubElement(target, f"{{{W}}}commentRangeEnd")
    cend.set(f"{{{W}}}id", "0")
    # порядок: cstart в начале, cend+ref в конце
    target.remove(cstart)
    target.insert(0, cstart)
    target.remove(cend)
    target.append(cend)

    entries["word/document.xml"] = _dump(docxml)
    entries["word/comments.xml"] = _comments_xml(author, text)
    _add_relationship(
        entries,
        "word/_rels/document.xml.rels",
        "rIdComments99",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
        "comments.xml",
    )
    _add_content_type(
        entries,
        "/word/comments.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
    )
    _write_zip(path, entries)
    return path


def _comments_xml(author: str, text: str) -> bytes:
    root = etree.Element(f"{{{W}}}comments")
    comment = etree.SubElement(root, f"{{{W}}}comment")
    comment.set(f"{{{W}}}id", "0")
    comment.set(f"{{{W}}}author", author)
    comment.set(f"{{{W}}}initials", "ИИ")
    comment.set(f"{{{W}}}date", "2026-08-12T10:00:00Z")
    p = etree.SubElement(comment, f"{{{W}}}p")
    r = etree.SubElement(p, f"{{{W}}}r")
    t = etree.SubElement(r, f"{{{W}}}t")
    t.text = text
    return _dump(root)


def make_docx_with_embedded_xlsx(path: Path, payload: bytes, prog_id: str = "Excel.Sheet.12") -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Перед встроенной таблицей")
    doc.save(str(path))

    entries = _read_zip(path)
    entries["word/embeddings/embedded.xlsx"] = payload

    docxml = _xml(entries["word/document.xml"])
    body = docxml.find(f"{{{W}}}body")
    w_p = etree.SubElement(body, f"{{{W}}}p")
    w_r = etree.SubElement(w_p, f"{{{W}}}r")
    w_obj = etree.SubElement(w_r, f"{{{W}}}object")
    docpr = etree.SubElement(w_obj, f"{{{W}}}docPr")
    docpr.set("id", "99")
    docpr.set("name", "Объект Excel")
    ole = etree.SubElement(w_obj, f"{{{O}}}OLEObject")
    ole.set("Type", "Embed")
    ole.set("ProgID", prog_id)
    ole.set("ShapeID", "shp99")
    ole.set(f"{{{R}}}id", "rIdEmbed99")
    entries["word/document.xml"] = _dump(docxml)

    _add_relationship(
        entries,
        "word/_rels/document.xml.rels",
        "rIdEmbed99",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject",
        "embeddings/embedded.xlsx",
    )
    _add_content_type(
        entries,
        "/word/embeddings/embedded.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    _write_zip(path, entries)
    return path


# ---------------------------------------------------------------- xlsx
def make_xlsx(path: Path, sheet: str = "Лист1", rows: list[list[object]] | None = None) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows or [["Номер VLAN", "Назначение"], [10, "Management"], [20, "DMZ"]]:
        ws.append(row)
    wb.save(str(path))
    return path


def xlsx_bytes(sheet: str = "Лист1", rows: list[list[object]] | None = None) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows or [["Код", "Описание"], [1, "Правило одно"], [2, "Правило два"]]:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- pdf
def make_pdf(path: Path, text: str = "Пример текста из pdf документа") -> Path:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    font = "Helvetica"
    for name in ("arial.ttf", "Vera.ttf"):
        try:
            pdfmetrics.registerFont(TTFont("DocParserFont", name))
            font = "DocParserFont"
            break
        except Exception:
            continue
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont(font, 12)
    c.drawString(72, 720, text)
    c.save()
    return path


def make_pdf_with_image(path: Path, image_bytes: bytes | None = None) -> Path:
    """PDF со страницей, содержащей текст и растровое изображение."""
    import io

    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    data = image_bytes or _tiny_png()
    c = canvas.Canvas(str(path), pagesize=letter)
    c.drawString(72, 720, "Текст страницы с картинкой")
    c.drawImage(ImageReader(io.BytesIO(data)), 72, 600, width=50, height=50)
    c.save()
    return path

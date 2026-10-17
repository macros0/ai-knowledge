"""The parser must retain evidence beyond the former aggregate text cutoff."""
from email.message import EmailMessage
from pathlib import Path
from random import Random
from string import ascii_letters, digits

import pytest
from docparser import parse_document_result
from docx import Document
from openpyxl import Workbook


@pytest.fixture(scope="module")
def large_body():
    # Compressible repeated characters can trigger the independent ZIP guard.
    return "".join(Random(42).choices(ascii_letters + digits, k=2_000_100)) + "TAILEVIDENCE"


@pytest.mark.parametrize("suffix", [".docx", ".xlsx", ".eml"])
def test_full_text_including_tail_survives_root_dispatch(tmp_path, large_body, suffix):
    path = tmp_path / ("full" + suffix)
    if suffix == ".docx":
        doc = Document()
        doc.add_paragraph(large_body)
        doc.save(path)
        expected_parts = [large_body]
    elif suffix == ".xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Evidence"])
        expected_parts = [large_body[start:start + 20_000] for start in range(0, len(large_body), 20_000)]
        for part in expected_parts:
            sheet.append([part])
        workbook.save(path)
        workbook.close()
    else:
        mail = EmailMessage()
        mail["From"] = "sender@example.test"
        mail.set_content(large_body)
        path.write_bytes(mail.as_bytes())
        expected_parts = [large_body]

    result = parse_document_result(path)
    text = "\n".join(block.text for block in result.blocks)

    assert all(part in text for part in expected_parts), "Source text or its tail was dropped"
    assert not any(warning["code"] == "text_limit_exceeded" for warning in result.warnings)


def test_aggregate_mail_text_retains_later_children_and_original_bytes(tmp_path, large_body):
    mail = EmailMessage()
    mail["From"] = "sender@example.test"
    mail.set_content(large_body)
    child = EmailMessage()
    child["From"] = "child@example.test"
    child.set_content("LATERCHILDEVIDENCE")
    child_bytes = child.as_bytes()
    mail.add_attachment(child_bytes, maintype="application", subtype="octet-stream", filename="child.eml")
    mail.add_attachment(b"original proof", maintype="application", subtype="octet-stream", filename="proof.bin")
    path = tmp_path / "aggregate.eml"
    path.write_bytes(mail.as_bytes())

    result = parse_document_result(path, attachments_dir=tmp_path / "attachments")
    text = "\n".join(block.text for block in result.blocks)
    originals = {block.meta["name"]: block for block in result.blocks if block.type == "attachment"}

    assert large_body in text, "Root text tail was dropped"
    assert "LATERCHILDEVIDENCE" in text, "Later source was dropped"
    assert originals["child.eml"].text
    assert originals["proof.bin"].text
    assert Path(originals["child.eml"].meta["saved_path"]).read_bytes() == child_bytes
    assert Path(originals["proof.bin"].meta["saved_path"]).read_bytes() == b"original proof"
    assert [source.source_id for source in result.sources] == ["root", "root/0", "root/1"]
    assert not any(warning["code"] == "text_limit_exceeded" for warning in result.warnings)

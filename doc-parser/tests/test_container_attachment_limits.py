"""Container admission must precede payload reads, including tiny attachments."""
import io
import zipfile
from pathlib import Path

import pytest
from docparser import parse_document, parse_document_result
from docparser.embedded import AttachmentBudget
from docparser.source_model import ParseContext
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import EmbeddedFile, NumberObject

from tests.mail_fixtures import unicode_msg


@pytest.mark.parametrize("name", ["message.eml", "message.msg", "sheet.xlsx", "message.BIN", "message"])
def test_xlsx_recurses_into_embedded_files_regardless_of_suffix(tmp_path, name):
    path = tmp_path / "parent.xlsx"
    workbook = Workbook()
    workbook.active.append(["PARENT TEXT"])
    workbook.save(path)
    expected = "CHILD EVIDENCE"
    if name.endswith(".msg"):
        payload = unicode_msg(body=expected)
    elif name.endswith(".xlsx"):
        child = Workbook()
        child.active.append([expected])
        output = io.BytesIO()
        child.save(output)
        payload = output.getvalue()
    else:
        payload = f"From: sender@example.test\r\n\r\n{expected}\r\n".encode()
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("xl/embeddings/", b"")
        archive.writestr(f"xl/embeddings/{name}", payload)
    result = parse_document_result(path, attachments_dir=tmp_path / "attachments")
    assert [(source.source_id, source.parent_source_id) for source in result.sources] == [
        ("root", None), ("root/0", "root"),
    ]
    assert any(expected in block.text and block.meta.get("source_id") == "root/0" for block in result.blocks)
    assert any("PARENT TEXT" in block.text for block in result.blocks)
    marker = next(block for block in result.blocks if block.type == "attachment")
    assert Path(marker.meta["saved_path"]).read_bytes() == payload
    assert result.warnings == []


@pytest.mark.parametrize("container,declare_size", [("xlsx", True), ("pdf", True), ("pdf", False)])
@pytest.mark.parametrize("limit", ["nodes", "bytes"])
def test_container_admission_precedes_payload_reads(tmp_path, monkeypatch, container, declare_size, limit):
    path = tmp_path / f"many.{container}"
    payload = b"From: sender@example.test\r\nSubject: Item\r\n\r\nUNIQUE BODY\r\n"
    reads = []
    if container == "xlsx":
        workbook = Workbook()
        workbook.active.append(["PARENT TEXT"])
        workbook.save(path)
        with zipfile.ZipFile(path, "a") as archive:
            for index in range(8):
                archive.writestr(f"xl/embeddings/item-{index}.bin", payload)
        original_read = zipfile.ZipFile.read

        def read(archive, name, *args, **kwargs):
            member = name.filename if isinstance(name, zipfile.ZipInfo) else name
            if member.startswith("xl/embeddings/"):
                reads.append(member)
            return original_read(archive, name, *args, **kwargs)

        monkeypatch.setattr(zipfile.ZipFile, "read", read)
    else:
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        for index in range(8):
            attachment = writer.add_attachment(f"item-{index}.eml", payload)
            if declare_size:
                attachment.size = NumberObject(len(payload))
        writer.write(path)
        original_content = EmbeddedFile.content

        def content(attachment):
            reads.append(attachment.name)
            return original_content.fget(attachment)

        monkeypatch.setattr(EmbeddedFile, "content", property(content))

    context = ParseContext(path.name)
    budget = AttachmentBudget(max_nodes=2) if limit == "nodes" else AttachmentBudget(total=0)
    blocks = parse_document(path, context=context, budget=budget)

    if limit == "nodes":
        assert len(reads) == 2
        assert len(context.sources) == 4  # Root, two admitted files, one remainder.
        assert sum(block.meta.get("extraction_status") == "skipped_count" for block in blocks) == 1
        assert {"code": "attachment_count_exceeded", "source_id": "root/2"} in context.warnings
        assert sum(block.text == "UNIQUE BODY" for block in blocks) == 2
    else:
        assert reads == []
        assert len(context.sources) == 9
        assert sum(block.meta.get("extraction_status") == "skipped_size" for block in blocks) == 8
        assert all(warning["code"] == "attachment_size_exceeded" for warning in context.warnings)
    if container == "xlsx":
        assert any("PARENT TEXT" in block.text for block in blocks)


@pytest.mark.parametrize("container", ["xlsx", "pdf"])
def test_broken_container_attachment_keeps_readable_sibling(tmp_path, container):
    path = tmp_path / f"broken-child.{container}"
    payload = b"From: sender@example.test\r\n\r\nSIBLING BODY\r\n"
    if container == "xlsx":
        workbook = Workbook()
        workbook.active.append(["PARENT TEXT"])
        workbook.save(path)
        with zipfile.ZipFile(path, "a") as archive:
            archive.writestr("xl/embeddings/broken.bin", b"broken bytes")
            archive.writestr("xl/embeddings/good.bin", payload)
            broken = archive.getinfo("xl/embeddings/broken.bin")
            offset = broken.header_offset + 30 + len(broken.filename.encode()) + len(broken.extra)
        raw = bytearray(path.read_bytes())
        raw[offset] ^= 1  # Preserve ZIP structure, invalidate only this member's CRC.
        path.write_bytes(raw)
    else:
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        broken = writer.add_attachment("broken.eml", payload)
        del broken.pdf_object["/EF"]
        writer.add_attachment("good.eml", payload)
        writer.write(path)

    context = ParseContext(path.name)
    blocks = parse_document(path, context=context)
    assert len(context.sources) == 3
    assert any(block.text == "SIBLING BODY" for block in blocks)
    assert sum(block.meta.get("extraction_status") == "skipped_parse" for block in blocks) == 1
    assert {"code": "attachment_parse_failed", "source_id": "root/0"} in context.warnings
    if container == "xlsx":
        assert any("PARENT TEXT" in block.text for block in blocks)

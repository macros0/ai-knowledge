"""Разбор XLSX: листы → Markdown-таблицы + встроенные объекты (xl/embeddings)."""
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from docparser.archive_guard import validate_member, validate_zip
from docparser.blocks import Block
from docparser.embedded import process_embedded


def parse_xlsx(
    path: str | Path,
    attachments_dir: str | Path | None = None,
    depth: int = 0,
    budget=None,
) -> list[Block]:
    # Validate the central directory before openpyxl starts reading XML parts.
    validate_zip(path)
    wb = load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[Block] = []
    try:
        for ws in wb.worksheets:
            blocks.append(Block("heading", f"Таблица: {ws.title}", level=1))
            md = _sheet_to_markdown(ws)
            if md:
                blocks.append(Block("table", md))
    finally:
        wb.close()

    for idx, (name, data) in enumerate(_embedded_files(path)):
        blocks.extend(process_embedded(data, name, "", "", attachments_dir, idx, depth=depth + 1, budget=budget))

    return blocks


def _sheet_to_markdown(ws) -> str:
    rows = []
    cell_count = 0
    max_rows = 250_000
    max_cells = 5_000_000
    for row in ws.iter_rows(values_only=True):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        if len(rows) >= max_rows or cell_count + len(row) > max_cells:
            raise ValueError("spreadsheet row/cell safety limit exceeded")
        cell_count += len(row)
        rows.append(["" if v is None else str(v).replace("\n", "<br>") for v in row])
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    for r in rows[1:]:
        lines.append("| " + " | ".join(r) + " |")
    return "\n".join(lines)


def _embedded_files(path: str | Path) -> list[tuple[str, bytes]]:
    results: list[tuple[str, bytes]] = []
    try:
        with zipfile.ZipFile(str(path)) as zf:
            for name in zf.namelist():
                if name.startswith("xl/embeddings/") and name.endswith(".bin"):
                    validate_member(zf.getinfo(name))
                    results.append((Path(name).name, zf.read(name)))
    except Exception:
        pass
    return results

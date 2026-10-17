"""Conservative header geometry; ambiguity never deletes physical rows."""

from __future__ import annotations

from dataclasses import dataclass, replace
from typing import Iterable

from openpyxl.utils.cell import get_column_letter

from docparser.table_model import CellAddress, CellRange, TableColumn, TableDescriptor, TableRow


@dataclass(frozen=True)
class HeaderDecision:
    header_rows: list[int]
    columns: list[TableColumn]
    method: str
    warnings: list[str]


def detect_table_regions(
    sheet_index: int,
    declared_ranges: list[CellRange],
    occupied_cells: Iterable[CellAddress],
    merges: list[CellRange],
) -> list[CellRange]:
    """Declared tables plus remaining contiguous column regions.

    A blank row alone cannot split a table. Blank columns split only if no
    merge bridges them. Store per-column bounds, never every occupied address.
    """
    bounds = {}
    for cell in occupied_cells:
        if any(r.contains(cell.row, cell.column) for r in declared_ranges):
            continue
        low, high = bounds.get(cell.column, (cell.row, cell.row))
        bounds[cell.column] = (min(low, cell.row), max(high, cell.row))
    for merge in merges:
        if any(r.contains(merge.min_row, merge.min_column) for r in declared_ranges):
            continue
        if merge.min_column not in bounds:
            continue
        for column in range(merge.min_column, merge.max_column + 1):
            low, high = bounds.get(column, (merge.min_row, merge.max_row))
            bounds[column] = (min(low, merge.min_row), max(high, merge.max_row))
    result = list(declared_ranges)
    group = []
    for column in sorted(bounds):
        if group and column != group[-1] + 1:
            result.append(
                CellRange(
                    sheet_index,
                    min(bounds[c][0] for c in group),
                    group[0],
                    max(bounds[c][1] for c in group),
                    group[-1],
                )
            )
            group = []
        group.append(column)
    if group:
        result.append(
            CellRange(
                sheet_index,
                min(bounds[c][0] for c in group),
                group[0],
                max(bounds[c][1] for c in group),
                group[-1],
            )
        )
    # Subtract declared rectangles from residual bounds. Cells above and below
    # a declared table must not make its data appear twice.
    residual = result[len(declared_ranges) :]
    for cut in declared_ranges:
        parts = []
        for area in residual:
            top, bottom = max(area.min_row, cut.min_row), min(area.max_row, cut.max_row)
            left, right = max(area.min_column, cut.min_column), min(area.max_column, cut.max_column)
            if top > bottom or left > right:
                parts.append(area)
                continue
            for r1, c1, r2, c2 in (
                (area.min_row, area.min_column, top - 1, area.max_column),
                (bottom + 1, area.min_column, area.max_row, area.max_column),
                (top, area.min_column, bottom, left - 1),
                (top, right + 1, bottom, area.max_column),
            ):
                if r1 <= r2 and c1 <= c2:
                    parts.append(CellRange(sheet_index, r1, c1, r2, c2))
        residual = parts
    if len(residual) + len(declared_ranges) > 1000:
        from docparser.archive_guard import ArchiveLimitError

        raise ArchiveLimitError("spreadsheet table safety limit exceeded")
    return list(declared_ranges) + residual


def _text_header(row):
    cells = [c for c in row.cells.values() if c.display_text and c.state == "value"]
    return len(cells) >= 2 and all(c.value_type == "string" and c.formula is None for c in cells)


def combine_numbered_regions(areas, sample_rows, merges, declared_ranges):
    """A shared ordinal row below merged headers can bridge blank columns.

    Numbering is evidence of a shared table, never a replacement for physical
    coordinates. Occasional mistakes in the printed numbering stay unchanged.
    Explicit Excel tables always retain their own boundaries.
    """
    residual = [area for area in areas if area not in declared_ranges]
    for row in sample_rows:
        cells = [
            (col, cell)
            for col, cell in sorted(row.cells.items())
            if cell.display_text and any(a.contains(row.row_number, col) for a in residual)
        ]
        if len(cells) < 6 or any(
            cell.value_type != "number"
            or cell.formula is not None
            or not isinstance(cell.raw_value, int)
            for _, cell in cells
        ):
            continue
        numbers = [cell.raw_value for _, cell in cells]
        # Ordinal labels must start at one, end at the occupied width and agree
        # at least 90% with that sequence; geometry supplies independent evidence.
        if (
            numbers[0] != 1
            or numbers[-1] != len(numbers)
            or sum(number == position for position, number in enumerate(numbers, 1)) * 10
            < len(numbers) * 9
        ):
            continue
        group = [
            area
            for area in residual
            if area.min_row <= row.row_number <= area.max_row
            and any(
                m.min_row < row.row_number
                and m.max_row == row.row_number - 1
                and m.max_row > m.min_row
                and area.min_column <= m.min_column <= area.max_column
                for m in merges
            )
        ]
        if len(group) < 2 or len({a.max_row for a in group}) != 1:
            continue
        combined = CellRange(
            group[0].sheet_index,
            min(a.min_row for a in group),
            min(a.min_column for a in group),
            group[0].max_row,
            max(a.max_column for a in group),
        )
        if any(
            not (a.max_column < combined.min_column or a.min_column > combined.max_column)
            for a in areas
            if a not in group
        ):
            continue
        areas = [area for area in areas if area not in group] + [combined]
        residual = [area for area in areas if area not in declared_ranges]
    return areas


def split_stacked_regions(area, rows, merges):
    """Split only where header geometry or a style/type boundary agrees.

    A gap or bold text by itself is insufficient. A full-width section title
    followed by labels can also start a table without a blank row.
    """
    from collections import deque

    source = (
        (
            replace(
                row,
                cells={
                    col: cell
                    for col, cell in row.cells.items()
                    if area.min_column <= col <= area.max_column
                },
            ),
            bool(bold_columns.intersection(range(area.min_column, area.max_column + 1))),
        )
        for row, bold_columns in rows
        if area.min_row <= row.row_number <= area.max_row
    )
    source = (
        (row, bold)
        for row, bold in source
        if any(cell.display_text or cell.formula for cell in row.cells.values())
    )
    window = deque()
    for _ in range(3):
        item = next(source, None)
        if item is not None:
            window.append(item)
    result, start, low, high, previous = [], area.min_row, None, None, None
    while window:
        row, bold = window[0]
        following, following_bold = window[1] if len(window) > 1 else (None, False)
        gap = previous is not None and row.row_number > previous + 1
        row_merges = [
            m
            for m in merges
            if m.min_row == row.row_number
            and area.min_column <= m.min_column <= m.max_column <= area.max_column
        ]
        full_title = sum(bool(c.display_text) for c in row.cells.values()) == 1 and any(
            m.min_column == area.min_column
            and m.max_column == area.max_column
            and m.max_row == m.min_row
            for m in row_merges
        )
        lower_numeric = following is not None and any(
            c.state == "value" and c.value_type != "string" for c in following.cells.values()
        )
        merged_header = any(
            m.max_row > m.min_row or m.max_column > m.min_column for m in row_merges
        )
        new_header = _text_header(row) and (merged_header or bold and lower_numeric)
        single_label = sum(bool(c.display_text) for c in row.cells.values()) == 1
        merged_caption = single_label and any(
            m.max_column > m.min_column and m.max_row == m.min_row for m in row_merges
        )
        section = (
            single_label
            and following is not None
            and _label_row(following)
            and (merged_caption or following_bold)
        )
        crossed_merge = any(
            m.min_row < row.row_number <= m.max_row
            and m.min_column <= area.max_column
            and m.max_column >= area.min_column
            for m in merges
        )
        section_after_body = full_title and section and row.row_number > start + 8
        if (gap and (new_header or section) or section_after_body) and not crossed_merge:
            result.append(CellRange(area.sheet_index, start, low, row.row_number - 1, high))
            start, low, high = row.row_number, None, None
            if len(result) >= 1000:
                from docparser.archive_guard import ArchiveLimitError

                raise ArchiveLimitError("spreadsheet table safety limit exceeded")
        for column, cell in row.cells.items():
            if cell.display_text or cell.formula:
                low = column if low is None else min(low, column)
                high = column if high is None else max(high, column)
        for merge in row_merges:
            low = merge.min_column if low is None else min(low, merge.min_column)
            high = merge.max_column if high is None else max(high, merge.max_column)
        previous = row.row_number
        window.popleft()
        item = next(source, None)
        if item is not None:
            window.append(item)
    if low is not None:
        result.append(CellRange(area.sheet_index, start, low, area.max_row, high))
    return result or [area]


def _label_row(row):
    cells = [
        cell for _, cell in sorted(row.cells.items()) if cell.display_text and cell.state == "value"
    ]
    strings = sum(cell.value_type == "string" for cell in cells)
    return (
        bool(cells)
        and cells[0].value_type == "string"
        and strings >= 2
        and strings * 2 >= len(cells)
        and all(cell.formula is None and cell.value_type in {"string", "number"} for cell in cells)
    )


def normalize_headers(
    descriptor: TableDescriptor,
    sample_rows: list[TableRow],
    merges: list[CellRange],
    *,
    bold_rows=frozenset(),
    declared_header_count=1,
    declared_column_names=(),
) -> HeaderDecision:
    area = descriptor.cell_range
    if descriptor.declared and declared_header_count == 0:
        # Declared names are metadata, not cells on the first physical row.
        columns = []
        for column in descriptor.columns:
            offset = column.index - area.min_column
            name = declared_column_names[offset] if offset < len(declared_column_names) else ""
            columns.append(replace(column, header_path=[name or column.letter], header_addresses=[]))
        return HeaderDecision([], columns, "explicit", [])
    sample = [
        r
        for r in sample_rows
        if any(c.state == "value" and c.display_text for c in r.cells.values())
    ][:8]
    by_row = {r.row_number: r for r in sample}
    candidates = list(sample)

    def level_merges(row_number):
        return [
            m
            for m in merges
            if m.min_row == row_number
            and area.min_column <= m.min_column <= m.max_column <= area.max_column
            and (m.max_column > m.min_column or m.max_row > m.min_row)
        ]

    def merged_title(row):
        cells = [col for col, cell in row.cells.items() if cell.display_text]
        return len(cells) == 1 and any(
            m.min_column == cells[0] and m.max_column > m.min_column and m.max_row == m.min_row
            for m in level_merges(row.row_number)
        )

    def header_width(row):
        indices = {col for col, cell in row.cells.items() if cell.display_text}
        for merge in level_merges(row.row_number):
            indices.update(range(merge.min_column, merge.max_column + 1))
        return len(indices)

    while len(candidates) >= 2:
        first, following = candidates[:2]
        sparse_preamble = header_width(first) * 2 < len(descriptor.columns) and header_width(
            following
        ) * 2 >= len(descriptor.columns)
        discard_title = merged_title(first) and (
            merged_title(following) or first.row_number not in bold_rows
        )
        if not (sparse_preamble or discard_title):
            break
        candidates.pop(0)
    def numbered_leaf(first, leaf, geometry):
        # A numeric-leading leaf is otherwise indistinguishable from data.
        # Require a merged key header, mixed named leaves and the immediately
        # following independent ordinal row; numbering stays physical data.
        cells = [cell for _, cell in sorted(leaf.cells.items()) if cell.display_text]
        strings = sum(cell.value_type == "string" for cell in cells)
        numbered = by_row.get(leaf.row_number + 1)
        ordinals = [cell for _, cell in sorted(numbered.cells.items()) if cell.display_text] if numbered else []
        leading = min(first.cells) if first.cells else 0
        return bool(
            cells and cells[0].value_type == "number" and strings >= 2
            and strings * 2 >= len(cells)
            and all(cell.formula is None and cell.value_type in {"string", "number"} for cell in cells)
            and any(m.min_column == leading and m.max_column == leading
                    and m.min_row == first.row_number and m.max_row >= leaf.row_number
                    for m in geometry)
            and len(ordinals) >= 6
            and all(cell.value_type == "number" and cell.formula is None
                    and isinstance(cell.raw_value, int) for cell in ordinals)
            and ordinals[0].raw_value == 1 and ordinals[-1].raw_value == len(ordinals)
            and sum(cell.raw_value == number for number, cell in enumerate(ordinals, 1)) * 10
                >= len(ordinals) * 9
        )

    headers, method = [], "unresolved"
    if descriptor.declared and area.min_row in by_row:
        headers, method = [area.min_row], "explicit"
    elif len(candidates) >= 2:
        first, following = candidates[:2]
        first_cells = [cell for _, cell in sorted(first.cells.items()) if cell.display_text]
        following_cells = [cell for _, cell in sorted(following.cells.items()) if cell.display_text]
        leading_boundary = (
            first_cells
            and following_cells
            and first_cells[0].value_type == "string"
            and following_cells[0].value_type != "string"
        )
        textual = first_cells and all(
            cell.value_type == "string" and cell.formula is None for cell in first_cells
        )
        if textual or _label_row(first) and leading_boundary:
            geometry = level_merges(first.row_number)
            has_body_boundary = any(
                any(
                    cell.state == "value" and cell.value_type != "string"
                    for cell in row.cells.values()
                )
                for row in candidates[1:]
            )
            caption_leaf = (
                merged_title(first)
                and _text_header(following)
                and not level_merges(following.row_number)
            )
            bounded_geometry = geometry and max(m.max_row for m in geometry) <= first.row_number + 7
            if (
                bounded_geometry
                and (has_body_boundary or caption_leaf)
                and (
                    max(m.max_row for m in geometry) > first.row_number
                    and (
                        following.row_number > max(m.max_row for m in geometry)
                        or _label_row(following)
                        or numbered_leaf(first, following, geometry)
                    )
                    or _text_header(following)
                )
            ):
                last_header = max(
                    [m.max_row for m in geometry]
                    + [
                        following.row_number
                        if any(m.max_column > m.min_column for m in geometry)
                        and (
                            _text_header(following)
                            or following.row_number <= max(m.max_row for m in geometry)
                        )
                        else first.row_number
                    ]
                )
                while last_header <= first.row_number + 7:
                    lower_merges = [
                        m for m in level_merges(last_header) if m.max_column > m.min_column
                    ]
                    lower = next((row for row in candidates if row.row_number > last_header), None)
                    if not lower_merges or lower is None or not _text_header(lower):
                        break
                    last_header = max(lower.row_number, max(m.max_row for m in lower_merges))
                signatures = [
                    tuple(
                        sorted(
                            (m.min_column, m.max_column)
                            for m in level_merges(row.row_number)
                            if m.max_column > m.min_column
                        )
                    )
                    for row in candidates
                    if row.row_number <= last_header
                ]
                repeated_groups = any(
                    left and left == right for left, right in zip(signatures, signatures[1:])
                )
                merged_leaf = any(m.max_column > m.min_column for m in level_merges(last_header))
                if last_header <= first.row_number + 7 and not (repeated_groups and merged_leaf):
                    headers = list(range(first.row_number, last_header + 1))
                    method = "structural"
            elif not geometry and (
                leading_boundary
                or textual
                and any(
                    cell.state == "value" and cell.value_type != "string"
                    for cell in following.cells.values()
                )
            ):
                headers, method = [first.row_number], "structural"
    columns = []
    for column in descriptor.columns:
        labels, addresses = [], []
        for row_number in headers:
            row = by_row.get(row_number)
            cell = row.cells.get(column.index) if row else None
            merge = next(
                (
                    m
                    for m in merges
                    if m.contains(row_number, column.index) and m.min_row in headers
                ),
                None,
            )
            if merge:
                anchor_row = by_row.get(merge.min_row)
                cell = anchor_row.cells.get(merge.min_column) if anchor_row else None
            if cell and cell.display_text and (not labels or labels[-1] != cell.display_text):
                labels.append(cell.display_text)
                addresses.append(cell.address)
        columns.append(
            replace(
                column,
                header_path=labels or [get_column_letter(column.index)],
                header_addresses=addresses,
            )
        )
    return HeaderDecision(
        headers, columns, method, ["header_ambiguous"] if method == "unresolved" else []
    )
